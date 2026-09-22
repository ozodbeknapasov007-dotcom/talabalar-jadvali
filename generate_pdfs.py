import os
import openpyxl
from reportlab import rl_config
# Ma'lumot o'zgarmagan bo'lsa PDF baytlari ham o'zgarmasligi shart.
# Aks holda ReportLab har safar yangi CreationDate/ModDate va /ID yozadi,
# natijada har bir qayta generatsiyada 16 ta PDF "o'zgargan" bo'lib
# git ga tushadi va har bir commit ~800 KB keraksiz ma'lumot bilan shishadi.
rl_config.invariant = 1
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# qayta_tekshiruv ichida bo'lsa, asosiy loyiha papkasi bir pog'ona yuqorida:
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(CURRENT_DIR) if os.path.basename(CURRENT_DIR) == 'qayta_tekshiruv' else CURRENT_DIR

# Fontlarni ro'yxatdan o'tkazish
try:
    pdfmetrics.registerFont(TTFont('Times', 'C:/Windows/Fonts/times.ttf'))
    pdfmetrics.registerFont(TTFont('Times-Bold', 'C:/Windows/Fonts/timesbd.ttf'))
    FONT_NORMAL = 'Times'
    FONT_BOLD = 'Times-Bold'
except Exception as e:
    FONT_NORMAL = 'Helvetica'
    FONT_BOLD = 'Helvetica-Bold'

GROUP_LEADERS = {
    "26-01": "Mirzayeva.D",
    "26-02": "Ochilov.D",
    "26-03": "To'rayeva.S",
    "26-04": "Hamdamova.M",
    "26-05": "Rayimova.X",
    "26-06": "Yuldashev.O",
    "26-07": "Asraliyev.A"
}

GROUP_TITLES = {
    "26-01": "Farmatsiya ishi",
    "26-02": "Hamshiralik ishi",
    "26-03": "Hamshiralik ishi",
    "26-04": "Hamshiralik ishi",
    "26-05": "Hamshiralik ishi",
    "26-06": "Hamshiralik ishi",
    "26-07": "Hamshiralik ishi"
}

def build_group_flowables(group_code, students, avail_width):
    g_students = [s for s in students if s.get('group') == group_code]
    g_students.sort(key=lambda x: str(x.get('ism', '')).lower())

    leader = GROUP_LEADERS.get(group_code, "—")
    g_title = GROUP_TITLES.get(group_code, "")

    elements = []
    count = len(g_students)

    # Talabalar soniga qarab optimal, katta va chiroyli shrift o'lchamlari
    if count <= 20:
        row_height = 25.0
        font_size = 10.5
        header_font_size = 11.0
        padding_size = 3.5
    elif count <= 27:
        row_height = 23.0
        font_size = 10.0
        header_font_size = 10.5
        padding_size = 3.0
    else:
        row_height = 22.0
        font_size = 10.0
        header_font_size = 10.5
        padding_size = 2.8

    title_style = ParagraphStyle(
        'DocTitle',
        fontName=FONT_BOLD,
        fontSize=16,
        leading=20,
        alignment=1,
        textColor=colors.HexColor('#0f172a')
    )

    info_left_style = ParagraphStyle(
        'InfoLeft',
        fontName=FONT_BOLD,
        fontSize=11,
        leading=14,
        alignment=0,
        textColor=colors.HexColor('#1e293b')
    )

    info_right_style = ParagraphStyle(
        'InfoRight',
        fontName=FONT_BOLD,
        fontSize=11,
        leading=14,
        alignment=2,
        textColor=colors.HexColor('#1e293b')
    )

    header_para_style = ParagraphStyle(
        'HeaderCell',
        fontName=FONT_BOLD,
        fontSize=header_font_size,
        leading=header_font_size + 2,
        alignment=1,
        textColor=colors.HexColor('#0f172a')
    )

    cell_name_style = ParagraphStyle(
        'CellName',
        fontName=FONT_BOLD,
        fontSize=font_size,
        leading=font_size + 2,
        alignment=0,
        textColor=colors.HexColor('#0f172a')
    )

    cell_center_style = ParagraphStyle(
        'CellCenter',
        fontName=FONT_NORMAL,
        fontSize=font_size,
        leading=font_size + 2,
        alignment=1,
        textColor=colors.HexColor('#334155')
    )

    header_table_data = [
        [
            Paragraph("Shahrisabz Tibbiyot Texnikumi", title_style),
            ""
        ],
        [
            Paragraph(f"Yo'nalish: <b>{g_title}</b> | Guruh: <b>{group_code}</b>", info_left_style),
            Paragraph(f"Guruh rahbari: <b>{leader}</b>", info_right_style)
        ]
    ]

    ht = Table(header_table_data, colWidths=[avail_width * 0.62, avail_width * 0.38])
    ht.setStyle(TableStyle([
        ('SPAN', (0, 0), (1, 0)),
        ('ALIGN', (0, 0), (1, 0), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 1), (-1, 1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(ht)

    w_no = 26
    w_sh = 65
    w_dob = 80
    w_imzo = 95
    w_name = avail_width - (w_no + w_sh + w_dob + w_imzo)

    col_widths = [w_no, w_sh, w_name, w_dob, w_imzo]

    headers = [
        Paragraph("№", header_para_style),
        Paragraph("Shartnoma", header_para_style),
        Paragraph("F.I.SH (To'liq)", header_para_style),
        Paragraph("Tug'ilgan sana", header_para_style),
        Paragraph("Imzo", header_para_style)
    ]

    data = [headers]

    for idx, st in enumerate(g_students, start=1):
        sh_val = str(st.get('shnum', '')).strip()
        if sh_val and sh_val not in ['—', '-', 'None', 'nan']:
            sh_num = f"#{sh_val}" if not sh_val.startswith('#') else sh_val
        else:
            sh_num = "—"
        row = [
            Paragraph(str(idx), cell_center_style),
            Paragraph(sh_num, cell_center_style),
            Paragraph(st.get('fio', st.get('ism', '')), cell_name_style),
            Paragraph(st.get('dob') or "—", cell_center_style),
            Paragraph("", cell_center_style)
        ]
        data.append(row)

    t = Table(data, colWidths=col_widths, rowHeights=[row_height + 3] + [row_height] * len(g_students))
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), padding_size),
        ('BOTTOMPADDING', (0, 0), (-1, -1), padding_size),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.6, colors.HexColor('#94a3b8')),
        ('BOX', (0, 0), (-1, -1), 1.0, colors.HexColor('#334155')),
        ('LINEBELOW', (0, 0), (-1, 0), 1.2, colors.HexColor('#0f172a')),
    ]

    for row_idx in range(1, len(data)):
        if row_idx % 2 == 0:
            t_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor('#f8fafc')))

    t.setStyle(TableStyle(t_style))
    elements.append(t)
    return elements

def create_single_group_pdf(group_code, students, output_pdf_path):
    os.makedirs(os.path.dirname(output_pdf_path), exist_ok=True)
    doc = SimpleDocTemplate(
        output_pdf_path,
        pagesize=A4,
        leftMargin=24,
        rightMargin=24,
        topMargin=20,
        bottomMargin=16
    )
    avail_width = A4[0] - 48
    elements = build_group_flowables(group_code, students, avail_width)
    doc.build(elements)
    return output_pdf_path

def create_all_groups_combined_pdf(groups, students, output_pdf_path):
    """
    Barcha guruhlarni 1 ta umumiy PDF hujjatga birlashtiradi.
    Har bir guruh navbati bilan alohida-alohida A4 varoqda chiqadi (PageBreak orqali).
    """
    os.makedirs(os.path.dirname(output_pdf_path), exist_ok=True)
    doc = SimpleDocTemplate(
        output_pdf_path,
        pagesize=A4,
        leftMargin=24,
        rightMargin=24,
        topMargin=20,
        bottomMargin=16
    )
    avail_width = A4[0] - 48
    elements = []
    for idx, g in enumerate(groups):
        g_elements = build_group_flowables(g, students, avail_width)
        elements.extend(g_elements)
        if idx < len(groups) - 1:
            elements.append(PageBreak())
    doc.build(elements)
    return output_pdf_path

def build_all_group_pdfs():
    excel_path = os.path.join(BASE_DIR, 'Talabalar_Toliq_Royxati.xlsx')
    if not os.path.exists(excel_path):
        return {}

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    students = []
    for r in range(2, ws.max_row + 1):
        ism = str(ws.cell(row=r, column=2).value or '').strip()
        shnum = str(ws.cell(row=r, column=5).value or '').strip()
        ota = str(ws.cell(row=r, column=7).value or '').strip()
        fish = str(ws.cell(row=r, column=8).value or '').strip()
        dob = str(ws.cell(row=r, column=13).value or '').strip()
        group = str(ws.cell(row=r, column=23).value or '').strip()
        if not ism and not shnum: continue
        full_fio = fish or f"{ism} {ota}".strip()
        students.append({
            'ism': ism,
            'ota': ota,
            'fio': full_fio,
            'dob': dob,
            'group': group,
            'shnum': shnum
        })

    groups = ["26-01", "26-02", "26-03", "26-04", "26-05", "26-06", "26-07"]
    generated_files = {}

    target_dirs = [
        os.path.join(BASE_DIR, 'pdf_jurnallar'),
        os.path.join(BASE_DIR, 'qayta_tekshiruv', 'pdf_jurnallar')
    ]

    for tdir in target_dirs:
        os.makedirs(tdir, exist_ok=True)
        # 1. Alohida guruh PDF lari
        for g in groups:
            out_file = os.path.join(tdir, f"Guruh_{g}.pdf")
            create_single_group_pdf(g, students, out_file)
            generated_files[g] = out_file

        # 2. Barcha 7 ta guruhni birlashtirgan YAGONA 1 ta A4 PDF (har bir guruh alohida varoqda)
        combined_file = os.path.join(tdir, "Barcha_Guruhlar_Jurnali.pdf")
        create_all_groups_combined_pdf(groups, students, combined_file)
        generated_files["ALL"] = combined_file

    print(f"[OK] Barcha {len(groups)} ta guruh uchun alohida va 1 ta yagona umumiy A4 PDF jurnallar yaratildi!")
    return generated_files

if __name__ == '__main__':
    build_all_group_pdfs()
