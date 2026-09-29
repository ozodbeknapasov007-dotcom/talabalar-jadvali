# -*- coding: utf-8 -*-
"""
QABUL - 2026.xlsx — administratorning "Qabul uchun shablon.xlsx" fayli bilan aynan bir xil qabul jadvali:
bitta "Лист1" sahifa, 1–2-qatorlar bo'sh, 3-qatorda sarlavha, 4-qatordan talabalar, 13 ustun (ruscha yo'q).

Qatorlar va rasmiy tarjimalar portal eksporti bilan AYNAN BIR XIL bo'lishi uchun
web/lib/qabul.ts dan olinadi (node web/scripts/qabul-rows.mjs orqali). Bu skript faqat
Excel faylni yozadi. Xizmatdagi "Baza administratori" fayli ham shu yerdagi build_workbook() dan.

Manba: talabalar_bazasi.json. Bazada telefon bo'lmasa — shartnoma .docx dagi raqam.
Rasmiy guruhlar tartibida, har guruh ichida alifbo bo'yicha. Safdan chiqarilganlar kirmaydi.
"""
import json
import os
import subprocess
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from shubhali_malumotlar_audit import GROUPS, load_students  # noqa: E402

sys.stdout.reconfigure(encoding='utf-8')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROWS_CLI = os.path.join(ROOT, 'web', 'scripts', 'qabul-rows.mjs')
QABUL_FILE = 'QABUL - 2026'  # web/lib/qabul.ts dagi QABUL_FILE bilan bir xil
OUT_SHABLON_2 = os.path.join(ROOT, 'hisobotlar', 'qabul', f'{QABUL_FILE}.xlsx')
OUT_DESKTOP = os.path.join(r'c:\Users\user\Desktop', f'{QABUL_FILE}.xlsx')
OUT_DIR_GROUPS = os.path.join(ROOT, 'hisobotlar', 'qabul', 'guruhlar')

# web/lib/qabul.ts dagi QABUL_HEADERS / QABUL_WIDTHS / QABUL_SHEET bilan bir xil
HEADERS = [
    '№', 'F.I.O', 'JSHSHIR', 'Pasport seriya raqami',
    'Tel raqam(pastdagi shablondagidek kiritilsin)', '2-Tel raqam(pastdagi shablondagidek kiritilsin)',
    'Yashash viloyat+tumani', 'UY manzili',
    'Avval o`qigan muassasa nomi(Uzbek tilida)', 'Avval o`qigan muassasa nomi(Ingliz tilida)',
    'Maktab, kollej va HK', 'Avval olgan diplom seriya+raqami', 'Boshlagan va tugatgan yili',
]
WIDTHS = [3.14, 34.43, 17.29, 13.43, 17.86, 22.57, 21.71, 17.71, 36.0, 26.29, 13.86, 16.43, 11.0]
SHEET = 'Лист1'
REVIEW_COL = {'viloyat': 7, 'makEn': 10}  # 1 dan boshlangan ustun


def qabul_rows(students: list[dict]) -> list[dict]:
    """web/lib/qabul.ts → qabulRow() natijalari (tartib saqlanadi)."""
    items = []
    for s in students:
        student = {k: v for k, v in s.items() if not k.startswith('_')}
        n = s['_n']
        if n.get('tel_manba') not in ('baza', ''):
            student['extraPhones'] = n['phones']
        items.append({'student': student, 'istisno': s.get('_ist') or None})
    res = subprocess.run(['node', ROWS_CLI], input=json.dumps(items, ensure_ascii=False).encode('utf-8'),
                         capture_output=True, check=True)
    return json.loads(res.stdout.decode('utf-8'))


def row_cells(r: dict, idx: int) -> list:
    num9 = lambda t: int(t) if len(t) == 9 and t.isdigit() else t  # noqa: E731
    return [idx, r['fio'], r['pinfl'], r['pv'], num9(r['tel1']), num9(r['tel2']), r['viloyat'], r['manzil'],
            r['makUz'], r['makEn'], r['eduType'], r['diplom'], r['yillar']]


def populate_sheet(ws, rows: list[dict]):
    """Admin shabloni: 3-qatorda sarlavha, 4-qatordan talabalar (rows tartibi saqlanadi)."""
    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font = Font(name='Times New Roman', size=12)
    fill_review = PatternFill('solid', start_color='FFF2CC')

    for c, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[3].height = 47.25
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(3, c, h)
        cell.font, cell.border = font, border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for idx, r in enumerate(rows, 1):
        rr = idx + 3
        ws.row_dimensions[rr].height = 15.75
        review_cols = {REVIEW_COL[k] for k in r['review'] if k in REVIEW_COL}
        for c, val in enumerate(row_cells(r, idx), 1):
            cell = ws.cell(rr, c, val if val != '' else None)
            cell.font, cell.border = font, border
            if c in review_cols:
                cell.fill = fill_review


def sort_rows(rows: list[dict]) -> list[dict]:
    """Rasmiy guruhlar tartibida, har guruh ichida F.I.O alifbosi bo'yicha."""
    order = {g: i for i, g in enumerate(GROUPS)}
    return sorted(rows, key=lambda x: (order.get(x['group'], len(order)), x['fio'].lower()))


def build_workbook(rows: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    populate_sheet(ws, sort_rows(rows))
    return wb


def official_rows() -> list[dict]:
    students = load_students()
    return qabul_rows([s for s in students if s.get('group') in GROUPS])


def main():
    rows = official_rows()
    groups = [g for g in GROUPS if any(r['group'] == g for r in rows)]

    wb = build_workbook(rows)
    wb.save(OUT_SHABLON_2)
    print(f"Saved: {OUT_SHABLON_2} ({len(rows)} nafar)")
    try:
        wb.save(OUT_DESKTOP)
        print(f"Saved: {OUT_DESKTOP}")
    except Exception as e:
        print(f"Could not save to Desktop directly (may be open in Excel): {e}")

    os.makedirs(OUT_DIR_GROUPS, exist_ok=True)
    for g in groups:
        g_rows = [r for r in rows if r['group'] == g]
        build_workbook(g_rows).save(os.path.join(OUT_DIR_GROUPS, f"{QABUL_FILE} ({g}).xlsx"))
        print(f"  {g}: {len(g_rows)} nafar")
    print(f"Saved {len(groups)} individual group templates in: {OUT_DIR_GROUPS}")
    review = sum(1 for r in rows if r['review'])
    print(f"Qo'lda tekshirish kerak (sariq kataklar): {review} ta qator")


if __name__ == '__main__':
    main()
