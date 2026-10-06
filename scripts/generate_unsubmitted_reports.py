# -*- coding: utf-8 -*-
"""
Generate comprehensive reports (Excel & Word) for 1-kurs students who haven't submitted the survey.
"""
import os, sys, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STUDENTS_FILE = os.path.join(BASE_DIR, 'data', 'students.json')
SOROV_FILE = os.path.join(BASE_DIR, 'data', 'sorovnoma_1kurs.json')
OUT_DIR = os.path.join(BASE_DIR, 'hisobotlar')
DOWNLOADS_DIR = r"C:\Users\user\Downloads"

OFFICIAL_GROUP_LEADERS = {
    '26-01': 'Mirzayeva.D',
    '26-02': 'Ochilov.D',
    '26-03': 'A.Asraliyev',
    '26-04': 'Xamdamova.M',
    '26-05': 'Rayimova.X',
    '26-06': 'Yuldashev.O',
    '26-07': 'Asraliyev.A',
}

with open(STUDENTS_FILE, encoding='utf-8') as f:
    students = json.load(f)

with open(SOROV_FILE, encoding='utf-8') as f:
    sorov = json.load(f)

kurs1 = [s for s in students if str(s.get('group', '')).startswith('26-')]

# Deduplicate sorov by student row
sorov_by_row = {}
for k, v in sorov.items():
    if 'row' in v:
        sorov_by_row[v['row']] = v

# Collect unsubmitted students
groups_list = ['26-01', '26-02', '26-03', '26-04', '26-05', '26-06', '26-07']

unsubmitted_all = []
stats = []

for g in groups_list:
    g_students = [s for s in kurs1 if s.get('group') == g]
    topshirgan = [s for s in g_students if sorov_by_row.get(s['row'], {}).get('topshirgan') == True]
    topshirmagan = [s for s in g_students if sorov_by_row.get(s['row'], {}).get('topshirgan') != True]
    
    tel_bor = [s for s in topshirmagan if (s.get('tel_shaxsiy') or s.get('tel'))]
    tel_yoq = [s for s in topshirmagan if not (s.get('tel_shaxsiy') or s.get('tel'))]
    
    stats.append({
        'group': g,
        'leader': OFFICIAL_GROUP_LEADERS.get(g, '—'),
        'total': len(g_students),
        'submitted': len(topshirgan),
        'unsubmitted': len(topshirmagan),
        'tel_bor': len(tel_bor),
        'tel_yoq': len(tel_yoq),
        'pct': round(len(topshirgan) / max(len(g_students), 1) * 100, 1)
    })
    
    for s in topshirmagan:
        tel = s.get('tel_shaxsiy') or s.get('tel') or ''
        status = "Telefoni bor (bog'lanish mumkin)" if tel else "Telefoni yo'q (aniqlash zarur)"
        unsubmitted_all.append({
            'row': s.get('row'),
            'tr': s.get('tr'),
            'group': g,
            'leader': OFFICIAL_GROUP_LEADERS.get(g, '—'),
            'fish': s.get('fish', ''),
            'pv': s.get('pv', ''),
            'pinfl': s.get('pinfl', ''),
            'tel': tel,
            'status': status,
            'has_tel': bool(tel)
        })

print(f"Total unsubmitted students: {len(unsubmitted_all)}")

# ─────────────────────────────────────────────────────────────────────────────
# 1. EXCEL REPORT
# ─────────────────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
wb.remove(wb.active) # remove default sheet

# Styles
f_title = Font(name='Calibri', size=14, bold=True, color='1E3A8A')
f_sub = Font(name='Calibri', size=10, italic=True, color='64748B')
f_hdr = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
f_bold = Font(name='Calibri', size=10, bold=True, color='0F172A')
f_regular = Font(name='Calibri', size=10, color='1E293B')

fill_navy = PatternFill('solid', fgColor='1E3A8A')
fill_indigo = PatternFill('solid', fgColor='312E81')
fill_total = PatternFill('solid', fgColor='E2E8F0')

fill_red = PatternFill('solid', fgColor='FEE2E2')
f_red = Font(name='Calibri', size=10, bold=True, color='991B1B')

fill_yellow = PatternFill('solid', fgColor='FEF3C7')
f_yellow = Font(name='Calibri', size=10, bold=True, color='92400E')

fill_alt = PatternFill('solid', fgColor='F8FAFC')
fill_white = PatternFill('solid', fgColor='FFFFFF')

border_thin = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                     top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))
border_med = Border(left=Side(style='thin', color='94A3B8'), right=Side(style='thin', color='94A3B8'),
                    top=Side(style='medium', color='1E3A8A'), bottom=Side(style='medium', color='1E3A8A'))

# Sheet 1: UMUMIY STATISTIKA (DASHBOARD)
ws_dash = wb.create_sheet(title="Umumiy tahlil")
ws_dash.views.sheetView[0].showGridLines = True

ws_dash.cell(1, 1, value="1-KURS TALABALAR SO'ROVNOMA TOPSHIRISH HOLATI (GURUHLAR KESIMIDA)").font = f_title
ws_dash.cell(2, 1, value="Ma'lumotlar avtomatik shakllantirilgan • Sana: 2026-yil 6-oktyabr").font = f_sub

dash_headers = [
    "№", "Guruh", "Guruh rahbari", "Jami talaba", "Topshirganlar",
    "Foiz (%)", "Topshirmaganlar", "Shundan telefoni borlar", "Shundan telefoni yo'qlar"
]

ws_dash.row_dimensions[4].height = 28
for c_idx, h in enumerate(dash_headers, 1):
    c = ws_dash.cell(4, c_idx, value=h)
    c.fill = fill_navy
    c.font = f_hdr
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = border_thin

for idx, st in enumerate(stats, 1):
    r_idx = idx + 4
    ws_dash.row_dimensions[r_idx].height = 22
    vals = [
        idx, st['group'], st['leader'], st['total'], st['submitted'],
        f"{st['pct']}%", st['unsubmitted'], st['tel_bor'], st['tel_yoq']
    ]
    fill_row = fill_alt if idx % 2 == 0 else fill_white
    for c_idx, val in enumerate(vals, 1):
        cell = ws_dash.cell(r_idx, c_idx, value=val)
        cell.fill = fill_row
        cell.border = border_thin
        cell.font = f_bold if c_idx in (2, 4, 7) else f_regular
        if c_idx == 7 and st['unsubmitted'] > 0:
            cell.fill = fill_yellow
            cell.font = f_yellow
        if c_idx == 9 and st['tel_yoq'] > 0:
            cell.fill = fill_red
            cell.font = f_red
        cell.alignment = Alignment(horizontal='center' if c_idx in (1, 2, 4, 5, 6, 7, 8, 9) else 'left', vertical='center')

# Totals row
tot_row = len(stats) + 5
ws_dash.row_dimensions[tot_row].height = 24
tot_vals = [
    "", "JAMI", "", sum(s['total'] for s in stats), sum(s['submitted'] for s in stats),
    f"{round(sum(s['submitted'] for s in stats) / sum(s['total'] for s in stats) * 100, 1)}%",
    sum(s['unsubmitted'] for s in stats), sum(s['tel_bor'] for s in stats), sum(s['tel_yoq'] for s in stats)
]
for c_idx, val in enumerate(tot_vals, 1):
    cell = ws_dash.cell(tot_row, c_idx, value=val)
    cell.fill = fill_total
    cell.border = border_med
    cell.font = Font(name='Calibri', size=11, bold=True, color='0F172A')
    cell.alignment = Alignment(horizontal='center', vertical='center')

for col in ws_dash.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_dash.column_dimensions[col_letter].width = max(max_len + 3, 11)

# Helper function for student table
def populate_unsubmitted_sheet(ws, student_list, sheet_title):
    ws.views.sheetView[0].showGridLines = True
    ws.cell(1, 1, value=sheet_title).font = f_title
    ws.cell(2, 1, value="Guruh rahbarlari tomonidan ma'lumotlarni yig'ish va to'ldirish uchun hisobot").font = f_sub
    
    hdrs = [
        "T/R", "Guruh", "Guruh rahbari", "Familiya Ism Sharif",
        "Pasport seriya №", "JSHSHIR (PINFL)", "Mavjud telefoni", "Holati",
        "Guruh rahbari qaydi / Yangi telefon / Manzil"
    ]
    ws.row_dimensions[4].height = 28
    for c_idx, h in enumerate(hdrs, 1):
        c = ws_dash.cell(4, c_idx)
        c_dest = ws.cell(4, c_idx, value=h)
        c_dest.fill = fill_indigo
        c_dest.font = f_hdr
        c_dest.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c_dest.border = border_thin
        
    for tr, s in enumerate(student_list, 1):
        r_idx = tr + 4
        ws.row_dimensions[r_idx].height = 24
        vals = [
            tr, s['group'], s['leader'], s['fish'],
            s['pv'], s['pinfl'],
            s['tel'] or "Mavjud emas",
            s['status'],
            "" # bo'sh katak rahbar uchun
        ]
        row_fill = fill_yellow if s['has_tel'] else fill_red
        stat_font = f_yellow if s['has_tel'] else f_red
        
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(r_idx, c_idx, value=val)
            cell.fill = fill_white if c_idx == 9 else row_fill
            cell.border = border_thin
            if c_idx in (1, 2, 5, 6, 7):
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.font = f_regular
            elif c_idx == 8:
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.font = stat_font
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.font = f_bold
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.font = f_regular
                
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    ws.column_dimensions['I'].width = 38 # Rahbar qaydi ustuni kengroq

# Sheet 2: BARCHA TOPSHIRMAGANLAR
ws_all = wb.create_sheet(title="Barcha topshirmaganlar")
populate_unsubmitted_sheet(ws_all, unsubmitted_all, f"SO'ROVNOMA TOPSHIRMAGAN BARCHA 1-KURS TALABALARI ({len(unsubmitted_all)} NAFAR)")

# Sheet 3..9: Har bir guruh alohida varaqda
for g in groups_list:
    g_unsub = [s for s in unsubmitted_all if s['group'] == g]
    if not g_unsub: continue
    ws_grp = wb.create_sheet(title=f"Guruh {g}")
    leader_name = OFFICIAL_GROUP_LEADERS.get(g, '')
    populate_unsubmitted_sheet(ws_grp, g_unsub, f"{g} GURUHI — SO'ROVNOMA TOPSHIRMAGANLAR ({len(g_unsub)} NAFAR) • Rahbar: {leader_name}")

excel_path = os.path.join(OUT_DIR, "1-kurs_Sorovnoma_Topshirmaganlar_Hisoboti.xlsx")
wb.save(excel_path)
print(f"✅ Excel hisobot saqlandi: {excel_path}")

# Nusxasini Downloads papkasiga ham qo'yamiz
try:
    import shutil
    dl_excel = os.path.join(DOWNLOADS_DIR, "1-kurs_Sorovnoma_Topshirmaganlar_Hisoboti.xlsx")
    shutil.copy2(excel_path, dl_excel)
    print(f"✅ Downloads papkasiga nusxalandi: {dl_excel}")
except Exception as e:
    print(f"Downloads nusxalash xatosi: {e}")

# ─────────────────────────────────────────────────────────────────────────────
# 2. WORD DOCUMENT (DOCX)
# ─────────────────────────────────────────────────────────────────────────────
doc = docx.Document()

# Set standard margins (0.5 inch / 1.27 cm)
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run("1-KURS TALABALARINING SO'ROVNOMA TOPSHIRISH NATIJALARI")
title_run.bold = True
title_run.font.size = Pt(15)
title_run.font.name = 'Times New Roman'
title_run.font.color.rgb = RGBColor(30, 58, 138)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_p.add_run("So'rovnoma to'ldirmagan talabalar va guruh rahbarlari bo'yicha to'liq hisobot\n(Sana: 06.10.2026)")
sub_run.font.size = Pt(10)
sub_run.font.italic = True
sub_run.font.name = 'Times New Roman'
sub_run.font.color.rgb = RGBColor(100, 116, 139)

doc.add_paragraph() # Spacing

# Summary Table
dash_table = doc.add_table(rows=1, cols=7)
dash_table.alignment = WD_TABLE_ALIGNMENT.CENTER
dash_hdr_cells = dash_table.rows[0].cells
dash_hdrs = ["№", "Guruh", "Guruh rahbari", "Jami", "Topshirgan", "Topshirmagan", "Foiz"]
for i, h in enumerate(dash_hdrs):
    dash_hdr_cells[i].text = h
    set_cell_background(dash_hdr_cells[i], "1E3A8A")
    set_cell_margins(dash_hdr_cells[i], 120, 120, 100, 100)
    p = dash_hdr_cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

for idx, st in enumerate(stats, 1):
    row_cells = dash_table.add_row().cells
    vals = [str(idx), st['group'], st['leader'], str(st['total']), str(st['submitted']), str(st['unsubmitted']), f"{st['pct']}%"]
    bg = "FEF3C7" if st['unsubmitted'] > 0 else "FFFFFF"
    for i, v in enumerate(vals):
        row_cells[i].text = v
        set_cell_background(row_cells[i], bg if i == 5 and st['unsubmitted'] > 0 else ("F8FAFC" if idx % 2 == 0 else "FFFFFF"))
        set_cell_margins(row_cells[i], 80, 80, 100, 100)
        p = row_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1, 3, 4, 5, 6) else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            r.font.bold = (i in (1, 3, 5))

# Total row
tot_cells = dash_table.add_row().cells
tot_vals = ["", "JAMI", "", str(sum(s['total'] for s in stats)), str(sum(s['submitted'] for s in stats)), str(sum(s['unsubmitted'] for s in stats)), f"{round(sum(s['submitted'] for s in stats) / sum(s['total'] for s in stats) * 100, 1)}%"]
for i, v in enumerate(tot_vals):
    tot_cells[i].text = v
    set_cell_background(tot_cells[i], "E2E8F0")
    set_cell_margins(tot_cells[i], 100, 100, 100, 100)
    p = tot_cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.bold = True
        r.font.size = Pt(10)

doc.add_page_break()

# Detailed group-by-group sections
for g in groups_list:
    g_unsub = [s for s in unsubmitted_all if s['group'] == g]
    if not g_unsub: continue
    
    leader_name = OFFICIAL_GROUP_LEADERS.get(g, '')
    gp = doc.add_paragraph()
    gp.paragraph_format.space_before = Pt(14)
    gp.paragraph_format.space_after = Pt(6)
    
    grun = gp.add_run(f"📌 {g} GURUHI — SO'ROVNOMA TOPSHIRMAGAN TALABALAR ({len(g_unsub)} NAFAR)")
    grun.font.name = 'Times New Roman'
    grun.font.size = Pt(12)
    grun.font.bold = True
    grun.font.color.rgb = RGBColor(49, 46, 129)
    
    lead_run = gp.add_run(f"  |  Guruh rahbari: {leader_name}")
    lead_run.font.name = 'Times New Roman'
    lead_run.font.size = Pt(11)
    lead_run.font.italic = True
    
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    thdrs = ["№", "Talaba Familiyasi, Ismi, Otasining ismi", "Pasport №", "Telefon raqami", "Holati", "Guruh rahbari qaydi / Yangi manzil"]
    hcells = table.rows[0].cells
    for i, h in enumerate(thdrs):
        hcells[i].text = h
        set_cell_background(hcells[i], "312E81")
        set_cell_margins(hcells[i], 100, 100, 80, 80)
        p = hcells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    for tr, s in enumerate(g_unsub, 1):
        rcells = table.add_row().cells
        c_vals = [
            str(tr),
            s['fish'],
            s['pv'],
            s['tel'] or "Mavjud emas",
            "🔴 Telefoni yo'q" if not s['has_tel'] else "🟡 Telefoni bor",
            "" # bo'sh katak
        ]
        row_bg = "FEE2E2" if not s['has_tel'] else "FEF3C7"
        for i, val in enumerate(c_vals):
            rcells[i].text = val
            set_cell_background(rcells[i], "FFFFFF" if i == 5 else row_bg)
            set_cell_margins(rcells[i], 80, 80, 80, 80)
            p = rcells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (1, 5) else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                r.font.bold = (i in (1, 4))
                if i == 4 and not s['has_tel']:
                    r.font.color.rgb = RGBColor(153, 27, 27)

docx_path = os.path.join(OUT_DIR, "1-kurs_Sorovnoma_Topshirmaganlar_Hisoboti.docx")
doc.save(docx_path)
print(f"✅ Word hisobot saqlandi: {docx_path}")

try:
    dl_docx = os.path.join(DOWNLOADS_DIR, "1-kurs_Sorovnoma_Topshirmaganlar_Hisoboti.docx")
    shutil.copy2(docx_path, dl_docx)
    print(f"✅ Downloads papkasiga nusxalandi: {dl_docx}")
except Exception as e:
    print(f"Downloads Word nusxalash xatosi: {e}")
