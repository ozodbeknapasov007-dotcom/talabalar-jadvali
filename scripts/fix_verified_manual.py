# -*- coding: utf-8 -*-
"""
QO'LDA TEKSHIRILGAN TUZATISHLAR
================================
Har bir qiymat rasmiy hujjat rasmidan o'qilib, matematik tasdiqlangan:
  • Pasport MRZ — o'z nazorat raqamlari bilan
  • PINFL       — 14-raqam [7,3,1] nazorat algoritmi bilan

    python scripts/fix_verified_manual.py --dry-run
    python scripts/fix_verified_manual.py
"""
import openpyxl, os, sys, re, shutil, datetime

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, 'Talabalar_Toliq_Royxati.xlsx')
DRY = '--dry-run' in sys.argv

W = [7, 3, 1, 7, 3, 1, 7, 3, 1, 7, 3, 1, 7]


def pinfl_ok(p):
    return (re.fullmatch(r'\d{14}', p or '') is not None
            and sum(int(p[i]) * W[i] for i in range(13)) % 10 == int(p[13]))


# ustun raqamlari
C = {'ifo': 2, 'shnum': 5, 'ota': 7, 'fish': 8, 'pass_fish': 9, 'pass': 10,
     'pinfl': 11, 'ber': 12, 'dob': 13, 'cert_fish': 14, 'cert': 15, 'qr': 16,
     'mak': 17, 'tur': 18, 'yil': 19, 'status': 21, 'izoh': 22, 'guruh': 23,
     'match': 24}

FIXES = [
    {
        'ifo': 'Xurramova Shodiya', 'guruh': '26-02',
        'manba': 'Pasport AC1967100 + ATTESTAT UM 0438281 rasmlari (Xurramova Shodiya.docx)',
        'set': {
            'pass_fish': 'KHURRAMOVA SHODIYA',
            'cert_fish': 'XURRAMOVA SHODIYA AROL QIZI',
            'cert': 'UM 0438281',
            'mak': "29-sonli umumiy o'rta ta'lim maktabi",
            'tur': 'Shahodatnoma',
            'yil': '2020',
            'izoh': "Fayl qo'lda tuzatildi: Xurramova Shodiya.docx. "
                    "Pasport va attestat rasmlari bilan to'liq tasdiqlandi. "
                    "DIQQAT: shartnomada yo'nalish 'Farmatsiya' deb yozilgan, bazada 'Hamshiralik ishi'.",
            'match': 'Pasport: MOS | Shahodatnoma: MOS',
        },
        'tekshir': {'pass': 'AC1967100', 'pinfl': '62704025590042',
                    'dob': '27.04.2002', 'ber': '12.06.2019'},
    },
    {
        'ifo': 'Suyunova Maftuna', 'guruh': '26-02',
        'manba': 'Pasport AB7166586 MRZ + Kasb-hunar kolleji diplomi (Suyunova Maftuna 122.docx)',
        'set': {
            'pinfl': '61409005680027',
            'pass_fish': 'SUYUNOVA MAFTUNA',
            'cert_fish': 'SUYUNOVA MAFTUNA AKBAR QIZI',
            'cert': 'K 5491697',
            'tur': 'Diplom',
            'mak': "Kitob tumani agrobiznes va xizmat ko'rsatish kasb-hunar kolleji",
            'yil': '2019',
            'qr': '',
            'izoh': "PINFL pasport MRZ dan tuzatildi (edi: 61409005680028). "
                    "Ta'lim hujjati almashtirildi: avval boshqa talabaning (Rasulova Husnobod) "
                    "shahodatnomasi UM 01891815 yozilgan edi; haqiqiysi — kasb-hunar kolleji diplomi.",
            'match': 'Pasport: MOS | Diplom: MOS',
        },
        'tekshir': {'pass': 'AB7166586', 'dob': '14.09.2000', 'ber': '09.07.2017'},
    },
    {
        'ifo': 'Berdimurodova Nafosat', 'guruh': '26-05',
        'manba': 'ID-karta AD9709558 QR (MRZ) — kuchaytirilgan skanerlashda o\'qildi',
        'set': {
            'pinfl': '60904075680023',
            'pass_fish': 'BERDIMURODOVA NAFOSAT',
            'izoh': 'PINFL ID-karta QR dan tuzatildi (edi: 60904075590028).',
        },
        'tekshir': {'pass': 'AD9709558', 'dob': '09.04.2007'},
    },

    # ── ID-karta orqa tomonidagi bosma MRZ dan ko'z bilan o'qilgan ──
    # (QR rasm sifati pastligi sababli o'qilmadi; har bir PINFL nazorat
    #  raqami algoritmi bilan tasdiqlandi)
    {
        'ifo': 'Ravshanova Ruxshona', 'guruh': '26-01',
        'manba': 'ID-karta AE6085379 orqa tomoni, bosma MRZ (Ravshanova Ruxshona 160.docx)',
        'set': {
            'pinfl': '60204095680082',
            'pass_fish': 'RAVSHANOVA RUHSHONA',
            'izoh': 'PINFL ID-karta MRZ dan tuzatildi (edi: 60204095590028).',
        },
        'tekshir': {'pass': 'AE6085379', 'dob': '02.04.2009', 'ber': '28.01.2026'},
    },
    {
        'ifo': 'Hasanova Marjona', 'guruh': '26-02',
        'manba': 'ID-karta AE5685784 orqa tomoni, bosma MRZ (Hasanova Marjona 192.docx)',
        'set': {
            'pinfl': '60411085680047',
            'pass_fish': 'HASANOVA MARJONA',
            'izoh': 'PINFL ID-karta MRZ dan tuzatildi (edi: 60411085590028).',
        },
        'tekshir': {'pass': 'AE5685784', 'dob': '04.11.2008', 'ber': '06.01.2026'},
    },
    {
        'ifo': 'Xusanova Aziza', 'guruh': '26-04',
        'manba': 'ID-karta AE5015467 orqa tomoni, bosma MRZ (Xusanova Aziza 196.docx)',
        'set': {
            'pinfl': '60110085590092',
            'pass_fish': 'XUSANOVA AZIZA',
            'izoh': 'PINFL ID-karta MRZ dan tuzatildi (edi: 60110085590028).',
        },
        'tekshir': {'pass': 'AE5015467', 'dob': '01.10.2008', 'ber': '12.11.2025'},
    },
    {
        'ifo': 'Jurayeva Madina', 'guruh': '26-05',
        'manba': 'ID-karta AD9394305 orqa tomoni, bosma MRZ (Jurayeva Madina  215.docx)',
        'set': {
            'pinfl': '62512075730060',
            'pass_fish': 'JURAYEVA MADINA',
            'izoh': 'PINFL ID-karta MRZ dan tuzatildi (edi: 62512075590018).',
        },
        'tekshir': {'pass': 'AD9394305', 'dob': '25.12.2007', 'ber': '13.11.2024'},
    },
    {
        'ifo': 'Nazirova Intizor', 'guruh': '26-07',
        'manba': 'Biometrik pasport AB5359926 MRZ + kasb-hunar kolleji diplomi '
                 "(Nazirova  Intizor 212.docx)",
        'set': {
            'pinfl': '40211995680023',
            'pass_fish': 'NARZIEVA INTIZOR',
            'cert_fish': 'NARZIYEVA INTIZOR ISKANDAR QIZI',
            'izoh': "PINFL pasport MRZ dan tuzatildi (edi: 40211995680028). "
                    "❗FAMILIYA: pasportda 'NARZIYEVA' deb yozilgan, ro'yxatda 'Nazirova'. "
                    "Ro'yxat O'ZGARTIRILMADI — qaysi biri to'g'ri ekanini tasdiqlang.",
            'match': "Pasport: FAMILIYA FARQ QILADI (Narziyeva / Nazirova)",
        },
        'tekshir': {'pass': 'AB5359926', 'dob': '02.11.1999', 'ber': '12.12.2016'},
    },

    # ── TIKLASH: dashboarddagi "Qayta tahlil" tugmasi buzib yozgan qator ──
    # 18.09.2026, 16:14–16:28 oralig'ida /api/reanalyze_student AI javobini
    # to'g'ridan-to'g'ri 7, 8, 9 va 11-ustunlarga yozib yuborgan.
    # Uchala mustaqil manba "Dalerovna" deydi:
    #   1) manba jadval "to'liq 1-KURS 2026-2027.xlsx"
    #   2) shartnoma matni (To'ychiyeva  Farzona 280.docx)
    #   3) pasportning o'zi (Gemini 2.5 Pro: patronymic = DALEROVNA)
    {
        'ifo': 'Tuychiyeva Farzona', 'guruh': '26-04',
        'manba': "TIKLASH — manba jadval + shartnoma matni + pasport (uchalasi 'Dalerovna')",
        'set': {
            'ota': 'Dalerovna',
            'fish': 'Tuychiyeva Farzona Dalerovna',
            'pinfl': '',          # "Rasmda ko'rsatilmagan" degan AI matni tozalanadi
            'pass_fish': '',      # tekshirilgan zanjir qayta to'ldiradi
            'izoh': "TIKLANDI: dashboard 'Qayta tahlil' tugmasi otasining ismini "
                    "'Daler qizi' ga o'zgartirib, PINFL ga AI matnini yozib yuborgan edi. "
                    "Manba jadval, shartnoma matni va pasport — uchalasi 'Dalerovna' deydi.",
        },
        'tekshir': {'ifo': 'Tuychiyeva Farzona'},
    },
]

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

index = {}
for r in range(2, ws.max_row + 1):
    ifo = str(ws.cell(row=r, column=C['ifo']).value or '').strip()
    g = str(ws.cell(row=r, column=C['guruh']).value or '').strip()
    if ifo:
        index[(ifo.lower(), g)] = r

changed = 0
for fx in FIXES:
    row = index.get((fx['ifo'].lower(), fx['guruh']))
    print('=' * 72)
    print(f"  {fx['ifo']}  [{fx['guruh']}]")
    print('=' * 72)
    if not row:
        print('   ❌ bazada topilmadi'); continue
    print(f"   manba: {fx['manba']}")

    # Mavjud qiymatlar kutilganiga mos kelishini tekshiramiz
    ok = True
    for k, want in fx.get('tekshir', {}).items():
        cur = str(ws.cell(row=row, column=C[k]).value or '').strip()
        if cur != want:
            print(f"   ⚠️  nazorat: {k} bazada {cur!r}, kutilgan {want!r}")
            ok = False
    if ok:
        print('   ✅ nazorat qiymatlari mos (pasport/sana o\'zgarmagan)')

    for k, v in fx['set'].items():
        if k not in C:
            print(f"   ⚠️  noma'lum ustun kaliti: {k}")
            continue
        cur = str(ws.cell(row=row, column=C[k]).value or '').strip()
        if cur == v:
            continue
        eski = cur if cur else '(bosh)'
        print(f"   {k:10s}: {eski:38s} -> {v}")
        if k == 'pinfl' and v and not pinfl_ok(v):
            print('      ❌ PINFL nazorat raqami mos emas — YOZILMAYDI')
            continue
        if not DRY:
            ws.cell(row=row, column=C[k], value=v)
            if k == 'pinfl':
                ws.cell(row=row, column=C[k]).number_format = '@'
            changed += 1

    # Statusni qayta hisoblash
    if not DRY:
        pv = str(ws.cell(row=row, column=C['pass']).value or '').strip()
        pin = str(ws.cell(row=row, column=C['pinfl']).value or '').strip()
        dob = str(ws.cell(row=row, column=C['dob']).value or '').strip()
        ce = str(ws.cell(row=row, column=C['cert']).value or '').strip()
        ws.cell(row=row, column=C['status'],
                value='TOPILDI' if (pv and pin and dob and ce)
                else ('CHALA' if (pv or pin or ce) else 'FAYL_YOQ'))

print()
if DRY:
    print('🔍 DRY-RUN — hech narsa yozilmadi')
    sys.exit(0)

stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
bp = os.path.join(BASE_DIR, 'backup', f'Talabalar_Toliq_Royxati_QOLDA_TUZATISHDAN_OLDIN_{stamp}.xlsx')
shutil.copy2(EXCEL_PATH, bp)
wb.save(EXCEL_PATH)
print(f"✅ Zaxira: backup/{os.path.basename(bp)}")
print(f"💾 Saqlandi — {changed} ta katak o'zgartirildi")
