# -*- coding: utf-8 -*-
import json
import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# 1. Update manual_file_map.json
mf_path = os.path.join(BASE_DIR, 'scripts', 'manual_file_map.json')
with open(mf_path, 'r', encoding='utf-8') as f:
    mf = json.load(f)

mf['fayllar']['Abdiraimova Shahzoda|25-16'] = {
    'file': 'Abdiraimova_Shahzoda_AD8711360.docx',
    'sabab': "25-16 guruhiga qo'shildi, ID-karta AD8711360"
}
mf['fayllar']['Hayitova Xolida|25-16'] = {
    'file': 'Hayitova_Xolida_AD5343254.docx',
    'sabab': "25-16 guruhiga qo'shildi, ID-karta AD5343254"
}
mf['fayllar']["Shog'dorova Marjona|25-16"] = {
    'file': 'Shogdorova_Marjona_AD9143185.docx',
    'sabab': "25-16 guruhiga qo'shildi, ID-karta AD9143185"
}
mf['fayllar']['Shogdorova Marjona|25-16'] = {
    'file': 'Shogdorova_Marjona_AD9143185.docx',
    'sabab': "25-16 guruhiga qo'shildi, ID-karta AD9143185"
}

with open(mf_path, 'w', encoding='utf-8') as f:
    json.dump(mf, f, ensure_ascii=False, indent=2)
print("manual_file_map.json yangilandi")

# 2. Add to Excel
ex_path = os.path.join(BASE_DIR, 'data', 'Talabalar_Toliq_Royxati.xlsx')
wb = openpyxl.load_workbook(ex_path)
ws = wb.worksheets[0]

new_students = [
    {
        'ism': 'Abdiraimova Shahzoda',
        'ota': 'Dilshod qizi',
        'fish': 'Abdiraimova Shahzoda Dilshod qizi',
        'pass_fish': 'ABDIRAIMOVA SHAHZODA DILSHOD QIZI',
        'pv': 'AD8711360',
        'pinfl': '61511075720098',
        'ber': '20.09.2024',
        'dob': '15.11.2007',
        'mak': 'Chiroqchi tumani',
        'group': '25-16',
        'yon': 'Hamshiralik ishi',
    },
    {
        'ism': 'Hayitova Xolida',
        'ota': 'Asad qizi',
        'fish': 'Hayitova Xolida Asad qizi',
        'pass_fish': 'HAYITOVA XOLIDA ASAD QIZI',
        'pv': 'AD5343254',
        'pinfl': '61112065720083',
        'ber': '06.12.2023',
        'dob': '11.12.2006',
        'mak': 'Chiroqchi tumani',
        'group': '25-16',
        'yon': 'Hamshiralik ishi',
    },
    {
        'ism': "Shog'dorova Marjona",
        'ota': 'Eshtemir qizi',
        'fish': "Shog'dorova Marjona Eshtemir qizi",
        'pass_fish': "SHOG'DOROVA MARJONA ESHTEMIR QIZI",
        'pv': 'AD9143185',
        'pinfl': '63006075720080',
        'ber': '24.10.2024',
        'dob': '30.06.2007',
        'mak': 'Chiroqchi tumani',
        'group': '25-16',
        'yon': 'Hamshiralik ishi',
    }
]

# 3. Update verifications.json
vf_path = os.path.join(BASE_DIR, 'scripts', 'verifications.json')
with open(vf_path, 'r', encoding='utf-8') as f:
    vmap = json.load(f)

for s in new_students:
    nr = ws.max_row + 1
    ws.cell(row=nr, column=1, value=nr - 1)
    ws.cell(row=nr, column=2, value=s['ism'])
    ws.cell(row=nr, column=3, value=s['yon'])
    ws.cell(row=nr, column=4, value=None)
    ws.cell(row=nr, column=5, value=None)
    ws.cell(row=nr, column=6, value=None)
    ws.cell(row=nr, column=7, value=s['ota'])
    ws.cell(row=nr, column=8, value=s['fish'])
    ws.cell(row=nr, column=9, value=s['pass_fish'])
    ws.cell(row=nr, column=10, value=s['pv'])
    ws.cell(row=nr, column=11, value=s['pinfl'])
    ws.cell(row=nr, column=12, value=s['ber'])
    ws.cell(row=nr, column=13, value=s['dob'])
    ws.cell(row=nr, column=14, value=None)
    ws.cell(row=nr, column=15, value=None)
    ws.cell(row=nr, column=16, value=None)
    ws.cell(row=nr, column=17, value=s['mak'])
    ws.cell(row=nr, column=18, value="Umumiy o'rta maktab")
    ws.cell(row=nr, column=19, value=None)
    ws.cell(row=nr, column=20, value=None)
    ws.cell(row=nr, column=21, value='TOPILDI')
    ws.cell(row=nr, column=22, value='2025-2026 kontingenti: 2-kurs')
    ws.cell(row=nr, column=23, value=s['group'])
    ws.cell(row=nr, column=24, value="MOS (TO'LIQ)")
    ws.cell(row=nr, column=25, value='TASDIQLANDI')
    ws.cell(row=nr, column=28, value='KIRITILDI')

    vmap[str(nr)] = 'TASDIQLANDI'
    vmap[f"pinfl_{s['pinfl']}"] = 'TASDIQLANDI'
    vmap[f"baza_{nr}"] = 'KIRITILDI'
    vmap[f"baza_pinfl_{s['pinfl']}"] = 'KIRITILDI'
    print(f"Excelga qo'shildi: {s['fish']} (qator {nr})")

wb.save(ex_path)
with open(vf_path, 'w', encoding='utf-8') as f:
    json.dump(vmap, f, ensure_ascii=False, indent=2)
print("Talabalar Excel va verifications.json ga muvaffaqiyatli saqlandi!")
