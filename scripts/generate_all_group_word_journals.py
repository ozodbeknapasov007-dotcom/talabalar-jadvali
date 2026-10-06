# -*- coding: utf-8 -*-
"""
BARCHA GURUHLAR UCHUN RASMIY WORD DAVOMAT JURNALLARINI YARATISH SKRIPTI
========================================================================
Texnikumdagi barcha 19 ta akademik guruh (1-kurs, 2-kurs, 3-kurs) uchun
"Davomat jurnali 26-02.docx" shablonidan foydalanib, har bir guruhning
talabalari ma'lumotlari (FISH, telefon, buyruq, manzil, DOB, ota-ona, qatnov)
bilan to'liq to'ldirilgan Word (.docx) jurnallarini yaratadi.

Shuningdek, barcha jurnallarni bitta .zip arxivga jamlaydi.
"""

import os
import sys
import json
import zipfile
import shutil
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(BASE_DIR, 'hisobotlar', 'davomat_jurnallari')
os.makedirs(OUT_DIR, exist_ok=True)

TEMPLATE_PATH = os.path.join(BASE_DIR, 'Davomat jurnali 26-02.docx')
if not os.path.exists(TEMPLATE_PATH):
    print(f"XATOLIK: {TEMPLATE_PATH} shablon fayli topilmadi!")
    sys.exit(1)

# 1. Talabalar va guruhlar sozlamalarini yuklash
with open(os.path.join(BASE_DIR, 'data', 'students.json'), 'r', encoding='utf-8') as f:
    students = json.load(f)

with open(os.path.join(BASE_DIR, 'data', 'guruhlar.json'), 'r', encoding='utf-8') as f:
    group_settings = json.load(f)

# Barcha rasmiy guruhlarni aniqlash (1-kurs, 2-kurs, 3-kurs)
OFFICIAL_GROUPS = [
    '26-01', '26-02', '26-03', '26-04', '26-05', '26-06', # 1-kurs
    '25-16', '25-17', '25-18', '25-19', '25-20', '25-21', '25-22', '25-23', # 2-kurs
    '24-11', '24-12', '24-13', '24-15', '24-16' # 3-kurs
]

# 1-kurs buyruqlari
COURSE_1_ORDERS = {
    '26-01': ('T-40/1', '02.09.2026'),
    '26-02': ('T-40/2', '02.09.2026'),
    '26-03': ('T-40/2', '02.09.2026'),
    '26-04': ('T-40/2', '02.09.2026'),
    '26-05': ('T-40/2', '02.09.2026'),
    '26-06': ('T-40/2', '02.09.2026'),
}

def set_cell_text(cell, text, bold=False, italic=False, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(str(text or ''))
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

def clear_cell(cell):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)

def ensure_table_rows(table, required_rows):
    """Jadvalda kerakli qatorlar soni yetarli bo'lishini ta'minlash"""
    while len(table.rows) < required_rows:
        table.add_row()

def fill_group_doc(group_code):
    cfg = group_settings.get(group_code, {})
    leader = cfg.get('rahbar', '—')
    kurs = cfg.get('kurs', 1)
    yonalish = cfg.get('yonalish', 'Hamshiralik ishi')

    # Shu guruh talabalari
    group_students = [s for s in students if (s.get('group') or '').strip() == group_code]
    group_students_sorted = sorted(group_students, key=lambda x: str(x.get('fish', '')))
    total_st = len(group_students_sorted)

    print(f"-> Guruh {group_code}: {total_st} nafar talaba, Rahbar: {leader} ({kurs}-kurs)")

    doc = docx.Document(TEMPLATE_PATH)

    # 1. Paragraflarni moslashtirish (sarlavhalar va guruh rahbari)
    if len(doc.paragraphs) > 1:
        doc.paragraphs[1].text = f"{group_code}-guruh talabalari ro'yxati"
    if len(doc.paragraphs) > 16:
        doc.paragraphs[16].text = "bilan tanishdim:"
    if len(doc.paragraphs) > 17:
        doc.paragraphs[17].text = f"{group_code} - guruh rahbari:   {leader}"

    # 2. TABLE 0: Ro'yxat, Telefon, Qabul buyrug'i
    # Col 0: T/r, Col 1: F.I.Sh, Col 2: Telefon, Col 3: Buyruq va sana
    t0 = doc.tables[0]
    ensure_table_rows(t0, total_st + 1)
    for idx, s in enumerate(group_students_sorted, 1):
        row = t0.rows[idx]
        set_cell_text(row.cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], s.get('fish', ''), bold=True)
        set_cell_text(row.cells[2], s.get('tel_shaxsiy') or s.get('tel', ''), align=WD_ALIGN_PARAGRAPH.CENTER)
        
        # Buyruq raqami
        b_num = s.get('buyruq')
        b_sana = s.get('buyruq_sana')
        if not b_num and group_code in COURSE_1_ORDERS:
            b_num, b_sana = COURSE_1_ORDERS[group_code]
        
        b_txt = f"{b_num}, {b_sana}" if b_num and b_sana else (b_num or "")
        set_cell_text(row.cells[3], b_txt, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Ortiqcha qatorlarni tozalash
    for r_idx in range(total_st + 1, len(t0.rows)):
        for cell in t0.rows[r_idx].cells:
            clear_cell(cell)

    # 3. TABLE 1: Doimiy manzil va Tug'ilgan sana
    # Col 0: T/r, Col 1: F.I.Sh, Col 2: Doimiy manzil, Col 3: Tug'ilgan sana
    t1 = doc.tables[1]
    ensure_table_rows(t1, total_st + 1)
    for idx, s in enumerate(group_students_sorted, 1):
        row = t1.rows[idx]
        set_cell_text(row.cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], s.get('fish', ''), bold=True)
        manzil_txt = s.get('manzil_toliq') or s.get('manzil_tuman') or ""
        set_cell_text(row.cells[2], manzil_txt)
        set_cell_text(row.cells[3], s.get('dob', ''), align=WD_ALIGN_PARAGRAPH.CENTER)
    
    for r_idx in range(total_st + 1, len(t1.rows)):
        for cell in t1.rows[r_idx].cells:
            clear_cell(cell)

    # 4. TABLE 2: Yaqin qarindoshlar ma'lumotlari
    # Row 0: Title, Row 1: Header. Students start at row 2!
    t2 = doc.tables[2]
    ensure_table_rows(t2, total_st + 2)
    for idx, s in enumerate(group_students_sorted, 1):
        row = t2.rows[idx + 1]
        set_cell_text(row.cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        if s.get('tel_otaona'):
            kim = s.get('tel_otaona_kim') or "Otasi / Onasi"
            set_cell_text(row.cells[1], kim)
            set_cell_text(row.cells[2], s.get('tel_otaona'), align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            clear_cell(row.cells[1])
            clear_cell(row.cells[2])
        clear_cell(row.cells[3])
        clear_cell(row.cells[4])
    
    for r_idx in range(total_st + 2, len(t2.rows)):
        for cell in t2.rows[r_idx].cells:
            clear_cell(cell)

    # 5. TABLE 3: O'qish davridagi yashash manzili va ishlaydigan telefoni
    t3 = doc.tables[3]
    ensure_table_rows(t3, total_st + 1)
    for idx, s in enumerate(group_students_sorted, 1):
        row = t3.rows[idx]
        set_cell_text(row.cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        qat = s.get('qatnov') or ""
        manz = s.get('manzil_toliq') or ""
        if qat and manz:
            full_qat = f"{qat}: {manz}"
        elif qat:
            full_qat = qat
        elif manz:
            full_qat = manz
        else:
            full_qat = ""
        set_cell_text(row.cells[1], full_qat)
        set_cell_text(row.cells[2], s.get('tel_shaxsiy') or s.get('tel', ''), align=WD_ALIGN_PARAGRAPH.CENTER)
        clear_cell(row.cells[3])

    for r_idx in range(total_st + 1, len(t3.rows)):
        for cell in t3.rows[r_idx].cells:
            clear_cell(cell)

    # 6. TABLE 4: Ichki tartib qoidalari bilan tanishtirilganligi
    t4 = doc.tables[4]
    ensure_table_rows(t4, total_st + 1)
    for idx, s in enumerate(group_students_sorted, 1):
        row = t4.rows[idx]
        set_cell_text(row.cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], s.get('fish', ''), bold=True)
        set_cell_text(row.cells[2], "Tanishtirildi", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[3], "02.09.2026", align=WD_ALIGN_PARAGRAPH.CENTER)
        clear_cell(row.cells[4])

    for r_idx in range(total_st + 1, len(t4.rows)):
        for cell in t4.rows[r_idx].cells:
            clear_cell(cell)

    # 7. TABLES 5 to 22: Kunlik davomat jadvallari (18 ta jadval)
    for t_idx in range(5, min(23, len(doc.tables))):
        tbl = doc.tables[t_idx]
        # Student rows start at index 2 (row 0 and 1 are headers)
        ensure_table_rows(tbl, total_st + 2)
        for idx, s in enumerate(group_students_sorted, 1):
            r_idx = idx + 1
            row = tbl.rows[r_idx]
            set_cell_text(row.cells[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5)
            set_cell_text(row.cells[1], s.get('fish', ''), bold=False, size=8.5)

        for r_idx in range(total_st + 2, len(tbl.rows)):
            clear_cell(tbl.rows[r_idx].cells[0])
            clear_cell(tbl.rows[r_idx].cells[1])

    # Fayllarni saqlash
    target_filename = f"Davomat_jurnali_{group_code}.docx"
    out_file = os.path.join(OUT_DIR, target_filename)
    doc.save(out_file)

    # Agar 26-02 bo'lsa, mos ravishda asosiy yo'llarga ham saqlanadi
    if group_code == '26-02':
        doc.save(os.path.join(BASE_DIR, 'Davomat jurnali 26-02 (TO\'LDIRILGAN).docx'))
        doc.save(os.path.join(BASE_DIR, 'hisobotlar', 'Davomat jurnali 26-02 (TO\'LDIRILGAN).docx'))

    return out_file

def main():
    print("=" * 65)
    print("BARCHA GURUHLARNING WORD DAVOMAT JURNALLARINI YARATISH BOSHLANDI")
    print("=" * 65)

    created_files = []
    course_1_files = []

    for group_code in OFFICIAL_GROUPS:
        try:
            fpath = fill_group_doc(group_code)
            created_files.append((group_code, fpath))
            if group_code.startswith('26-'):
                course_1_files.append((group_code, fpath))
        except Exception as e:
            print(f"XATOLIK ({group_code}): {e}")

    print(f"\nJami {len(created_files)} ta guruh jurnallari Word (.docx) formatida yaratildi.")

    # 1. Barcha guruhlar jurnallari ZIP arxivi
    all_zip_path = os.path.join(BASE_DIR, 'hisobotlar', 'Barcha_Guruhlar_Davomat_Jurnallari_Word.zip')
    with zipfile.ZipFile(all_zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for grp, fpath in created_files:
            zipf.write(fpath, arcname=f"Davomat_jurnali_{grp}.docx")
    
    # nusxasini davomat_jurnallari ichiga ham joylaymiz
    shutil.copy2(all_zip_path, os.path.join(OUT_DIR, 'Barcha_Guruhlar_Davomat_Jurnallari_Word.zip'))

    # 2. Faqat 1-kurs guruhlari ZIP arxivi
    c1_zip_path = os.path.join(BASE_DIR, 'hisobotlar', '1-kurs_Guruhlar_Davomat_Jurnallari_Word.zip')
    with zipfile.ZipFile(c1_zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for grp, fpath in course_1_files:
            zipf.write(fpath, arcname=f"Davomat_jurnali_{grp}.docx")
    shutil.copy2(c1_zip_path, os.path.join(OUT_DIR, '1-kurs_Guruhlar_Davomat_Jurnallari_Word.zip'))

    print(f"\n✅ ZIP arxiv yaratildi: {all_zip_path} ({len(created_files)} ta fayl)")
    print(f"✅ 1-kurs ZIP arxiv yaratildi: {c1_zip_path} ({len(course_1_files)} ta fayl)")
    print("=" * 65)

if __name__ == '__main__':
    main()
