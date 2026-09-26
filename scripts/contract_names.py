# -*- coding: utf-8 -*-
"""
SHARTNOMA MATNIDAN ISM AJRATISH
================================
Shartnoma .docx fayllarida talabaning ismi o'zbek lotinida MATN ko'rinishida
yozilgan bo'ladi (rasm emas):

    YO'NALISHI:        | Hamshiralik ishi
    FAMILIYASI ISMI:   | Suyunova Maftuna
    OTASINI ISMI:      | Akbar qizi
    ...

Bu AI'dan MUSTAQIL, determinik manba — 151 ta fayldan 145 tasida mavjud.
AI o'qigan pasport ismini tekshirish uchun ishlatiladi.

DIQQAT: bu etalon EMAS — kollej xodimi terib kiritgan, xato bo'lishi mumkin
(masalan bazada "Turg'unova Jasmina", matnda "Turayeva Jasmina").
Shuning uchun u TASDIQLOVCHI manba, hal qiluvchi emas.
"""
import os
import re

try:
    import docx
except ImportError:
    docx = None

# "FAMILIYASI ISMI:" ko'rinishidagi yorliqlar (qiymat keyingi katakda turadi)
LBL_FAM = re.compile(r"(?i)^\s*FAMILIYA(?:SI)?\s*[,/]?\s*ISM(?:I)?\s*:?\s*$")
LBL_OTA = re.compile(r"(?i)^\s*OTASINI(?:NG)?\s*ISM(?:I)?\s*:?\s*$")

# "FAMILIYASI ISMI: Suyunova Maftuna" — yorliq va qiymat bitta katakda
INL_FAM = re.compile(r"(?i)FAMILIYA(?:SI)?\s*[,/]?\s*ISM(?:I)?\s*:\s*(.+)")
INL_OTA = re.compile(r"(?i)OTASINI(?:NG)?\s*ISM(?:I)?\s*:\s*(.+)")

# Yorliq bo'lib qolgan, qiymat bo'lmagan matnlar
LOOKS_LIKE_LABEL = re.compile(
    r"(?i)(yo.?nalishi|tugatgan|shartnoma|telefon|raqami|sanasi|ismi|familiya)")


def _cells(path):
    """Docx dagi barcha matn kataklarini tartib bilan qaytaradi"""
    if docx is None:
        return []
    try:
        d = docx.Document(path)
    except Exception:
        return []
    out = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                out.append(c.text)
    return [re.sub(r'\s+', ' ', t.replace('\xa0', ' ')).strip()
            for t in out if t and t.strip()]


def _clean_value(v):
    v = re.sub(r'\s+', ' ', str(v or '').replace('\xa0', ' ')).strip(' :|')
    # ba'zi kataklarda qiymatdan keyin keyingi yorliq ham qo'shilib ketadi
    v = re.split(r"(?i)\b(YO.?NALISHI|TUGATGAN|SHARTNOMA|TELEFON)\b", v)[0]
    return v.strip(' :|')


def read_contract_name(path):
    """
    Shartnoma faylidan (familiya_ism, otasining_ismi) ni qaytaradi.
    Topilmasa ('', '') qaytaradi.
    """
    cells = _cells(path)
    if not cells:
        return '', ''

    fam = ota = ''

    # 1-usul: yorliq katagi + keyingi katakdagi qiymat
    for i, c in enumerate(cells):
        if not fam and LBL_FAM.match(c):
            for j in range(i + 1, min(i + 3, len(cells))):
                cand = _clean_value(cells[j])
                if cand and not LOOKS_LIKE_LABEL.match(cand):
                    fam = cand
                    break
        if not ota and LBL_OTA.match(c):
            for j in range(i + 1, min(i + 3, len(cells))):
                cand = _clean_value(cells[j])
                if cand and not LOOKS_LIKE_LABEL.match(cand):
                    ota = cand
                    break

    # 2-usul: yorliq va qiymat bitta katakda
    if not fam or not ota:
        joined = ' | '.join(cells)
        if not fam:
            m = INL_FAM.search(joined)
            if m:
                fam = _clean_value(m.group(1).split('|')[0])
        if not ota:
            m = INL_OTA.search(joined)
            if m:
                ota = _clean_value(m.group(1).split('|')[0])

    return fam, ota


def read_contract_fields(path):
    """
    Kengaytirilgan variant — shartnomadagi boshqa maydonlarni ham qaytaradi.
    Tekshiruvda yo'nalish/maktab nomuvofiqligini aniqlash uchun foydali.
    """
    cells = _cells(path)
    out = {'fam_ism': '', 'ota': '', 'yonalish': '', 'maktab': '',
           'bitirgan': '', 'shnum': '', 'shsana': '', 'tel': ''}
    if not cells:
        return out

    out['fam_ism'], out['ota'] = read_contract_name(path)

    LBLS = [
        ('yonalish', r"(?i)^\s*YO.?NALISHI\s*:?\s*$"),
        ('maktab',   r"(?i)^\s*TUGATGAN\s*O.?QISH\s*JOYI\s*:?\s*$"),
        ('bitirgan', r"(?i)^\s*TUGATGAN\s*YILI\s*:?\s*$"),
        ('shnum',    r"(?i)^\s*SHARTNOMA\s*RAQAMI\s*:?\s*$"),
        ('shsana',   r"(?i)^\s*SHARTNOMA\s*SANASI\s*:?\s*$"),
        ('tel',      r"(?i)^\s*TELEFON\s*RAQAMI\s*:?\s*$"),
    ]
    for key, pat in LBLS:
        rx = re.compile(pat)
        for i, c in enumerate(cells):
            if rx.match(c) and i + 1 < len(cells):
                out[key] = _clean_value(cells[i + 1])
                break
    return out


if __name__ == '__main__':
    import sys
    sys.stdout.reconfigure(encoding='utf-8')
    BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    args = sys.argv[1:]
    if not args:
        print("Ishlatish: python scripts/contract_names.py \"<fayl nomi>.docx\" [...]")
        sys.exit(0)
    for fn in args:
        p = fn if os.path.isabs(fn) else os.path.join(BASE, 'hujjatlar', 'shartnomalar', fn)
        d = read_contract_fields(p)
        print('=' * 66)
        print(os.path.basename(p))
        for k, v in d.items():
            if v:
                print(f'   {k:10s}: {v}')
