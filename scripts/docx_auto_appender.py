# -*- coding: utf-8 -*-
"""
YANGI DOCX FAYLLARNI EXCELGA AVTOMATIK QO'SHUVCHI MODUL
======================================================
"""
import openpyxl, docx, os, sys, re, io
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')

def parse_docx_and_append(docx_path, excel_path='Talabalar_Toliq_Royxati.xlsx'):
    if not os.path.exists(docx_path):
        return None
    
    fname = os.path.basename(docx_path)
    doc = docx.Document(docx_path)
    
    # Jadvaldan ma'lumotlarni o'qish
    data = {
        'ism': '',
        'yonalis': 'Hamshiralik ishi',
        'tolov': "To'lov qildi",
        'shnum': '',
        'shsan': '',
        'ota': '',
        'fish': '',
        'maktab': '',
        'edutur': "Umumiy o'rta maktab",
        'yil': '',
        'tel': '',
        'docx': fname,
        'docx_path': os.path.abspath(docx_path)
    }
    
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if len(cells) >= 2:
                k = cells[0].lower()
                v = cells[1]
                if 'familiya' in k or 'ism' in k: data['ism'] = v
                elif 'ota' in k: data['ota'] = v
                elif 'yo’nalish' in k or "yo'nalish" in k: data['yonalis'] = v
                elif 'o’qish joyi' in k or "o'qish joyi" in k: data['maktab'] = v
                elif 'tugatgan yili' in k: data['yil'] = v
                elif 'shartnoma raqami' in k: data['shnum'] = v
                elif 'shartnoma sanasi' in k: data['shsan'] = v
                elif 'telefon' in k: data['tel'] = v

    if not data['ism']:
        # Fayl nomidan ism olish
        m_name = re.match(r'^([A-Za-z\'\s]+)\s*\d*', fname.replace('.docx', ''))
        if m_name:
            data['ism'] = m_name.group(1).strip()

    if data['ota'] and data['ism']:
        data['fish'] = f"{data['ism']} {data['ota']}"
    else:
        data['fish'] = data['ism']

    # Excelga qo'shish
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    # Takrorlanishni tekshirish (agar shu shartnoma raqami yoki ism allaqachon bo'lsa)
    exists = False
    for r in range(2, ws.max_row+1):
        cur_shnum = str(ws.cell(row=r, column=5).value or '').strip()
        cur_ism = str(ws.cell(row=r, column=2).value or '').strip()
        if (data['shnum'] and cur_shnum == str(data['shnum'])) or (cur_ism and cur_ism.lower() == data['ism'].lower()):
            exists = True
            break
    
    if not exists:
        next_row = ws.max_row + 1
        ws.cell(row=next_row, column=1, value=next_row - 1)
        ws.cell(row=next_row, column=2, value=data['ism'])
        ws.cell(row=next_row, column=3, value=data['yonalis'])
        ws.cell(row=next_row, column=4, value=data['tolov'])
        ws.cell(row=next_row, column=5, value=data['shnum'])
        ws.cell(row=next_row, column=6, value=data['shsan'])
        ws.cell(row=next_row, column=7, value=data['ota'])
        ws.cell(row=next_row, column=8, value=data['fish'])
        ws.cell(row=next_row, column=18, value=data['maktab'])
        ws.cell(row=next_row, column=19, value=data['edutur'])
        ws.cell(row=next_row, column=20, value=data['yil'])
        ws.cell(row=next_row, column=21, value=data['tel'])
        ws.cell(row=next_row, column=22, value=f'=HYPERLINK("{data["docx_path"]}", "📄 {data["docx"]}")')
        ws.cell(row=next_row, column=23, value="YANGI")
        ws.cell(row=next_row, column=24, value="Telegramdan yangi yuklandi")
        wb.save(excel_path)
        print(f"[QO'SHILDI] {data['ism']} (Sh. #{data['shnum']}) Excelga qo'shildi!")
        return data
    else:
        print(f"[MAVJUD] {data['ism']} allaqachon bazada bor.")
        return None
