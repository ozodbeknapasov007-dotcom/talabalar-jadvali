# -*- coding: utf-8 -*-
"""
HUJJAT KAMCHILIKLARI AUDITI
============================
Talabalar_Toliq_Royxati.xlsx dagi har bir talabani ustunma-ustun tekshiradi
va qaysi hujjatida nima yetishmasligini toifalarga ajratadi.

Chiqish: ekranga hisobot + HUJJAT_KAMCHILIKLARI.csv
"""
import csv, io, os, re, sys, collections
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, 'data', 'Talabalar_Toliq_Royxati.xlsx')
OUT = os.path.join(BASE, 'hisobotlar', 'tekshiruvlar', 'HUJJAT_KAMCHILIKLARI.csv')

C = dict(ifo=2, shartnoma=5, otasi=7, pasfio=9, pasraqam=10, pinfl=11,
         berilgan=12, tugilgan=13, shfio=14, shseriya=15, qr=16,
         maktab=17, turi=18, bitirgan=19, tel=20, status=21, izoh=22,
         guruh=23, moslik=24, tasdiq=25)

W = [7, 3, 1, 7, 3, 1, 7, 3, 1, 7, 3, 1, 7]


def bosh(ws, r, k):
    v = ws.cell(r, C[k]).value
    s = '' if v is None else str(v).strip()
    return s if s and s.lower() != 'none' else ''


def main():
    ws = openpyxl.load_workbook(XLSX).active
    talabalar = []

    for r in range(2, ws.max_row + 1):
        ism = bosh(ws, r, 'ifo')
        if not ism:
            continue
        k = []                                   # kamchiliklar
        og = []                                  # ogohlantirishlar

        izoh = bosh(ws, r, 'izoh')
        status = bosh(ws, r, 'status')

        # ── shaxsni tasdiqlovchi hujjat ──
        pin, pas = bosh(ws, r, 'pinfl'), bosh(ws, r, 'pasraqam')
        if not pas:
            k.append('Pasport/ID raqami yo‘q')
        if not pin or not re.fullmatch(r'\d{14}', pin):
            k.append('JSHSHIR yo‘q' if not pin else 'JSHSHIR formati noto‘g‘ri')
        elif sum(int(pin[i]) * W[i] for i in range(13)) % 10 != int(pin[13]):
            k.append('JSHSHIR nazorat raqami XATO')
        if not bosh(ws, r, 'tugilgan'):
            k.append('Tug‘ilgan sana yo‘q')
        if not bosh(ws, r, 'berilgan'):
            og.append('pasport berilgan sanasi yo‘q')

        # PINFL ichidagi sana pasportdagi sana bilan mosmi
        tug = bosh(ws, r, 'tugilgan')
        if re.fullmatch(r'\d{14}', pin or '') and re.fullmatch(r'\d{2}\.\d{2}\.\d{4}', tug):
            asr = {'1': 1800, '2': 1800, '3': 1900, '4': 1900, '5': 2000, '6': 2000}.get(pin[0])
            if asr:
                p_sana = f'{pin[1:3]}.{pin[3:5]}.{asr + int(pin[5:7])}'
                if p_sana != tug:
                    k.append(f'JSHSHIR sanasi ({p_sana}) tug‘ilgan sanaga ({tug}) mos emas')

        # ── ta'lim hujjati ──
        ser, turi = bosh(ws, r, 'shseriya'), bosh(ws, r, 'turi')
        mak, bit = bosh(ws, r, 'maktab'), bosh(ws, r, 'bitirgan')
        if not ser:
            k.append('Shahodatnoma/diplom seriyasi yo‘q')
        if not turi:
            og.append('ta’lim hujjati turi ko‘rsatilmagan')
        if not mak:
            k.append('Tugatgan o‘qish joyi yo‘q')
        if not bit:
            k.append('Bitirgan yili yo‘q')
        if ser and not bosh(ws, r, 'qr'):
            og.append('e-shahodatnoma QR havolasi yo‘q')

        # ── shartnoma ──
        if not bosh(ws, r, 'shartnoma'):
            k.append('Shartnoma raqami yo‘q')
        if not bosh(ws, r, 'tel'):
            og.append('telefon raqami yo‘q')
        if not bosh(ws, r, 'otasi'):
            og.append('otasining ismi yo‘q')

        # ── ism mosligi ──
        mos = bosh(ws, r, 'moslik')
        if 'MOS EMAS' in mos.upper():
            k.append('Ism mosligi: ' + mos)
        elif mos.upper() == 'TEKSHIRILMAGAN' or not mos:
            og.append('ism mosligi tekshirilmagan')

        # ── fayl holati ──
        if status == 'FAYL_YOQ':
            k.insert(0, 'SHARTNOMA FAYLI YO‘Q')
        elif status == 'HUJJAT XATO':
            k.insert(0, 'FAYL BOSHQA TALABANING MA’LUMOTINI SAQLAYDI')
        elif status == 'CHALA':
            og.append('fayl chala — hamma rasm yuklanmagan')

        talabalar.append(dict(
            qator=r, ism=ism, guruh=bosh(ws, r, 'guruh') or '—',
            shartnoma=bosh(ws, r, 'shartnoma') or '—',
            tasdiq=bosh(ws, r, 'tasdiq'), status=status,
            kamchilik=k, ogoh=og, izoh=izoh))

    # ───────────────────── hisobot ─────────────────────
    muam = [t for t in talabalar if t['kamchilik']]
    faqat_og = [t for t in talabalar if not t['kamchilik'] and t['ogoh']]
    toza = [t for t in talabalar if not t['kamchilik'] and not t['ogoh']]

    print('=' * 80)
    print('  HUJJAT KAMCHILIKLARI AUDITI')
    print('=' * 80)
    print(f"  Jami talaba              : {len(talabalar)}")
    print(f"  ✅ Hujjatlari to'liq      : {len(toza)}")
    print(f"  ⚠️  Faqat mayda eslatma   : {len(faqat_og)}")
    print(f"  ❌ Jiddiy kamchilik bor   : {len(muam)}")
    print()

    sanoq = collections.Counter(x for t in talabalar for x in t['kamchilik'])
    print('-' * 80)
    print('  KAMCHILIK TURLARI BO‘YICHA')
    print('-' * 80)
    for nom, n in sanoq.most_common():
        print(f"  {n:>4} ta   {nom}")
    print()

    print('-' * 80)
    print(f"  JIDDIY KAMCHILIGI BOR {len(muam)} TALABA")
    print('-' * 80)
    for t in sorted(muam, key=lambda x: (x['guruh'], x['ism'])):
        print(f"\n  {t['ism']}   [{t['guruh']} · shartnoma {t['shartnoma']} · {t['tasdiq']}]")
        for x in t['kamchilik']:
            print(f"       ❌ {x}")
        for x in t['ogoh']:
            print(f"       ·  {x}")
        if t['izoh']:
            print(f"       izoh: {t['izoh'][:150]}")

    if faqat_og:
        print()
        print('-' * 80)
        print(f"  MAYDA ESLATMA ({len(faqat_og)} talaba)")
        print('-' * 80)
        for t in sorted(faqat_og, key=lambda x: (x['guruh'], x['ism'])):
            print(f"  {t['ism']:<28} [{t['guruh']}]  {'; '.join(t['ogoh'])}")

    # ───────────────────── CSV ─────────────────────
    with io.open(OUT, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Qator', 'F.I.O', 'Guruh', 'Shartnoma', 'Tasdiq', 'Fayl holati',
                    'Jiddiy kamchiliklar', 'Eslatmalar', 'Izoh'])
        for t in sorted(talabalar, key=lambda x: (not x['kamchilik'], x['guruh'], x['ism'])):
            w.writerow([t['qator'], t['ism'], t['guruh'], t['shartnoma'], t['tasdiq'],
                        t['status'], ' | '.join(t['kamchilik']), ' | '.join(t['ogoh']),
                        t['izoh']])
    print()
    print('=' * 80)
    print(f"  CSV saqlandi: {os.path.basename(OUT)}")
    print('=' * 80)


if __name__ == '__main__':
    main()
