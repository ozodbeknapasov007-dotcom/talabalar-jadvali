# -*- coding: utf-8 -*-
"""
Shartnoma .docx dagi telefon raqamlarini asosiy bazaga (Talabalar_Toliq_Royxati.xlsx, 20-ustun) yozish.

Faqat telefoni BO'SH talabalar to'ldiriladi — mavjud raqamlarga tegilmaydi.
Talaba satri shartnoma raqami (5-ustun) + ism (2-ustun) bo'yicha topiladi.
Keyin: python qayta_tekshiruv/03_hisobot_yasat.py (JSON / hisobotlarni qayta yaratadi).
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from shubhali_malumotlar_audit import clean_text, load_students  # noqa: E402

sys.stdout.reconfigure(encoding='utf-8')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = os.path.join(ROOT, 'data', 'Talabalar_Toliq_Royxati.xlsx')
COL_ISM, COL_SHNUM, COL_TEL = 2, 5, 20


def main(dry_run: bool = False):
    found = [s for s in load_students() if s['_n'].get('tel_manba') not in ('baza', '')]
    wb = openpyxl.load_workbook(EXCEL)
    ws = wb.active
    written, skipped = 0, []
    for s in found:
        ism = clean_text(s.get('ism')).lower()
        shnum = str(s.get('shnum') or '').strip()
        rows = [r for r in range(2, ws.max_row + 1)
                if clean_text(ws.cell(r, COL_ISM).value).lower() == ism
                and str(ws.cell(r, COL_SHNUM).value or '').strip() == shnum]
        if len(rows) != 1:
            skipped.append(f"{s['fish']} ({len(rows)} ta satr)")
            continue
        cell = ws.cell(rows[0], COL_TEL)
        if str(cell.value or '').strip():
            skipped.append(f"{s['fish']} (Excelda telefon bor)")
            continue
        cell.value = ' / '.join(s['_n']['phones'])
        written += 1
    if not dry_run:
        wb.save(EXCEL)
    print(f"{'[sinov] ' if dry_run else ''}Yozildi: {written} ta talaba | o'tkazib yuborildi: {len(skipped)}")
    for x in skipped:
        print('  -', x)


if __name__ == '__main__':
    main(dry_run='--sinov' in sys.argv)
