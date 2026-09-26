# -*- coding: utf-8 -*-
"""
BAZANI QAYTA QURISH: "to'liq 1-KURS 2026-2027.xlsx" dan
========================================================
- Talabalar_Toliq_Royxati.xlsx dagi barcha eski qatorlar o'chiriladi (sarlavha qoladi)
- Manba fayldan talabalar ko'chiriladi, GURUHI "N" bo'lganlar TASHLAB KETILADI
- Ismlar, pasport, shahodatnoma va yo'nalishlar standart ko'rinishga keltiriladi
- Ishga tushirishdan oldin avtomatik zaxira nusxa olinadi
"""

import openpyxl
import os
import re
import sys
import shutil
import datetime
from collections import Counter

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TARGET = os.path.join(BASE_DIR, 'data', 'Talabalar_Toliq_Royxati.xlsx')
SOURCE = os.path.join(BASE_DIR, 'data', 'manba', "to'liq 1-KURS 2026-2027.xlsx")
BACKUP_DIR = os.path.join(BASE_DIR, 'arxiv', 'backup')

HEADER_ROW = 12          # manbadagi sarlavha qatori
DATA_START = 13          # manbadagi birinchi talaba qatori
SKIP_GROUP = 'N'         # hisobga olinmaydigan guruh

DRY_RUN = '--dry-run' in sys.argv   # faqat ko'rsatadi, hech narsa yozmaydi


def nbsp_clean(v):
    """None, NBSP va ortiqcha bo'shliqlarni tozalaydi"""
    if v is None:
        return ''
    s = str(v).replace('\xa0', ' ').replace('​', '')
    return re.sub(r'\s+', ' ', s).strip()


def clean_uz_name(text):
    """O'zbekcha ism-familiyani standart ko'rinishga keltiradi (loyiha uslubida)"""
    text = nbsp_clean(text)
    if not text:
        return ''
    text = re.sub(r"[`‘’ʻʼ´]", "'", text)

    out = []
    for w in text.split(' '):
        wl = w.lower()
        if wl in ('qizi', 'kizi', 'qzy'):
            out.append('qizi')
        elif wl in ("o'g'li", "ogli", "ugli", "o'gli", "uli"):
            out.append("o'g'li")
        elif "'" in w:
            parts = w.split("'")
            out.append("'".join(p.capitalize() if i == 0 else p.lower()
                                for i, p in enumerate(parts)))
        else:
            out.append(w.capitalize())
    return ' '.join(out)


def _split_doc_id(v):
    """'ae2261219' -> ('AE', '2261219'); mos kelmasa (None, tozalangan matn)"""
    s = nbsp_clean(v)
    if not s:
        return (None, '')
    m = re.match(r"^([A-Za-z]{1,3})\s*[-–]?\s*(\d{5,10})$", s)
    if m:
        return (m.group(1).upper(), m.group(2))
    return (None, s.upper() if len(s) <= 14 else s)


def clean_passport(v):
    """Pasport / ID-karta — bazadagi uslub: bo'shliqsiz ('AC1967100')"""
    ser, num = _split_doc_id(v)
    return f"{ser}{num}" if ser else num


def clean_cert(v):
    """Shahodatnoma / Diplom — bazadagi uslub: bo'shliq bilan ('UM 0438281')"""
    ser, num = _split_doc_id(v)
    return f"{ser} {num}" if ser else num


def clean_pinfl(v):
    """PINFL — faqat raqamlar, matn ko'rinishida saqlanadi"""
    s = re.sub(r'\D', '', nbsp_clean(v))
    return s


def clean_date(v):
    """Sanani DD.MM.YYYY ko'rinishiga keltiradi"""
    if isinstance(v, datetime.datetime):
        return v.strftime('%d.%m.%Y')
    s = nbsp_clean(v)
    if not s:
        return ''
    m = re.match(r'^(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})$', s)
    if m:
        return f"{int(m.group(1)):02d}.{int(m.group(2)):02d}.{m.group(3)}"
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    return s


def clean_sentence(v):
    """Muassasa nomi kabi matnlar — birinchi harf katta"""
    s = nbsp_clean(v)
    if not s:
        return ''
    return s[0].upper() + s[1:]


def detect_doc_type(maktab, cert):
    """Ta'lim muassasasi turini aniqlaydi"""
    m = (maktab or '').lower()
    if any(k in m for k in ('kollej', 'texnikum', 'litsey', 'bilim yurti')):
        return 'Diplom'
    if 'maktab' in m:
        return 'Shahodatnoma'
    return 'Shahodatnoma' if cert else ''


def calc_status(pv, pinfl, dob, cert):
    if pv and pinfl and dob and cert:
        return 'TOPILDI'
    if pv or cert or pinfl:
        return 'CHALA'
    return 'FAYL_YOQ'


# ══════════════════════════════════════════════════════════
# 1. ZAXIRA NUSXA
# ══════════════════════════════════════════════════════════
if DRY_RUN:
    print("🔍 DRY-RUN rejimi — hech qanday fayl o'zgartirilmaydi\n")
else:
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    backup_path = os.path.join(BACKUP_DIR, f'Talabalar_Toliq_Royxati_QAYTA_QURISHDAN_OLDIN_{stamp}.xlsx')
    shutil.copy2(TARGET, backup_path)
    print(f"✅ Zaxira nusxa: backup/{os.path.basename(backup_path)}")

# ══════════════════════════════════════════════════════════
# 2. MANBANI O'QISH
# ══════════════════════════════════════════════════════════
src_wb = openpyxl.load_workbook(SOURCE, data_only=True)
src_ws = src_wb.active

raw = list(src_ws.iter_rows(min_row=DATA_START, max_row=src_ws.max_row, values_only=True))
print(f"📖 Manbada jami qator: {len(raw)}")

students = []
skipped_n = 0
skipped_empty = 0
seen = {}
duplicates = []

for row in raw:
    kelgan, guruh, _tr, fam, ism_, ota, _fish, dob, pv, pinfl, ber, cert, maktab, yon, yil, shnum, wfile = row[:17]

    guruh = nbsp_clean(guruh)
    fam = clean_uz_name(fam)
    ism_ = clean_uz_name(ism_)

    if guruh == SKIP_GROUP:
        skipped_n += 1
        continue
    if not fam and not ism_:
        skipped_empty += 1
        continue

    ifo = f"{fam} {ism_}".strip()
    ota_c = clean_uz_name(ota)
    fish = f"{ifo} {ota_c}".strip()

    pv_c = clean_passport(pv)
    pinfl_c = clean_pinfl(pinfl)
    dob_c = clean_date(dob)
    ber_c = clean_date(ber)
    cert_c = clean_cert(cert)
    maktab_c = clean_sentence(maktab)
    yon_c = clean_sentence(yon)
    yil_c = nbsp_clean(yil)
    shnum_c = nbsp_clean(shnum)

    key = (ifo.lower(), guruh)
    rec = {
        'guruh': guruh, 'ifo': ifo, 'ota': ota_c, 'fish': fish, 'yon': yon_c,
        'shnum': shnum_c, 'kelgan': kelgan, 'pv': pv_c, 'pinfl': pinfl_c,
        'dob': dob_c, 'ber': ber_c, 'cert': cert_c, 'maktab': maktab_c,
        'yil': yil_c, 'wfile': nbsp_clean(wfile),
        'doc_tur': detect_doc_type(maktab_c, cert_c),
    }
    rec['status'] = calc_status(pv_c, pinfl_c, dob_c, cert_c)

    if key in seen:
        # Dublikat: ma'lumoti to'liqrog'ini qoldiramiz
        old = seen[key]
        old_score = sum(1 for k in ('pv', 'pinfl', 'dob', 'cert', 'shnum') if old[k])
        new_score = sum(1 for k in ('pv', 'pinfl', 'dob', 'cert', 'shnum') if rec[k])
        duplicates.append(f"{ifo} ({guruh})")
        if new_score > old_score:
            students[students.index(old)] = rec
            seen[key] = rec
        continue

    seen[key] = rec
    students.append(rec)

# Guruh, keyin alifbo tartibida
students.sort(key=lambda s: (s['guruh'], s['ifo'].lower()))

print(f"⏭️  '{SKIP_GROUP}' guruhi tashlab ketildi: {skipped_n} ta")
if skipped_empty:
    print(f"⏭️  Bo'sh qatorlar: {skipped_empty} ta")
if duplicates:
    print(f"⚠️  Dublikat birlashtirildi ({len(duplicates)} ta): {', '.join(duplicates)}")
print(f"✅ Qo'shiladigan talabalar: {len(students)} ta")

# ══════════════════════════════════════════════════════════
# 3. ESKI BAZANI TOZALASH
# ══════════════════════════════════════════════════════════
wb = openpyxl.load_workbook(TARGET)
ws = wb.active
old_count = ws.max_row - 1

if DRY_RUN:
    print(f"🗑️  (dry-run) O'chirilishi kerak bo'lgan eski qatorlar: {old_count} ta")
    print("\n  NAMUNA — birinchi 5 ta yangi qator:")
    for s in students[:5]:
        print(f"   {s['guruh']} | {s['fish']:42s} | #{s['shnum']:4s} | "
              f"{s['pv']:10s} | {s['pinfl']:15s} | {s['cert']:12s} | {s['status']}")
else:
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    print(f"🗑️  Eski baza tozalandi: {old_count} ta qator o'chirildi")

# ══════════════════════════════════════════════════════════
# 4. YANGI MA'LUMOTLARNI YOZISH
# ══════════════════════════════════════════════════════════
for i, s in enumerate([] if DRY_RUN else students, start=1):
    r = i + 1
    ws.cell(row=r, column=1, value=i)                   # T/R
    ws.cell(row=r, column=2, value=s['ifo'])            # I.F.O
    ws.cell(row=r, column=3, value=s['yon'])            # Yo'nalishi
    ws.cell(row=r, column=4, value='')                  # To'lov statusi (manbada yo'q)
    ws.cell(row=r, column=5, value=s['shnum'])          # Shartnoma raqami
    ws.cell(row=r, column=6,                            # Sanasi (kelgan sanasi)
            value=s['kelgan'].strftime('%Y-%m-%d %H:%M:%S')
            if isinstance(s['kelgan'], datetime.datetime) else nbsp_clean(s['kelgan']))
    ws.cell(row=r, column=7, value=s['ota'])            # Otasining ismi
    ws.cell(row=r, column=8, value=s['fish'])           # To'liq F.I.SH
    ws.cell(row=r, column=9, value='')                  # Passport bo'yicha F.I.SH (manbada yo'q)
    ws.cell(row=r, column=10, value=s['pv'])            # Passport / ID
    c = ws.cell(row=r, column=11, value=s['pinfl'])     # PINFL
    c.number_format = '@'
    ws.cell(row=r, column=12, value=s['ber'])           # Berilgan sanasi
    ws.cell(row=r, column=13, value=s['dob'])           # Tug'ilgan sanasi
    ws.cell(row=r, column=14, value='')                 # Shahodatnoma bo'yicha F.I.SH
    ws.cell(row=r, column=15, value=s['cert'])          # Shahodatnoma / Diplom
    ws.cell(row=r, column=16, value='')                 # QR havolasi
    ws.cell(row=r, column=17, value=s['maktab'])        # Tugatgan o'qish joyi
    ws.cell(row=r, column=18, value=s['doc_tur'])       # Muassasa turi
    ws.cell(row=r, column=19, value=s['yil'])           # Bitirgan yili
    ws.cell(row=r, column=20, value='')                 # Telefon (manbada yo'q)
    ws.cell(row=r, column=21, value=s['status'])        # Qidiruv holati
    ws.cell(row=r, column=22, value='')                 # Izoh
    ws.cell(row=r, column=23, value=s['guruh'])         # Guruh

if not DRY_RUN:
    wb.save(TARGET)
    print(f"💾 Saqlandi: {os.path.basename(TARGET)} ({len(students)} ta talaba)")

# ══════════════════════════════════════════════════════════
# 5. HISOBOT
# ══════════════════════════════════════════════════════════
print('\n' + '=' * 58)
print('  YAKUNIY HOLAT')
print('=' * 58)
gc = Counter(s['guruh'] for s in students)
for g in sorted(gc):
    print(f"  Guruh {g}: {gc[g]:3d} ta")
print('-' * 58)
st = Counter(s['status'] for s in students)
for k in ('TOPILDI', 'CHALA', 'FAYL_YOQ'):
    print(f"  {k:10s}: {st.get(k, 0):3d} ta")
print('-' * 58)
for y, n in Counter(s['yon'] or '(ko\'rsatilmagan)' for s in students).most_common():
    print(f"  {y}: {n} ta")
print('-' * 58)
print(f"  Shartnoma raqami bor : {sum(1 for s in students if s['shnum'])} ta")
print(f"  Word fayli ko'rsatilgan: {sum(1 for s in students if s['wfile'])} ta")
print('=' * 58)
