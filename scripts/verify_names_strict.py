# -*- coding: utf-8 -*-
"""
QAT'IY ISM TEKSHIRUVI
=====================
Har bir talabaning ismini UCH manbada yonma-yon solishtiradi:
  1. Sizning ro'yxatingiz (Excel: I.F.O + Otasining ismi)
  2. Pasport / ID-karta MRZ (transliteratsiya qilingan, apostrofsiz, BOSH HARF)
  3. e-shahodatnoma.uz rasmiy PDF (to'liq o'zbek lotinida, otasining ismi bilan)

MUHIM: MRZ xalqaro ICAO standartida yoziladi — apostrof (') tashlanadi,
ba'zi harflar almashadi (q->K, o'->O). Shuning uchun MRZ bilan "harfma-harf"
mos kelishi SHART EMAS. Skript transliteratsiya bilan izohlanadigan farqni
haqiqiy xatodan ajratadi.

Shahodatnoma PDF esa to'liq o'zbek lotinida — u bilan 100% mos kelishi SHART.

    python scripts/verify_names_strict.py
"""
import openpyxl, os, sys, re, json, io, unicodedata

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, 'data', 'Talabalar_Toliq_Royxati.xlsx')
RESULT_JSON = os.path.join(BASE_DIR, 'arxiv', 'scratch', 'verify_results.json')
CACHE_DIR = os.path.join(BASE_DIR, 'arxiv', 'scratch', 'qr_cache')
REPORT = os.path.join(BASE_DIR, 'hisobotlar', 'tekshiruvlar', 'ISMLAR_TEKSHIRUVI.txt')

try:
    import pypdf
except ImportError:
    pypdf = None


def nb(v):
    return re.sub(r'\s+', ' ', str(v if v is not None else '').replace('\xa0', ' ')).strip()


def apos(s):
    """Barcha apostrof ko'rinishlarini bittaga keltiradi"""
    return re.sub(r"[`‘’ʻʼ´ʹ']", "'", nb(s))


def cmp_key(s):
    """Shahodatnoma bilan qat'iy solishtirish uchun: apostrof birxil, kichik harf"""
    return apos(s).lower()


def mrz_forms(name):
    """
    O'zbek lotinidagi ismning MRZ da uchrashi mumkin bo'lgan ko'rinishlari.
    ICAO MRZ faqat A-Z ishlatadi: apostrof tashlanadi, ayrim harflar almashadi.
    """
    s = apos(name).upper().replace("'", '')
    out = {s}
    # Ma'lum transliteratsiya juftliklari (ikkala yo'nalishda ham sinaladi)
    subs = [('Q', 'K'), ('X', 'H'), ('TS', 'S'), ('YO', 'O'), ('YU', 'U'),
            ('SH', 'S'), ('CH', 'C'), ('Ō', 'O'), ('Ğ', 'G')]
    for a, b in subs:
        for v in list(out):
            if a in v:
                out.add(v.replace(a, b))
            if b in v:
                out.add(v.replace(b, a))
    # G' -> G allaqachon apostrof tashlanishi bilan hal bo'lgan
    return out


def levenshtein(a, b):
    if a == b:
        return 0
    if not a or not b:
        return max(len(a), len(b))
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


# Transliteratsiya juftliklari: hujjatlarda bir xil ism turlicha yozilishi mumkin
TRANSLIT = [('q', 'k'), ('x', 'h'), ('sh', 's'), ('ch', 'c'),
            ('yo', 'o'), ('yu', 'u'), ('ts', 's')]


def canon(w):
    """
    So'zni kanonik shaklga keltiradi: hujjatlarda bir xil tovush turlicha
    yozilishi mumkin bo'lgan harflarni bitta vakilga birlashtiradi.
      x/h -> h     (Shoxjaxon / Shohjaxon, Axmadova / Ahmadova)
      q/k -> k     (Arziqulova / Arzikulova)
      o'/u/o -> o  (O'roqova / Uroqova, Ruhshona / Rohshona)
    Har bir harf alohida almashtiriladi, shuning uchun so'z ichidagi
    qisman farq ham to'g'ri aniqlanadi.
    """
    w = apos(w).lower()
    w = w.replace("o'", 'o').replace("g'", 'g').replace("'", '')
    w = w.replace('x', 'h').replace('q', 'k')
    w = w.replace('u', 'o')
    return w


def word_kind(bw, cw):
    """
    Ikki so'z qanday farq qiladi?
      'mos'      — aynan bir xil
      'apostrof' — faqat tutuq belgisi (') farq qiladi
      'translit' — harf almashinuvi bilan izohlanadi (q/k, x/h, o'/u)
      'boshqa'   — haqiqiy boshqa yozuv
    """
    b, c = apos(bw).lower(), apos(cw).lower()
    if b == c:
        return 'mos'
    if b.replace("'", '') == c.replace("'", ''):
        return 'apostrof'
    if canon(b) == canon(c):
        return 'translit'
    return 'boshqa'


RANK = {'mos': 0, 'apostrof': 1, 'translit': 2, 'boshqa': 3}


def classify_fish(base, cert):
    """To'liq F.I.SH ni so'zma-so'z tahlil qiladi -> (umumiy_turi, [so'z tafsiloti])"""
    bws, cws = apos(base).split(), apos(cert).split()
    if len(bws) != len(cws):
        return 'boshqa', [(' '.join(bws), ' '.join(cws), 'boshqa')]
    det = [(b, c, word_kind(b, c)) for b, c in zip(bws, cws)]
    worst = max(det, key=lambda d: RANK[d[2]])[2]
    return worst, det


def fmt_detail(det):
    """Farq qilgan so'zlarni ko'rsatadi"""
    parts = []
    for b, c, k in det:
        if k == 'mos':
            continue
        parts.append(f"{b} → {c}  ({k})")
    return '; '.join(parts)


def raw_cert_name(url):
    """Keshlangan PDF dan F.I.SH satrini ASL holida (o'zgartirmasdan) oladi"""
    if not pypdf:
        return ''
    key = re.sub(r'\W+', '_', url)[-120:]
    f = os.path.join(CACHE_DIR, key + '.pdf')
    if not os.path.exists(f):
        return ''
    try:
        reader = pypdf.PdfReader(f)
        text = "\n".join((p.extract_text() or '') for p in reader.pages)
    except Exception:
        return ''
    for l in [x.strip() for x in text.split('\n') if x.strip()]:
        if re.search(r"(?i)\b(qizi|o'g'li|o‘g‘li|ogli|ugli|kizi)\.?$", l) and len(l.split()) >= 3:
            return l
    return ''


# ══════════════════════════════════════════════════════════
results = json.load(open(RESULT_JSON, encoding='utf-8'))
by_row = {r['row']: r for r in results}
ws = openpyxl.load_workbook(EXCEL_PATH).active

rows = []
for r in range(2, ws.max_row + 1):
    ifo = nb(ws.cell(row=r, column=2).value)
    if not ifo:
        continue
    rows.append({
        'row': r, 'ifo': ifo,
        'ota': nb(ws.cell(row=r, column=7).value),
        'fish': nb(ws.cell(row=r, column=8).value),
        'guruh': nb(ws.cell(row=r, column=23).value),
    })

cert_ok, cert_bad, cert_none = [], [], []
mrz_ok, mrz_translit, mrz_bad, mrz_none = [], [], [], []

for x in rows:
    res = by_row.get(x['row'], {})
    qr = res.get('qr') or {}
    trusted = res.get('trusted', True)
    x['trusted'] = trusted
    x['file'] = res.get('file')

    # ── SHAHODATNOMA (qat'iy) ──────────────────────────
    url = qr.get('qr')
    raw = raw_cert_name(url) if url else ''
    x['cert_raw'] = raw
    if not raw:
        cert_none.append(x)
    else:
        base_full = x['fish'] or f"{x['ifo']} {x['ota']}".strip()
        x['cert_base'] = base_full
        if cmp_key(raw) == cmp_key(base_full):
            cert_ok.append(x)
        else:
            kind, det = classify_fish(base_full, raw)
            x['cert_kind'], x['cert_det'] = kind, det
            cert_bad.append(x)

    # ── PASPORT MRZ ────────────────────────────────────
    mrz = nb(qr.get('mrz_name'))
    x['mrz_raw'] = mrz
    if not mrz:
        mrz_none.append(x)
    else:
        base_two = ' '.join(apos(x['ifo']).split()[:2])
        # MRZ da apostrof umuman bo'lmaydi — shuning uchun uni hisobga olmaymiz
        if mrz.upper().replace(' ', '') == apos(base_two).upper().replace("'", '').replace(' ', ''):
            mrz_ok.append(x)
        elif canon(mrz).replace(' ', '') == canon(base_two).replace(' ', ''):
            mrz_translit.append(x)
        else:
            x['mrz_base'] = base_two
            x['mrz_dist'] = levenshtein(mrz.upper().replace(' ', ''),
                                        apos(base_two).upper().replace("'", '').replace(' ', ''))
            mrz_bad.append(x)

# ══════════════════════════════════════════════════════════
L = []


def p(s=''):
    L.append(s)
    print(s)


p('=' * 80)
p('  QAT\'IY ISM TEKSHIRUVI — RO\'YXAT / PASPORT / SHAHODATNOMA')
p('=' * 80)
p(f"  Jami talaba                          : {len(rows)}")
p('-' * 80)
p('  SHAHODATNOMA PDF (100% mos kelishi shart):')
p(f"     ✅ harfma-harf mos                : {len(cert_ok)}")
p(f"     ❌ haqiqiy farq                   : "
  f"{len([x for x in cert_bad if x.get('cert_kind') == 'boshqa' and x['trusted']])}")
p(f"     ⚠️  tutuq belgisi (') farqi        : "
  f"{len([x for x in cert_bad if x.get('cert_kind') == 'apostrof' and x['trusted']])}")
p(f"     ≈  transliteratsiya (q/k, x/h)    : "
  f"{len([x for x in cert_bad if x.get('cert_kind') == 'translit' and x['trusted']])}")
p(f"     🛑 karantin (boshqa odam)         : {len([x for x in cert_bad if not x['trusted']])}")
p(f"     —  shahodatnoma o'qilmagan        : {len(cert_none)}")
p('-' * 80)
p('  PASPORT / ID-KARTA MRZ:')
p(f"     ✅ aynan mos                      : {len(mrz_ok)}")
p(f"     ≈  transliteratsiya bilan izohli  : {len(mrz_translit)}")
p(f"     ❌ FARQ BOR                       : {len(mrz_bad)}")
p(f"     —  pasport o'qilmagan             : {len(mrz_none)}")
p('=' * 80)

if cert_bad:
    GROUPS = [
        ('boshqa', '❌ HAQIQIY FARQ — ism boshqacha yozilgan',
         'Bu haqiqiy xato: ro\'yxatdagi yozuv shahodatnomaga umuman mos emas.'),
        ('apostrof', "⚠️  TUTUQ BELGISI (') FARQI",
         "Ro'yxatda apostrof tushib qolgan yoki ortiqcha — rasmiy hujjatda bor."),
        ('translit', '≈ TRANSLITERATSIYA FARQI (q/k, x/h, o\'/u)',
         'Bir ism ikki hujjatda turlicha yozilgan — qaysi birini asos qilishni siz hal qilasiz.'),
    ]
    for kind, title, note in GROUPS:
        grp = [x for x in cert_bad if x.get('cert_kind') == kind and x['trusted']]
        if not grp:
            continue
        p('')
        p('█' * 80)
        p(f'  {title}  ({len(grp)} ta)')
        p(f'     {note}')
        p('█' * 80)
        for x in grp:
            p(f"\n  [{x['guruh']}] qator {x['row']}")
            p(f"      ro'yxatda    : {x['cert_base']}")
            p(f"      shahodatnoma : {x['cert_raw']}")
            p(f"      farq         : {fmt_detail(x['cert_det'])}")

    kar_grp = [x for x in cert_bad if not x['trusted']]
    if kar_grp:
        p('')
        p('█' * 80)
        p(f'  🛑 KARANTIN — butunlay boshqa odamning shahodatnomasi ({len(kar_grp)} ta)')
        p('█' * 80)
        for x in kar_grp:
            p(f"\n  [{x['guruh']}] qator {x['row']}  (fayl: {x['file']})")
            p(f"      ro'yxatda    : {x['cert_base']}")
            p(f"      shahodatnoma : {x['cert_raw']}")

if mrz_bad:
    p('')
    p('█' * 80)
    p('  ❌ PASPORTDAGI F.I.SH RO\'YXATGA MOS EMAS')
    p('     (transliteratsiya bilan izohlab bo\'lmaydigan farqlar)')
    p('█' * 80)
    for x in sorted(mrz_bad, key=lambda y: -y['mrz_dist']):
        kar = '   [KARANTIN — fayl boshqa odamniki]' if not x['trusted'] else ''
        p(f"\n  [{x['guruh']}] qator {x['row']}  (farq: {x['mrz_dist']} harf){kar}")
        p(f"      ro'yxatda : {x['mrz_base']}")
        p(f"      pasportda : {x['mrz_raw']}")
        p(f"      fayl      : {x['file']}")

if mrz_translit:
    p('')
    p('-' * 80)
    p(f"  ≈ TRANSLITERATSIYA BILAN IZOHLANADI ({len(mrz_translit)} ta) — xato emas")
    p("    MRZ da apostrof tashlanadi va ayrim harflar almashadi (q->K, x->H)")
    p('-' * 80)
    for x in mrz_translit:
        p(f"      [{x['guruh']}] {x['ifo']:30s} pasportda: {x['mrz_raw']}")

p('')
p('-' * 80)
p(f"  TEKSHIRIB BO'LMAGANLAR")
p('-' * 80)
p(f"  Shahodatnomasi o'qilmagan: {len(cert_none)} ta")
p(f"  Pasporti o'qilmagan      : {len(mrz_none)} ta")
both = [x for x in cert_none if not nb((by_row.get(x['row'], {}).get('qr') or {}).get('mrz_name'))]
p(f"  Ikkalasi ham yo'q        : {len(both)} ta")
for x in both:
    p(f"      [{x['guruh']}] {x['ifo']:30s} fayl={x['file'] or '-'}")

p('')
p('=' * 80)

with open(REPORT, 'w', encoding='utf-8') as fh:
    fh.write('\n'.join(L))
print(f"\n📄 Hisobot: ISMLAR_TEKSHIRUVI.txt")

# ── Excel / dashboard uchun mashina o'qiydigan natija ──────
STATUS = {}
for x in cert_ok:
    STATUS.setdefault(x['row'], {})['cert'] = 'mos'
for x in cert_bad:
    STATUS.setdefault(x['row'], {})['cert'] = (
        'boshqa_odam' if not x['trusted'] else x.get('cert_kind', 'boshqa'))
for x in mrz_ok:
    STATUS.setdefault(x['row'], {})['pass'] = 'mos'
for x in mrz_translit:
    STATUS.setdefault(x['row'], {})['pass'] = 'translit'
for x in mrz_bad:
    STATUS.setdefault(x['row'], {})['pass'] = (
        'boshqa_odam' if not x['trusted'] else 'boshqa')

out = {}
for x in rows:
    st = STATUS.get(x['row'], {})
    out[str(x['row'])] = {
        'ifo': x['ifo'],
        'pass_fish': x.get('mrz_raw', ''),
        'cert_fish': x.get('cert_raw', ''),
        'pass_holat': st.get('pass', ''),
        'cert_holat': st.get('cert', ''),
        'trusted': x.get('trusted', True),
    }
json.dump(out, open(os.path.join(BASE_DIR, 'arxiv', 'scratch', 'name_check.json'), 'w',
                    encoding='utf-8'), ensure_ascii=False, indent=1)
print('💾 scratch/name_check.json')
