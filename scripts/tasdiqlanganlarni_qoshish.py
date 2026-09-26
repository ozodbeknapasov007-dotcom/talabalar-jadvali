# -*- coding: utf-8 -*-
"""
TASDIQLANGAN 22 TALABANI ASOSIY RO'YXATGA QO'SHISH
===================================================
Manba : TEKSHIRUV/data.js  (OCR qilingan hujjat ma'lumotlari)
        Downloads/tekshiruv_natija.csv  (operator qarori: TASDIQLANGAN / RAD ETILGAN)
Nishon : Talabalar_Toliq_Royxati.xlsx  ->  "Barcha Talabalar" varag'i

Faqat TASDIQLANGAN yozuvlar yoziladi. RAD ETILGANlarga izoh qo'yiladi.
Har bir o'zgarish (eski -> yangi) ekranga chiqariladi.
--dry  bayrog'i bilan ishga tushirilsa hech narsa saqlanmaydi (faqat hisobot).
"""
import csv, io, os, re, shutil, sys, datetime
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, 'data', 'Talabalar_Toliq_Royxati.xlsx')
CSVP = os.path.join(os.path.expanduser('~'), 'Downloads', 'tekshiruv_natija.csv')
DRY = '--dry' in sys.argv

# ustun raqamlari (1 dan)
C = dict(tr=1, ifo=2, yonalish=3, tolov=4, shartnoma=5, sana=6, otasi=7, toliq=8,
         pasfio=9, pasraqam=10, pinfl=11, berilgan=12, tugilgan=13, shfio=14,
         shseriya=15, qr=16, maktab=17, turi=18, bitirgan=19, tel=20,
         status=21, izoh=22, guruh=23, moslik=24, tasdiq=25)


# ─── data.js ni o'qish ────────────────────────────────────────────────────
def yozuvlar():
    src = io.open(os.path.join(BASE, 'hujjatlar', 'TEKSHIRUV', 'data.js'), encoding='utf-8').read()
    out = []
    for blok in src.split('\n{\n')[1:]:
        def sub(nom):
            m = re.search(nom + r':\{(.*?)\},?\n', blok, re.S)
            return m.group(1) if m else ''
        def g(key, ichida):
            m = re.search(r'\b' + key + r':"([^"]*)"', ichida)
            return m.group(1).strip() if m else ''
        pas, huj, sh = sub('pasport'), sub('hujjat'), sub('shartnoma')
        out.append(dict(
            fayl=g('fayl', blok), royxat=g('royxat', blok),
            familiya=g('familiya', pas).split(' (')[0], ism=g('ism', pas),
            otasi=g('otasi', pas), tugilgan=g('tugilgan', pas),
            pinfl=g('jshshir', pas), pasraqam=g('raqam', pas),
            pasberilgan=g('berilgan', pas),
            h_turi=g('turi', huj), h_raqam=g('raqam', huj),
            h_muassasa=g('muassasa', huj), h_berilgan=g('berilgan', huj),
            h_yil=g('yil', huj),
            s_raqam=g('raqam', sh), s_sana=g('sana', sh), s_tel=g('tel', sh),
            s_yonalish=g('yonalish', sh),
        ))
    return out


def norm(s):
    s = (s or '').lower().replace('’', "'").replace('‘', "'").replace('`', "'")
    for a, b in [('kh', 'x'), ('x', 'h'), ('q', 'k'), ("o'", 'o'), ("g'", 'g'),
                 ("'", ''), ('y', 'i'), ('sh', 's'), ('ch', 'c'), ('ye', 'e')]:
        s = s.replace(a, b)
    return re.sub(r'[^a-z]', '', s)


def bosh(s):
    """«YUSUFJON QIZI» -> «Yusufjon qizi», «O‘TKIR QIZI» -> «O‘tkir qizi»,
    «NIGMATOVNA» -> «Nigmatovna». .title() apostrofdan keyin ham bosh harf
    qo'yib yuboradi, shuning uchun qo'lda."""
    s = (s or '').strip().lower()
    if not s:
        return ''
    return s[0].upper() + s[1:]


def familiya_ism(f, i):
    return f"{bosh(f)} {bosh(i)}".strip()


def yil(s):
    m = re.search(r'(19|20)\d{2}', s or '')
    return int(m.group(0)) if m else None


def hujjat_turi(t):
    t = (t or '').lower()
    if 'attestat' in t: return 'Attestat'
    if 'diplom' in t:   return 'Diplom'
    if 'shahodat' in t: return 'Shahodatnoma'
    return None


def main():
    qaror = {}
    for r in csv.DictReader(io.open(CSVP, encoding='utf-8-sig')):
        qaror[r['Fayl']] = (r['Holat'], (r['Izoh'] or '').strip())

    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Barcha Talabalar']

    # mavjud qatorlar indeksi: familiya|ism -> qator
    idx = {}
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(r, C['ifo']).value or '').split()
        if len(v) >= 2:
            idx.setdefault(norm(v[0]) + '|' + norm(v[1]), r)

    ozg, yangi, rad, muam = [], [], [], []

    for y in yozuvlar():
        holat, izoh = qaror.get(y['fayl'], ('', ''))
        nom = (y['familiya'] + ' ' + y['ism']).strip()

        if holat != 'TASDIQLANGAN':
            k = norm(y['royxat'].split()[0]) + '|' + norm(y['royxat'].split()[1])
            r = idx.get(k)
            rad.append((y['royxat'], r, y['fayl']))
            if r and not DRY:
                ws.cell(r, C['status']).value = 'HUJJAT XATO'
                ws.cell(r, C['izoh']).value = (
                    f"{y['fayl']} fayli boshqa talabaning ma'lumotini saqlaydi — "
                    f"hujjatlar qaytadan so'ralsin (tekshiruv 20.09.2026)")
            continue

        # ---- qatorni topish: avval pasport ismi, keyin kanal ro'yxati ----
        r = idx.get(norm(y['familiya']) + '|' + norm(y['ism']))
        if not r and y['royxat']:
            p = y['royxat'].split()
            if len(p) >= 2:
                r = idx.get(norm(p[0]) + '|' + norm(p[1]))
        if not r:
            r = ws.max_row + 1
            ws.cell(r, C['tr']).value = r - 1
            ws.cell(r, C['ifo']).value = familiya_ism(y['familiya'], y['ism'])
            yangi.append((nom, r))

        # ---- yoziladigan qiymatlar ----
        toliq = (familiya_ism(y['familiya'], y['ism']) + ' ' + bosh(y['otasi'])).strip()
        qiy = {
            'ifo':      familiya_ism(y['familiya'], y['ism']),
            'otasi':    bosh(y['otasi']),
            'toliq':    toliq,
            'pasfio':   toliq,
            'pasraqam': y['pasraqam'],
            'pinfl':    y['pinfl'] if re.fullmatch(r'\d{14}', y['pinfl']) else None,
            'berilgan': y['pasberilgan'],
            'tugilgan': y['tugilgan'],
            'tel':      y['s_tel'] if '+' not in y['s_tel'] else None,
            'status':   'TOPILDI',
            'tasdiq':   'TASDIQLANDI',
        }
        # shartnoma raqami — fayl nomi bilan ziddiyat bo'lsa yozilmaydi
        if '(' not in y['s_raqam'] and '!' not in y['s_raqam']:
            qiy['shartnoma'] = y['s_raqam']
        else:
            muam.append((nom, f"shartnoma raqami ziddiyatli: {y['s_raqam']} — qo‘lda kiriting"))
        if y['s_sana']:
            try:
                qiy['sana'] = datetime.datetime.strptime(y['s_sana'], '%d.%m.%Y')
            except ValueError:
                pass
        if y['h_raqam']:
            qiy['shfio']    = toliq
            qiy['shseriya'] = y['h_raqam'].replace('№', '').replace('  ', ' ').strip()
            qiy['maktab']   = y['h_muassasa']
            qiy['turi']     = hujjat_turi(y['h_turi'])
            qiy['bitirgan'] = yil(y['h_berilgan']) or yil(y['h_yil'])
            qiy['moslik']   = 'Pasport: MOS | Shahodatnoma: MOS'
        else:
            muam.append((nom, 'diplom/shahodatnoma rasmi yo‘q — ta’lim hujjati to‘ldirilmadi'))
            qiy['status'] = 'CHALA'

        for k, v in qiy.items():
            if v in (None, ''):
                continue
            hoz = ws.cell(r, C[k]).value
            hoz_s = '' if hoz is None else str(hoz).strip()
            if hoz_s == str(v).strip():
                continue
            ozg.append((nom, k, hoz_s or '(bo‘sh)', str(v)))
            if not DRY:
                ws.cell(r, C[k]).value = v
        if izoh and not DRY:
            eski = str(ws.cell(r, C['izoh']).value or '')
            if izoh not in eski:
                ws.cell(r, C['izoh']).value = (eski + ' | ' if eski else '') + 'Operator: ' + izoh

    # ───────── hisobot ─────────
    print('=' * 78)
    print(f"  ASOSIY RO'YXATNI YANGILASH{'  (DRY RUN — saqlanmadi)' if DRY else ''}")
    print('=' * 78)
    if yangi:
        print(f"\n  YANGI QO'SHILDI ({len(yangi)}):")
        for n, r in yangi:
            print(f"     + {n}  ->  {r}-qator")
    print(f"\n  O'ZGARISHLAR ({len(ozg)}):")
    oxirgi = None
    for n, k, a, b in ozg:
        if n != oxirgi:
            print(f"\n   {n}")
            oxirgi = n
        print(f"      {k:<10} {a[:34]:<36} ->  {b[:40]}")
    if rad:
        print(f"\n  RAD ETILGAN ({len(rad)}) — ma'lumot yozilmadi, izoh qo'yildi:")
        for n, r, f in rad:
            print(f"     ✕ {n:<32} {r}-qator   ({f})")
    if muam:
        print(f"\n  QO'LDA TO'LDIRISH KERAK ({len(muam)}):")
        for n, m in muam:
            print(f"     ! {n:<32} {m}")

    if not DRY:
        zax = XLSX.replace('.xlsx', f".backup_{datetime.datetime.now():%Y%m%d_%H%M}.xlsx")
        shutil.copy2(XLSX, zax)
        wb.save(XLSX)
        print(f"\n  Zaxira : {os.path.basename(zax)}")
        print(f"  Saqlandi: {os.path.basename(XLSX)}")
    print('=' * 78)


if __name__ == '__main__':
    main()
