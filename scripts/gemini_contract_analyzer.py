# -*- coding: utf-8 -*-
"""
GEMINI FLASH AI VISION & CONTRACT ANALYZER (V2 - RETRY BILAN)
=============================================================
"""

import os
import sys
import json
import base64
import docx
import urllib.request
import time

sys.stdout.reconfigure(encoding='utf-8')

GEMINI_API_KEY = "AQ.Ab8RN6KLUpk8Ag9edtF89-0kv32SYAY0gPevAsvOoZTI1dIt7A"
MODEL_NAME = "gemini-flash-latest"
API_URL = f"https://generativelanguage.googleapis.com/v1beta/models/{MODEL_NAME}:generateContent?key={GEMINI_API_KEY}"

def analyze_docx_with_gemini(docx_path, max_retries=3):
    if not os.path.exists(docx_path):
        return None
    
    fname = os.path.basename(docx_path)
    doc = docx.Document(docx_path)
    
    # 1. Matnlarni yig'ish
    doc_texts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t: doc_texts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                doc_texts.append(" | ".join(cells))
    full_text = "\n".join(doc_texts)

    # 2. Rasmlarni ajratish (Inline base64)
    image_parts = []
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            blob = rel.target_part.blob
            if len(blob) > 5000:
                b64 = base64.b64encode(blob).decode('utf-8')
                mime = "image/jpeg"
                if rel.target_ref.lower().endswith(".png"): mime = "image/png"
                image_parts.append({
                    "inline_data": {
                        "mime_type": mime,
                        "data": b64
                    }
                })

    # 3. Gemini Prompt
    prompt = f"""Siz O'zbekiston ta'lim shartnomalari va rasmiy hujjatlarini (Pasport, ID-karta, Diplom, Shahodatnoma) tahlil qiluvchi mutaxassis Sun'iy Intellektsiz.
Quyida shartnoma matni va unga biriktirilgan hujjat rasmlari (Pasport/ID-karta va Diplom/Shahodatnoma) berilgan.

Fayl nomi: {fname}
Shartnoma matni:
{full_text}

Iltimos, matn va rasmlarni sinchiklab ko'rib chiqib, quyidagi JSON formatida aniq ma'lumotlarni qaytaring (faqat toza JSON, hech qanday qo'shimcha so'zsiz):
{{
  "ism": "Talabaning Familiyasi va Ismi (masalan: Olimjonova Zuxra yoki Raxmonova Durdona)",
  "ota": "Otasining ismi (masalan: Zarif qizi yoki Shokir qizi)",
  "fish": "To'liq F.I.SH (masalan: Olimjonova Zuxra Zarif qizi)",
  "yonalis": "Ta'lim yo'nalishi (masalan: Hamshiralik ishi)",
  "tolov": "To'lov qildi",
  "shnum": "Shartnoma raqami (faqat raqami, masalan: 201)",
  "shsan": "Shartnoma sanasi (masalan: 19.08.2026)",
  "maktab": "Tugatgan muassasasi to'liq nomi (masalan: 4-sonli umumiy o'rta ta'lim maktabi yoki Shahrisabz Iqtisodiyot Kolleji)",
  "edutur": "Muassasa turi: 'Umumiy o\\'rta maktab' yoki 'Kollej' yoki 'Texnikum'",
  "yil": "Bitirgan yili (4 xonali yil, masalan: 2026 yoki 2014)",
  "tel": "Telefon raqamlari",
  "pass_tur": "'ID-karta' yoki 'Biometrik pasport'",
  "pass_ser": "Pasport seriyasi (2 ta bosh harf, masalan: AE yoki AD yoki AC)",
  "pass_num": "Pasport raqami (7 ta raqam, masalan: 4652546)",
  "pinfl": "14 xonali JSHSHIR (PINFL) raqami",
  "dob": "Tug'ilgan sanasi (kun.oy.yil formatda, masalan: 03.12.2008)",
  "passber": "Pasport berilgan sanasi",
  "cert_tur": "'Shahodatnoma' yoki 'Diplom'",
  "cert_ser": "Ta'lim hujjati seriyasi (masalan: UM yoki K yoki PT yoki IM yoki DB)",
  "cert_num": "Ta'lim hujjati raqami (masalan: 03728034 yoki 3359681)",
  "cert_full": "Ta'lim hujjati to'liq (masalan: UM № 03728034 yoki K № 3359681)"
}}
"""

    parts = [{"text": prompt}]
    parts.extend(image_parts[:4])

    payload = {
        "contents": [{"parts": parts}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json"
        }
    }

    for attempt in range(1, max_retries + 1):
        try:
            req = urllib.request.Request(
                API_URL,
                data=json.dumps(payload).encode('utf-8'),
                headers={'Content-Type': 'application/json'}
            )
            with urllib.request.urlopen(req, timeout=45) as resp:
                data = json.loads(resp.read().decode('utf-8'))
                candidates = data.get("candidates", [])
                if candidates:
                    res_text = candidates[0]["content"]["parts"][0]["text"]
                    res_json = json.loads(res_text)
                    res_json["docx"] = fname
                    res_json["docx_path"] = os.path.abspath(docx_path)
                    return res_json
        except Exception as e:
            print(f"[AI Urinish #{attempt}] {fname}: {e}")
            if attempt < max_retries:
                time.sleep(2)
    return None
