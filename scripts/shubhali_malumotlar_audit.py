# -*- coding: utf-8 -*-
"""
SHUBHALI MA'LUMOTLAR TEKSHIRUVI
===============================
talabalar_bazasi.json dagi barcha talabalarni tekshirib, shubhali / xato
ma'lumotlarni SHUBHALI_MALUMOTLAR_HISOBOTI.xlsx fayliga chiqaradi.

Tekshiruvlar:
  * JSHSHIR (PINFL): formati, nazorat raqami, tug'ilgan sana va jins mosligi
  * Pasport: formati, turi (ID-karta / biometrik) va seriya mosligi, berilgan sana
  * Takrorlar: bir xil PINFL / pasport / shahodatnoma raqami ikki talabada
  * Ta'lim hujjati: seriya va hujjat turi mosligi, bitirgan yoshi
  * F.I.O: otasining ismi, pasport / shahodatnomadagi ism bilan farq
  * Telefon: raqam uzunligi va operator kodi
  * Saqlashdagi xatolar: "+" belgilari, ortiqcha bo'shliqlar

Boshqa skriptlar normalize_student() ni import qilib ishlatadi
(generate_qabul_shablon.py).
"""
import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from datetime import date, datetime

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_telefonlar import docx_phones  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_JSON = os.path.join(ROOT, 'data', 'talabalar_bazasi.json')
OUT_XLSX = os.path.join(ROOT, 'hisobotlar', 'tekshiruvlar', 'SHUBHALI_MALUMOTLAR_HISOBOTI.xlsx')
# Foydalanuvchi tasdiqlagan istisnolar (masalan, eksternat) — JSHSHIR bo'yicha
# (portal eksporti ham shu faylni o'qiydi)
EXCEPTIONS_JSON = os.path.join(ROOT, 'web', 'lib', 'qabul-istisnolar.json')

# 1-kurs rasmiy guruhlari
GROUPS = ['26-01', '26-02', '26-03', '26-04', '26-05', '26-06', '26-07']
TODAY = date.today()

JIDDIY, ORTA, PAST = 'Jiddiy', "O'rta", 'Past'
SEVERITY_ORDER = {JIDDIY: 0, ORTA: 1, PAST: 2}

W = [7, 3, 1, 7, 3, 1, 7, 3, 1, 7, 3, 1, 7]
# O'zbekiston mobil va Qashqadaryo atrofi statsionar kodlari
# 87 — shartnomalarda 20+ oilada uchraydi, haqiqiy kod deb qabul qilingan
PHONE_CODES = {'20', '33', '50', '55', '61', '62', '65', '66', '67', '69', '70', '71', '72', '73',
               '74', '75', '76', '77', '78', '79', '87', '88', '90', '91', '93', '94', '95', '97', '98', '99'}
SCHOOL_SERIES = ('UM', "O'R-SH", "O'R-A", 'DB', 'UM-O', 'A', 'AT')
COLLEGE_SERIES = ('K', 'PT', 'D', 'L', 'KX', 'TT')
APOS = re.compile(r"[‘’ʻʼ`´]")


# ---------------------------------------------------------------- normalize
def clean_text(v) -> str:
    s = str(v or '')
    s = s.replace('+', ' ')
    s = APOS.sub("'", s)
    return re.sub(r'\s+', ' ', s).strip()


def norm_doc(v) -> str:
    s = clean_text(v).upper().replace('№', ' ')
    s = re.sub(r'\s+', ' ', s).strip()
    m = re.fullmatch(r"([A-Z'\-]+?)\s*(\d{5,9})", s)
    return f"{m.group(1)} {m.group(2)}" if m else s


def split_phones(v) -> list[str]:
    out = []
    for part in re.split(r'[;,/]+', str(v or '')):
        d = re.sub(r'\D', '', part)
        if len(d) == 12 and d.startswith('998'):
            d = d[3:]
        if d:
            out.append(d)
    return out


def normalize_student(s: dict) -> dict:
    """Shablon uchun tozalangan nusxa: '+' belgilari, bo'shliqlar, formatlar."""
    n = dict(s)
    for k in ('ism', 'ota', 'fish', 'mak', 'yil', 'pass_fish', 'cert_fish', 'dob', 'ber'):
        n[k] = clean_text(s.get(k))
    n['pinfl'] = re.sub(r'\D', '', str(s.get('pinfl') or ''))
    n['pv'] = re.sub(r'\s+', '', clean_text(s.get('pv'))).upper()
    n['sh_doc'] = norm_doc(s.get('sh_doc'))
    n['phones'] = split_phones(s.get('tel'))
    return n


def fio_key(v) -> str:
    return clean_text(v).lower().replace("'", '')


# ------------------------------------------------------------------- PINFL
def pinfl_ok(p: str) -> bool:
    return sum(int(p[i]) * W[i] for i in range(13)) % 10 == int(p[13])


def pinfl_dob(p: str):
    cent = {'1': 1800, '2': 1800, '3': 1900, '4': 1900, '5': 2000, '6': 2000}.get(p[0])
    if not cent:
        return None
    try:
        return date(cent + int(p[5:7]), int(p[3:5]), int(p[1:3]))
    except ValueError:
        return None


def pinfl_fix_candidates(p: str) -> list[str]:
    """Bitta raqam o'zgartirilganda to'g'ri va real yoshni beradigan variantlar."""
    out = []
    for i in range(14):
        for d in '0123456789':
            if d == p[i]:
                continue
            c = p[:i] + d + p[i + 1:]
            if pinfl_ok(c):
                dob = pinfl_dob(c)
                if dob and 15 <= (TODAY.year - dob.year) <= 60:
                    out.append(c)
    return out


def parse_date(v):
    try:
        return datetime.strptime(clean_text(v), '%d.%m.%Y').date()
    except ValueError:
        return None


def is_female_name(ota: str, fam: str):
    o = ota.lower()
    if 'qizi' in o or o.endswith(('ovna', 'evna', 'yevna')):
        return True
    if "o'g'li" in o or 'ogli' in o or o.endswith(('ovich', 'evich', 'yevich')):
        return False
    f = fam.lower()
    if f.endswith(('ova', 'eva', 'yeva')):
        return True
    if f.endswith(('ov', 'ev', 'yev')):
        return False
    return None


# ------------------------------------------------------------------ checks
def audit(students: list[dict]) -> list[dict]:
    from generate_qabul_shablon import qabul_rows
    qrows = {id(s): r for s, r in zip(students, qabul_rows(students))}
    issues = []

    def add(s, sev, field, value, problem, advice=''):
        issues.append({'sev': sev, 'group': s['group'], 'fio': clean_text(s.get('fish')),
                       'field': field, 'value': str(value or ''), 'problem': problem,
                       'advice': advice, 'shnum': s.get('shnum', ''), 'pinfl': s['_n']['pinfl']})

    for s in students:
        n = s['_n']
        raw_pinfl = str(s.get('pinfl') or '')
        dob = parse_date(s.get('dob'))
        fam = n['fish'].split(' ')[0] if n['fish'] else ''

        # --- saqlash xatolari ("+" va bo'shliqlar)
        plus_fields = [k for k in ('ism', 'ota', 'mak', 'sh_doc', 'pv', 'pinfl') if '+' in str(s.get(k) or '')]
        if plus_fields:
            add(s, ORTA, ', '.join(plus_fields), '; '.join(str(s.get(k)) for k in plus_fields),
                "Ma'lumot bazaga '+' belgilari bilan saqlangan (bo'shliq o'rniga)",
                "Portalda qayta saqlang; shablonda bo'shliq bilan tozalab yozildi")

        # --- F.I.O
        ota = n['ota']
        if not ota:
            add(s, ORTA, 'Otasining ismi', ota, "Otasining ismi kiritilmagan", 'Pasportdan to\'ldiring')
        elif not re.search(r"(qizi|o'g'li|ogli|ovna|evna|ovich|evich)$", ota.lower()):
            add(s, PAST, 'Otasining ismi', ota, "Otasining ismi 'qizi / o'g'li' bilan tugamagan", 'Pasport bilan solishtiring')
        if n['ota'] and 'muroulla' in n['ota'].lower():
            add(s, ORTA, 'Otasining ismi', ota, "Ism g'ayrioddiy yozilgan (ehtimol 'Murodulla')", 'Pasport bilan solishtiring')
        for key, label in (('pass_fish', 'Pasportdagi F.I.O'), ('cert_fish', 'Shahodatnomadagi F.I.O')):
            other = n[key]
            if other and other not in ('Mavjud',) and fio_key(other) != fio_key(n['fish']):
                add(s, ORTA, label, other, f"{label} bazadagi F.I.O ({n['fish']}) bilan mos emas",
                    'Hujjat asl nusxasi bilan tekshiring')
        if 'BOSHQA ODAM' in str(s.get('name_match') or '').upper():
            add(s, ORTA, 'Pasport', n['pv'], "Avvalgi tekshiruvda pasport 'BOSHQA ODAM' deb belgilangan",
                'Pasport nusxasi shu talabaga tegishliligini tasdiqlang')

        # --- PINFL
        p = n['pinfl']
        if not p:
            add(s, JIDDIY, 'JSHSHIR', raw_pinfl, 'JSHSHIR kiritilmagan', 'Pasportdan kiriting')
        elif len(p) != 14:
            add(s, JIDDIY, 'JSHSHIR', raw_pinfl, f'JSHSHIR {len(p)} xonali (14 bo\'lishi kerak)', 'Pasportdan qayta tekshiring')
        else:
            if raw_pinfl != p:
                add(s, PAST, 'JSHSHIR', raw_pinfl, "JSHSHIR ichida bo'shliq / ortiqcha belgi bor",
                    f"To'g'ri yozilishi: {p} (shablonda tozalandi)")
            if not pinfl_ok(p):
                cands = pinfl_fix_candidates(p)
                adv = ('Ehtimoliy to\'g\'ri variant(lar): ' + ', '.join(cands[:4])) if cands else 'Pasportdan qayta kiriting'
                add(s, JIDDIY, 'JSHSHIR', p, "JSHSHIR nazorat raqami (14-raqam) mos emas — raqamlardan biri xato terilgan", adv)
            pdob = pinfl_dob(p)
            if pdob and dob and pdob != dob:
                add(s, JIDDIY, "Tug'ilgan sana", s.get('dob'),
                    f"Tug'ilgan sana JSHSHIRdagi sanaga ({pdob:%d.%m.%Y}) mos emas", 'Pasport bilan solishtiring')
            fem = is_female_name(ota, fam)
            if fem is not None and p[0] in '3456' and fem != (p[0] in '246'):
                add(s, JIDDIY, 'JSHSHIR', p,
                    f"Jins mos emas: F.I.O bo'yicha {'ayol' if fem else 'erkak'}, JSHSHIR 1-raqami ({p[0]}) bo'yicha {'ayol' if p[0] in '246' else 'erkak'}",
                    'JSHSHIR yoki F.I.O xato — pasportni tekshiring')

        # --- yosh (real bo'lmasa, sanaga bog'liq keyingi tekshiruvlar o'tkazib yuboriladi)
        if dob:
            age = (TODAY - dob).days // 365
            if age < 15 or age > 55:
                add(s, JIDDIY, "Tug'ilgan sana", s.get('dob'), f"Yoshi {age} — real emas", 'Tug\'ilgan yil / JSHSHIR xato terilgan')
                dob = None
        elif s.get('dob'):
            add(s, ORTA, "Tug'ilgan sana", s.get('dob'), "Tug'ilgan sana formati noto'g'ri", 'dd.mm.yyyy')

        # --- pasport
        pv = n['pv']
        if not pv:
            add(s, JIDDIY, 'Pasport', '', 'Pasport seriya raqami kiritilmagan', 'Pasport nusxasini oling')
        else:
            if not re.fullmatch(r'[A-Z]{2}\d{7}', pv):
                add(s, JIDDIY, 'Pasport', s.get('pv'), "Pasport formati noto'g'ri (2 harf + 7 raqam bo'lishi kerak)", 'Pasportdan qayta kiriting')
            elif str(s.get('pv') or '') != pv:
                add(s, PAST, 'Pasport', s.get('pv'), "Pasport raqamida bo'shliq / kichik harf bor", f"To'g'ri yozilishi: {pv} (shablonda tozalandi)")
            pt = s.get('pass_type') or ''
            if pt == 'ID-karta' and pv[:2] in ('AA', 'AB', 'AC', 'FA'):
                add(s, ORTA, 'Pasport turi', f'{pt} / {pv}', 'Seriya biometrik pasportga xos, turi esa ID-karta', 'Turini tekshiring')
            if pt == 'Biometrik Pasport' and pv[:2] in ('AD', 'AE'):
                add(s, ORTA, 'Pasport turi', f'{pt} / {pv}', 'Seriya ID-kartaga xos, turi esa biometrik pasport', 'Turini tekshiring')
        ber = parse_date(s.get('ber'))
        ber_bad = False
        if ber and dob:
            if ber <= dob:
                ber_bad = True
                add(s, JIDDIY, 'Pasport berilgan sana', s.get('ber'),
                    f"Pasport berilgan sana tug'ilgan sanaga ({s.get('dob')}) teng yoki undan oldin", 'Pasportdan to\'g\'ri sanani kiriting')
            elif (ber - dob).days / 365.25 < 16 and (s.get('pass_type') == 'ID-karta'):
                add(s, ORTA, 'Pasport berilgan sana', s.get('ber'), "ID-karta 16 yoshdan oldin berilgan bo'lib chiqyapti", 'Sanalarni tekshiring')
        if ber and ber > TODAY:
            add(s, JIDDIY, 'Pasport berilgan sana', s.get('ber'), 'Pasport berilgan sana kelajakda', 'Sanani tekshiring')
        if ber and not ber_bad and (TODAY - ber).days > 3652:
            add(s, ORTA, 'Pasport berilgan sana', s.get('ber'), "Pasport muddati o'tgan (10 yildan ortiq)", 'Yangi pasport nusxasini oling')

        # --- ta'lim hujjati
        doc, tur, mak, yil = n['sh_doc'], s.get('doc_tur') or '', n['mak'], n['yil']
        if not doc:
            add(s, ORTA, 'Diplom / shahodatnoma', '', 'Ta\'lim hujjati seriya-raqami kiritilmagan', 'Hujjat nusxasidan kiriting')
        else:
            series = doc.split(' ')[0]
            if tur in ('Shahodatnoma', 'Attestat') and series in COLLEGE_SERIES:
                add(s, ORTA, 'Hujjat turi', f'{tur} / {doc}', 'Seriya kollej diplomiga xos, turi esa shahodatnoma', 'Hujjat turini tekshiring')
            if tur == 'Diplom' and series in ('UM', "O'R-SH", 'UM-O'):
                add(s, ORTA, 'Hujjat turi', f'{tur} / {doc}', 'Seriya maktab shahodatnomasiga xos, turi esa diplom', 'Hujjat turini tekshiring')
        if tur not in ('Shahodatnoma', 'Diplom', 'Attestat'):
            add(s, ORTA, 'Hujjat turi', tur, f"Ta'lim hujjati turi noto'g'ri kiritilgan ('{tur}')", 'Shahodatnoma yoki Diplom deb belgilang')
        if not mak or mak.lower() in ("noma'lum", 'nomalum'):
            add(s, ORTA, "O'qigan muassasa", mak, "Avval o'qigan muassasa nomi kiritilmagan", 'Hujjatdan kiriting')
        elif tur == 'Shahodatnoma' and re.search(r'kollej|texnikum', mak.lower()):
            add(s, ORTA, "O'qigan muassasa", mak, 'Muassasa kollej, lekin hujjat turi shahodatnoma', 'Hujjat turini tekshiring')
        if not re.fullmatch(r'\d{4}', yil):
            add(s, ORTA, 'Tugatgan yili', yil, 'Tugatgan yili kiritilmagan', 'Hujjatdan kiriting')
        elif dob:
            grad_age = int(yil) - dob.year
            if int(yil) > TODAY.year:
                add(s, JIDDIY, 'Tugatgan yili', yil, 'Tugatgan yili kelajakda', 'Yilni tekshiring')
            elif (tur in ('Shahodatnoma', 'Attestat') and not (15 <= grad_age <= 20)
                  and 'bitirish_yoshi' not in s['_ist'].get('otkazish', [])):
                add(s, ORTA, 'Tugatgan yili', yil,
                    f"Maktabni {grad_age} yoshda tugatgan bo'lib chiqyapti (odatda 17-18)", 'Yil yoki tug\'ilgan sanani tekshiring')
            elif tur == 'Diplom' and grad_age < 17:
                add(s, ORTA, 'Tugatgan yili', yil, f'Kollejni {grad_age} yoshda tugatgan bo\'lib chiqyapti', 'Yilni tekshiring')

        # --- telefon
        for ph in n['phones']:
            if len(ph) != 9:
                add(s, ORTA, 'Telefon', s.get('tel'), f"Telefon raqami {len(ph)} xonali ({ph}) — 9 bo'lishi kerak", 'Raqamni tekshiring')
            elif ph[:2] not in PHONE_CODES:
                add(s, ORTA, 'Telefon', s.get('tel'), f"Telefon {ph}: '{ph[:2]}' operator kodi mavjud emas", 'Raqamni tekshiring')

        # --- qabul shabloni: yashash hududi (G ustuni) va muassasa tarjimasi
        qr = qrows[id(s)]
        if p and not qr['viloyat']:
            add(s, PAST, 'Yashash viloyat+tumani', f'PINFL hudud kodi {p[7:10]}' if len(p) == 14 else '',
                "Yashash hududi aniqlanmadi (muassasa nomida tuman yo'q, PINFL hudud kodi noma'lum)",
                "Shablonda bo'sh qoldirildi — qo'lda kiriting")
        if 'makEn' in qr['review']:
            add(s, ORTA, "Muassasa tarjimasi", qr['makUz'],
                f"Muassasa nomi uchun tarjima qoidasi topilmadi — umumiy tarjima berildi: {qr['makEn']} / {qr['makRu']}",
                "web/lib/qabul.ts dagi FIELDS / OVERRIDES ga qo'shing")

        # --- holat
        if s.get('status') in ('chala', 'yoq'):
            add(s, PAST, 'Holat', s.get('status'), "Hujjatlar to'liq emas (status: " + s.get('status') + ')', '')
        if s.get('verified') == 'KUTILMOQDA':
            add(s, PAST, 'Tasdiq', 'KUTILMOQDA', 'Ma\'lumotlar hali hujjat bilan tasdiqlanmagan', 'Hujjat nusxasi bilan solishtirib tasdiqlang')

    # --- takrorlar
    for key, label in (('pinfl', 'JSHSHIR'), ('pv', 'Pasport'), ('sh_doc', 'Diplom / shahodatnoma')):
        seen = defaultdict(list)
        for s in students:
            v = s['_n'][key]
            if v:
                seen[v].append(s)
        for v, lst in seen.items():
            if len(lst) > 1:
                names = ', '.join(clean_text(x['fish']) for x in lst)
                for s in lst:
                    add(s, JIDDIY, label, v, f'{label} bir nechta talabada takrorlangan: {names}', 'Qaysi biriga tegishli ekanini aniqlang')
    phones = defaultdict(list)
    for s in students:
        for ph in set(s['_n']['phones']):
            phones[ph].append(s)
    for ph, lst in phones.items():
        if len(lst) > 1:
            names = ', '.join(clean_text(x['fish']) for x in lst)
            for s in lst:
                add(s, PAST, 'Telefon', ph, f'Bir xil telefon raqami: {names}', 'Aka-uka/opa-singil bo\'lishi mumkin — tekshiring')

    issues.sort(key=lambda i: (SEVERITY_ORDER[i['sev']], i['group'], i['fio']))
    return issues


# ------------------------------------------------------------- today diff
def load_baseline():
    """Bugungi birinchi o'zgarishdan oldingi talabalar_bazasi.json (git)."""
    try:
        rev = subprocess.run(['git', 'rev-list', '-1', f'--before={TODAY:%Y-%m-%d} 00:00', 'HEAD', '--', 'data/talabalar_bazasi.json', 'talabalar_bazasi.json'],
                             cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
        if not rev:
            return None, None
        # 26.09.2026 dan keyin data/ da, undan oldin ildizda edi
        for rel in ('data/talabalar_bazasi.json', 'talabalar_bazasi.json'):
            r = subprocess.run(['git', 'show', f'{rev}:{rel}'], cwd=ROOT, capture_output=True)
            if r.returncode == 0:
                return rev[:7], json.loads(r.stdout.decode('utf-8'))['students']
        return None, None
    except Exception as e:
        print('Baseline olinmadi:', e)
        return None, None


DIFF_FIELDS = [('fish', 'F.I.O'), ('group', 'Guruh'), ('pinfl', 'JSHSHIR'), ('pv', 'Pasport'), ('dob', "Tug'ilgan sana"),
               ('ber', 'Pasport berilgan'), ('sh_doc', 'Diplom/shahodatnoma'), ('mak', 'Muassasa'), ('yil', 'Tugatgan yili'),
               ('doc_tur', 'Hujjat turi'), ('tel', 'Telefon'), ('status', 'Holat'), ('verified', 'Tasdiq')]


def today_changes(current, baseline):
    if baseline is None:
        return []
    def k(s):
        return str(s.get('shnum') or '') + '|' + fio_key(s.get('ism'))
    old = {k(s): s for s in baseline}
    old_by_name = {fio_key(s.get('ism')): s for s in baseline}
    rows = []
    for s in current:
        o = old.get(k(s)) or old_by_name.get(fio_key(s.get('ism')))
        if o is None:
            rows.append((s, 'Yangi talaba', '', '', clean_text(s.get('fish'))))
            continue
        for f, label in DIFF_FIELDS:
            if str(o.get(f) or '') != str(s.get(f) or ''):
                rows.append((s, label, f, str(o.get(f) or ''), str(s.get(f) or '')))
    return rows


# ------------------------------------------------------------------ excel
FONT = 'Arial'
THIN = Side(style='thin', color='B7B7B7')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill('solid', start_color='1F3864')
SEV_FILL = {JIDDIY: PatternFill('solid', start_color='F8CBAD'),
            ORTA: PatternFill('solid', start_color='FFE699'),
            PAST: PatternFill('solid', start_color='DDEBF7')}


def write_table(ws, headers, rows, widths, start_row=1, sev_col=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row, c, h)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[start_row].height = 30
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if sev_col and c == sev_col:
                cell.fill = SEV_FILL.get(v, PatternFill())
                cell.font = Font(name=FONT, size=10, bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='top')
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def build_report(active, expelled, issues, changes, base_rev):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Xulosa'

    changed_keys = {(c[0]['group'], clean_text(c[0].get('fish'))) for c in changes}
    fio_issue_count = defaultdict(int)
    for i in issues:
        fio_issue_count[(i['group'], i['fio'])] += 1

    # --- 2. Shubhali ma'lumotlar
    ws_i = wb.create_sheet("Shubhali ma'lumotlar")
    rows = []
    for n, i in enumerate(issues, 1):
        rows.append([n, i['sev'], i['group'], i['fio'], i['shnum'], i['field'], i['value'], i['problem'], i['advice'],
                     'Ha' if (i['group'], i['fio']) in changed_keys else ''])
    write_table(ws_i, ['№', 'Daraja', 'Guruh', 'F.I.O', 'Shartnoma №', 'Maydon', 'Kiritilgan qiymat', 'Muammo',
                       'Tavsiya', 'Bugun o\'zgartirilgan'],
                rows, [5, 9, 9, 34, 10, 18, 26, 52, 40, 11], sev_col=2)
    n_issues = len(rows)

    # --- 3. Talabalar kesimida
    ws_s = wb.create_sheet('Talabalar kesimida')
    per = defaultdict(list)
    for i in issues:
        per[(i['group'], i['fio'])].append(i)
    srows = []
    for (g, fio), lst in sorted(per.items(), key=lambda kv: (min(SEVERITY_ORDER[x['sev']] for x in kv[1]), kv[0])):
        worst = min((x['sev'] for x in lst), key=lambda v: SEVERITY_ORDER[v])
        srows.append([0, worst, g, fio, len(lst),
                      '\n'.join(f"• [{x['sev']}] {x['problem']}" for x in lst),
                      'Ha' if (g, fio) in changed_keys else ''])
    for n, r in enumerate(srows, 1):
        r[0] = n
    write_table(ws_s, ['№', 'Eng yuqori daraja', 'Guruh', 'F.I.O', 'Muammolar soni', 'Muammolar', 'Bugun o\'zgartirilgan'],
                srows, [5, 11, 9, 34, 10, 90, 11], sev_col=2)

    # --- 4. Bugungi o'zgarishlar
    ws_c = wb.create_sheet("Bugungi o'zgarishlar")
    crows = []
    for n, (s, label, f, old, new) in enumerate(changes, 1):
        crows.append([n, s['group'], clean_text(s.get('fish')), label, old, new,
                      "Ha" if fio_issue_count.get((s['group'], clean_text(s.get('fish')))) else ''])
    write_table(ws_c, ['№', 'Guruh', 'F.I.O', 'Maydon', 'Oldingi qiymat', 'Yangi qiymat', 'Shubha bor'],
                crows, [5, 9, 34, 20, 34, 40, 10])

    # --- 5. Chiqarilganlar (shablonga kirmaydi)
    ws_e = wb.create_sheet('Chiqarilganlar')
    write_table(ws_e, ['№', 'F.I.O', 'Guruh / holat', 'JSHSHIR'],
                [[n, clean_text(s.get('fish')), s['group'], s['_n']['pinfl']] for n, s in enumerate(expelled, 1)],
                [5, 36, 34, 18])

    # --- 7. Telefonlar manbasi (shablonga qaysi raqam qayerdan olingani)
    ws_t = wb.create_sheet('Telefonlar manbasi')
    trows = []
    for s in sorted(active, key=lambda x: (x['group'], clean_text(x.get('fish')))):
        n = s['_n']
        src = n.get('tel_manba', '')
        trows.append([0, s['group'], clean_text(s.get('fish')), ', '.join(n['phones']),
                      'Bazadan' if src == 'baza' else (f'Shartnomadan: {src}' if src else 'Topilmadi')])
    for k, r in enumerate(trows, 1):
        r[0] = k
    write_table(ws_t, ['№', 'Guruh', 'F.I.O', 'Telefon(lar)', 'Manba'], trows, [5, 9, 36, 24, 52])

    # --- 6. Tasdiqlangan istisnolar (tekshiruvdan chiqarilgan holatlar)
    ws_x = wb.create_sheet('Tasdiqlangan istisnolar')
    xrows = [[n, s['group'], clean_text(s.get('fish')), s['_n']['pinfl'], s['_ist'].get('sabab', ''),
              ', '.join(s['_ist'].get('otkazish', []))]
             for n, s in enumerate([s for s in active if s['_ist']], 1)]
    write_table(ws_x, ['№', 'Guruh', 'F.I.O', 'JSHSHIR', 'Sabab', "O'tkazilgan tekshiruv"],
                xrows, [5, 9, 34, 18, 50, 20])

    # --- 1. Xulosa (formulalar "Shubhali ma'lumotlar" varag'iga bog'langan)
    sh = "'Shubhali ma''lumotlar'"
    rng_sev = f"{sh}!$B$2:$B${max(n_issues + 1, 2)}"
    rng_grp = f"{sh}!$C$2:$C${max(n_issues + 1, 2)}"
    ws['A1'] = "Shubhali ma'lumotlar tekshiruvi — hisobot"
    ws['A1'].font = Font(name=FONT, size=14, bold=True, color='1F3864')
    ws['A2'] = f"Manba: talabalar_bazasi.json · Tekshiruv sanasi: {TODAY:%d.%m.%Y} · Rasmiy guruhlar: {', '.join(GROUPS)}"
    ws['A2'].font = Font(name=FONT, size=9, italic=True, color='595959')

    ws['A4'], ws['B4'] = "Tekshirilgan talabalar", len(active)
    ws['A5'], ws['B5'] = "Shubhali ma'lumoti bor talabalar", f"=COUNTA('Talabalar kesimida'!$D$2:$D${max(len(srows) + 1, 2)})"
    ws['A6'], ws['B6'] = "Jami topilgan muammolar", f"=COUNTA({sh}!$A$2:$A${max(n_issues + 1, 2)})"
    ws['A7'], ws['B7'] = "Bugun o'zgartirilgan maydonlar", f"=COUNTA('Bugungi o''zgarishlar'!$A$2:$A${max(len(crows) + 1, 2)})"
    for r in range(4, 8):
        ws.cell(r, 1).font = Font(name=FONT, size=11)
        ws.cell(r, 2).font = Font(name=FONT, size=11, bold=True)
        ws.cell(r, 2).alignment = Alignment(horizontal='center')

    ws['A9'] = 'Daraja'
    for c, g in enumerate(GROUPS, 2):
        ws.cell(9, c, g)
    ws.cell(9, len(GROUPS) + 2, 'Jami')
    for c in range(1, len(GROUPS) + 3):
        cell = ws.cell(9, c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal='center')
        cell.border = BORDER
    for r, sev in enumerate((JIDDIY, ORTA, PAST), 10):
        ws.cell(r, 1, sev).fill = SEV_FILL[sev]
        for c, g in enumerate(GROUPS, 2):
            # SUMPRODUCT: "26-01" kabi guruh nomlarini COUNTIFS sana deb talqin qilmasligi uchun
            ws.cell(r, c, f'=SUMPRODUCT(({rng_sev}=$A{r})*({rng_grp}={get_column_letter(c)}$9))')
        last = get_column_letter(len(GROUPS) + 1)
        ws.cell(r, len(GROUPS) + 2, f'=SUM(B{r}:{last}{r})')
        for c in range(1, len(GROUPS) + 3):
            ws.cell(r, c).font = Font(name=FONT, bold=(c == 1 or c == len(GROUPS) + 2))
            ws.cell(r, c).border = BORDER
            ws.cell(r, c).alignment = Alignment(horizontal='center')

    ws['A14'] = 'Darajalar izohi'
    ws['A14'].font = Font(name=FONT, bold=True)
    legend = [
        (JIDDIY, "Ma'lumot aniq xato yoki ziddiyatli (JSHSHIR nazorat raqami, sana/jins mos emas, takror, real bo'lmagan yosh). Qabulga yuborishdan oldin albatta tuzatish kerak."),
        (ORTA, "Ma'lumot to'liq emas yoki shubhali (hujjat turi, muassasa, telefon kodi, '+' belgilari, ismlar farqi). Hujjat bilan solishtirish kerak."),
        (PAST, "Format kamchiligi yoki eslatma (bo'shliqlar, tasdiqlanmagan, bir xil telefon). Shablonda avtomatik tozalandi."),
    ]
    for r, (sev, txt) in enumerate(legend, 15):
        ws.cell(r, 1, sev).fill = SEV_FILL[sev]
        ws.cell(r, 1).font = Font(name=FONT, bold=True)
        ws.cell(r, 2, txt).font = Font(name=FONT, size=10)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 30
    ws['A19'] = ("Izoh: 'Bugun o'zgartirilgan' — bugun kiritilgan/o'zgartirilgan ma'lumotlar, "
                 f"git'dagi {base_rev or '—'} holatiga nisbatan. Telefon raqami bazada faqat ayrim talabalarda bor — "
                 "bo'sh telefonlar muammo sifatida sanalmagan.")
    ws['A19'].font = Font(name=FONT, size=9, italic=True, color='595959')
    ws.column_dimensions['A'].width = 34
    for c in range(2, len(GROUPS) + 3):
        ws.column_dimensions[get_column_letter(c)].width = 11

    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT_XLSX)
    return len(srows)


def load_exceptions() -> dict:
    if not os.path.exists(EXCEPTIONS_JSON):
        return {}
    with open(EXCEPTIONS_JSON, encoding='utf-8') as f:
        return {k: v for k, v in json.load(f).items() if not k.startswith('_')}


def load_students():
    with open(DB_JSON, encoding='utf-8') as f:
        students = json.load(f)['students']
    exceptions = load_exceptions()
    for s in students:
        s['_n'] = normalize_student(s)
        s['_ist'] = exceptions.get(s['_n']['pinfl'], {})
        # Bazada telefon bo'lmasa — shartnoma .docx dagi "TELEFON RAQAMI" bo'limidan
        s['_n']['tel_manba'] = 'baza' if s['_n']['phones'] else ''
        if not s['_n']['phones']:
            phones, src = docx_phones(s)
            if phones:
                s['_n']['phones'], s['_n']['tel_manba'] = phones, src
    return students


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    students = load_students()
    active = [s for s in students if s.get('group') in GROUPS]
    expelled = [s for s in students if s.get('group') not in GROUPS]
    issues = audit(active)
    base_rev, baseline = load_baseline()
    changes = today_changes(active, baseline)
    n_students = build_report(active, expelled, issues, changes, base_rev)

    by = defaultdict(int)
    for i in issues:
        by[i['sev']] += 1
    print(f"Tekshirildi: {len(active)} talaba (chiqarilgan: {len(expelled)})")
    print(f"Shubhali talabalar: {n_students} | muammolar: {len(issues)} "
          f"(Jiddiy {by[JIDDIY]}, O'rta {by[ORTA]}, Past {by[PAST]})")
    print(f"Bugungi o'zgarishlar: {len(changes)} ta maydon (asos: {base_rev})")
    print(f"Saqlandi: {OUT_XLSX}")
    return issues


if __name__ == '__main__':
    main()
