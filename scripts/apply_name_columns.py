# -*- coding: utf-8 -*-
"""
PASPORT VA SHAHODATNOMADAGI F.I.SH NI ALOHIDA USTUNLARGA YOZISH
================================================================
Foydalanuvchining o'z ro'yxati (2-ustun "I.F.O") O'ZGARTIRILMAYDI.
Rasmiy hujjatlardagi yozuv alohida ustunlarga qo'yiladi, shunda ikkalasini
yonma-yon ko'rib, qaysi biri to'g'ri ekanini o'zingiz hal qilasiz.

  9-ustun  : Passport bo'yicha F.I.SH      (ID-karta MRZ dan)
  14-ustun : Shahodatnoma bo'yicha F.I.SH  (e-shahodatnoma.uz PDF dan)
  24-ustun : Ism mosligi                   (MOS / TRANSLIT / FARQ / BOSHQA ODAM)

    python scripts/apply_name_columns.py --dry-run
    python scripts/apply_name_columns.py
"""
import openpyxl, os, sys, json, shutil, datetime, re

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, 'Talabalar_Toliq_Royxati.xlsx')
NAME_JSON = os.path.join(BASE_DIR, 'scratch', 'name_check.json')

DRY = '--dry-run' in sys.argv

COL_PASS_FISH = 9
COL_CERT_FISH = 14
COL_MATCH = 24

# Ichki kod -> Excelda ko'rinadigan matn
LABEL = {
    'mos': 'MOS',
    'translit': 'TRANSLIT',          # q/k, x/h — imlo konvensiyasi farqi
    'apostrof': "TUTUQ BELGISI",     # ' tushib qolgan
    'boshqa': 'FARQ',                # haqiqiy boshqacha yozilgan
    'boshqa_odam': 'BOSHQA ODAM',    # hujjat umuman boshqa kishiniki
    '': '',
}
RANK = {'': 0, 'mos': 1, 'translit': 2, 'apostrof': 3, 'boshqa': 4, 'boshqa_odam': 5}

if not os.path.exists(NAME_JSON):
    print('❌ Avval ishga tushiring: python scripts/verify_names_strict.py')
    sys.exit(1)

data = json.load(open(NAME_JSON, encoding='utf-8'))

# Qo'lda tasdiqlangan (qulflangan) qatorlar — ularga yozmaymiz
MANUAL_MAP = os.path.join(BASE_DIR, 'scripts', 'manual_file_map.json')
locked = {}
if os.path.exists(MANUAL_MAP):
    try:
        _cfg = json.load(open(MANUAL_MAP, encoding='utf-8'))
        locked = {(k.partition('|')[0].strip().lower(), k.partition('|')[2].strip()): v
                  for k, v in _cfg.get('qulflangan', {}).items()}
    except Exception as e:
        print(f"⚠️  manual_file_map.json o'qilmadi: {e}")

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

# Sarlavhalar
ws.cell(row=1, column=COL_PASS_FISH, value="Passport bo'yicha F.I.SH")
ws.cell(row=1, column=COL_CERT_FISH, value="Shahodatnoma bo'yicha F.I.SH")
ws.cell(row=1, column=COL_MATCH, value="Ism mosligi (Pasport / Shahodatnoma)")

n_pass = n_cert = n_mark = 0
stats = {}

n_lock = 0
for r in range(2, ws.max_row + 1):
    ifo = str(ws.cell(row=r, column=2).value or '').strip()
    if not ifo:
        continue
    guruh = str(ws.cell(row=r, column=23).value or '').strip()
    if (ifo.lower(), guruh) in locked:
        n_lock += 1
        continue
    d = data.get(str(r), {})
    pf = (d.get('pass_fish') or '').strip()
    cf = (d.get('cert_fish') or '').strip()
    ph = d.get('pass_holat') or ''
    ch = d.get('cert_holat') or ''

    if pf:
        ws.cell(row=r, column=COL_PASS_FISH, value=pf)
        n_pass += 1
    if cf:
        ws.cell(row=r, column=COL_CERT_FISH, value=cf)
        n_cert += 1

    # Umumiy holat — ikki manbadan eng jiddiyrog'i
    worst = max([ph, ch], key=lambda k: RANK.get(k, 0)) if (ph or ch) else ''
    if worst:
        parts = []
        if ph:
            parts.append(f"Pasport: {LABEL.get(ph, ph)}")
        if ch:
            parts.append(f"Shahodatnoma: {LABEL.get(ch, ch)}")
        ws.cell(row=r, column=COL_MATCH, value=' | '.join(parts))
        n_mark += 1
    else:
        ws.cell(row=r, column=COL_MATCH, value='TEKSHIRILMAGAN')
    stats[worst or 'tekshirilmagan'] = stats.get(worst or 'tekshirilmagan', 0) + 1

print('=' * 66)
print('  ALOHIDA USTUNLARGA YOZISH')
print('=' * 66)
print(f"  9-ustun  Passport bo'yicha F.I.SH     : {n_pass} ta")
print(f"  14-ustun Shahodatnoma bo'yicha F.I.SH : {n_cert} ta")
print(f"  24-ustun Ism mosligi                  : {n_mark} ta")
print(f"  🔒 Qo'lda tasdiqlangan (tegilmadi)    : {n_lock} ta")
print('-' * 66)
for k, v in sorted(stats.items(), key=lambda kv: -RANK.get(kv[0], 0)):
    print(f"  {LABEL.get(k, k) or k:16s}: {v:3d} ta")
print('=' * 66)

if DRY:
    print('\n🔍 DRY-RUN — hech narsa yozilmadi')
    sys.exit(0)

stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
bp = os.path.join(BASE_DIR, 'backup', f'Talabalar_Toliq_Royxati_ISM_USTUNLARIDAN_OLDIN_{stamp}.xlsx')
shutil.copy2(EXCEL_PATH, bp)
print(f"\n✅ Zaxira: backup/{os.path.basename(bp)}")

wb.save(EXCEL_PATH)
print(f"💾 Saqlandi: {os.path.basename(EXCEL_PATH)}")
print("\nDIQQAT: 2-ustun (sizning ro'yxatingiz) o'zgartirilmadi.")
