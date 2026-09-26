# -*- coding: utf-8 -*-
"""
6 TA MUAMMO BO'YICHA TO'LIQ TUZATISH SKRIPTI
============================================
Case 1 (T/R 32): Saidaliyeva -> Saydalimova Ruxshona Sharofiddin qizi (#74)
Case 2 (T/R 54): Mustafayeva Madina (26-03) -> Mustafoyeva 282 uzildi (boshqa shaxs)
Case 3 (T/R 91): Qodirova Sabina (26-04) -> Qodirova Shabnam 137 uzildi (boshqa shaxs)
Case 4 (T/R 152): Quchqarova Zebiniso Zebiniso qizi -> Otasining ismi Farhod qizi (#187)
Case 5 (T/R 165): Nazirova Intizor -> Narziyeva Intizor Iskandar qizi (#212)
Case 6 (T/R 95): Tuychiyeva Farzona -> Shahodatnoma F.I.SH "Mavjud" o'rniga haqiqiy F.I.SH (#280)
"""

import openpyxl
import json
import os
import shutil
import datetime
import sys

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TS = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
BACKUP_DIR = os.path.join(BASE_DIR, 'arxiv', 'backup')
os.makedirs(BACKUP_DIR, exist_ok=True)

# 1. ZAXIRA NUSXALARI
f_main = os.path.join(BASE_DIR, 'data', 'Talabalar_Toliq_Royxati.xlsx')
f_orig = os.path.join(BASE_DIR, 'data', 'manba', "to'liq 1-KURS 2026-2027.xlsx")
shutil.copy2(f_main, os.path.join(BACKUP_DIR, f'Talabalar_Toliq_Royxati_before_6cases_{TS}.xlsx'))
shutil.copy2(f_orig, os.path.join(BACKUP_DIR, f'to_liq_1-KURS_before_6cases_{TS}.xlsx'))
print(f"✅ Zaxira nusxalari yaratildi: {TS}")

# 2. TALABALAR_TOLIQ_ROYXATI.XLSX NI YANGILASH
wb = openpyxl.load_workbook(f_main)
ws = wb['Barcha Talabalar']

# Helper to find row by T/R in given worksheet
def find_row_by_tr(sheet, tr_val):
    for r in range(2, sheet.max_row + 1):
        if str(sheet.cell(r, 1).value or '').strip() == str(tr_val):
            return r
    return None

# Helper to update row in Talabalar_Toliq_Royxati
def apply_update_to_sheet(sheet, r, updates):
    for col_idx, val in updates.items():
        sheet.cell(row=r, column=col_idx).value = val

# CASE 1: T/R 32 (Row 33 in Barcha Talabalar, 26-02)
# Saydalimova Ruxshona Sharofiddin qizi (#74)
c1_updates = {
    2: 'Saydalimova Ruxshona',
    5: '74',
    7: 'Sharofiddin qizi',
    8: 'Saydalimova Ruxshona Sharofiddin qizi',
    9: 'SAYDALIMOVA RUXSHONA',
    10: 'AE3352294',
    11: '61807085590034',
    12: '10.07.2025',
    13: '18.07.2008',
    14: 'SAYDALIMOVA RUXSHONA SHAROFIDDIN QIZI',
    15: 'UM 03754198',
    16: 'https://e-shahodatnoma.uz/api/public/getcertpdf/4087114/3616453/43b2fbe0-3232-4e4c-a111-e6fb3ff34111?preview=true',
    17: "29-sonli umumiy o'rta ta'lim maktabi",
    18: 'Shahodatnoma',
    19: '2026',
    21: 'TOPILDI',
    22: "Familiyadagi xato tuzatildi: Saydalimova Ruxshona Sharofiddin qizi (ID-karta AE3352294, Shahodatnoma UM 03754198, shartnoma #74)",
    24: 'Pasport: MOS | Shahodatnoma: MOS',
    25: 'TASDIQLANDI'
}

# CASE 2: T/R 54 (Row 55 in Barcha Talabalar, 26-03)
# Mustafayeva Madina Maxmadiyevna -> uzildi
c2_updates = {
    9: None,
    10: None,
    11: None,
    12: None,
    13: None,
    14: None,
    15: None,
    16: None,
    17: None,
    18: None,
    19: None,
    21: 'FAYL_YOQ',
    22: "Fayl yo'q (Mustafoyeva Madina 282.docx fayli boshqa shaxs — Ortiq qiziga tegishli bo'lgani uchun uzildi)",
    24: 'TEKSHIRILMAGAN',
    25: 'KUTILMOQDA'
}

# CASE 3: T/R 91 (Row 92 in Barcha Talabalar, 26-04)
# Qodirova Sabina Suxrob qizi -> uzildi
c3_updates = {
    9: None,
    10: None,
    11: None,
    12: None,
    13: None,
    14: None,
    15: None,
    16: None,
    17: None,
    18: None,
    19: None,
    21: 'FAYL_YOQ',
    22: "Fayl yo'q (Qodirova Shabnam 137.docx fayli boshqa shaxs — Shabnam O'tkir qiziga tegishli bo'lgani uchun uzildi)",
    24: 'TEKSHIRILMAGAN',
    25: 'KUTILMOQDA'
}

# CASE 4: T/R 152 (Row 153 in Barcha Talabalar, 26-06)
# Quchqarova Zebiniso Farhod qizi (#187)
c4_updates = {
    7: 'Farhod qizi',
    8: 'Quchqarova Zebiniso Farhod qizi',
    9: 'QUCHQAROVA ZEBINISO',
    10: 'AE1143288',
    11: '62911085630012',
    12: '10.01.2025',
    13: '29.11.2008',
    21: 'CHALA',
    22: "Otasining ismi ID-karta AE1143288 bo'yicha 'Farhod qizi' deb tuzatildi (#187). Shahodatnoma kutilmoqda.",
    24: 'Pasport: MOS',
    25: 'TASDIQLANDI'
}

# CASE 5: T/R 165 (Row 166 in Barcha Talabalar, 26-07)
# Narziyeva Intizor Iskandar qizi (#212)
c5_updates = {
    2: 'Narziyeva Intizor',
    8: 'Narziyeva Intizor Iskandar qizi',
    9: 'NARZIEVA INTIZOR',
    10: 'AB5359926',
    11: '40211995680023',
    12: '12.12.2016',
    13: '02.11.1999',
    14: 'NARZIYEVA INTIZOR ISKANDAR QIZI',
    15: 'K 5047841',
    17: "Kitob san'at va xizmat ko'rsatish kasb-hunar kolleji",
    18: 'Kasb-hunar kolleji diplomi',
    19: '2018',
    21: 'TOPILDI',
    22: "Familiya biometrik pasport AB5359926 va kasb-hunar kolleji diplomi K 5047841 bo'yicha 'Narziyeva' deb tasdiqlandi (#212)",
    24: 'Pasport: MOS | Diplom: MOS',
    25: 'TASDIQLANDI'
}

# CASE 6: T/R 95 (Row 96 in Barcha Talabalar, 26-04)
# Tuychiyeva Farzona Dalerovna (#280)
c6_updates = {
    9: 'TUYCHIYEVA FARZONA',
    14: 'TUYCHIYEVA FARZONA DALEROVNA',
    15: 'UM 00064269',
    17: "34-sonli umumiy o'rta ta'lim maktabini",
    18: 'Shahodatnoma',
    19: '2026',
    22: "ID-karta AE4114421 va Shahodatnoma UM 00064269 bo'yicha to'liq tasdiqlandi (#280)",
    24: 'Pasport: MOS | Shahodatnoma: MOS',
    25: 'TASDIQLANDI'
}

all_cases = [
    (32, c1_updates, '26-02'),
    (54, c2_updates, '26-03'),
    (91, c3_updates, '26-04'),
    (152, c4_updates, '26-06'),
    (165, c5_updates, '26-07'),
    (95, c6_updates, '26-04'),
]

for tr_val, updates, group_name in all_cases:
    # 1. Barcha Talabalar sheetida
    r_all = find_row_by_tr(ws, tr_val)
    if r_all:
        apply_update_to_sheet(ws, r_all, updates)
        print(f"  [Barcha Talabalar] T/R {tr_val} (Row {r_all}) yangilandi.")
    else:
        print(f"  ⚠️ [Barcha Talabalar] T/R {tr_val} topilmadi!")

    # 2. Tegishli guruh sheetida
    if group_name in wb.sheetnames:
        g_sheet = wb[group_name]
        # Guruh sheetida qidirish (T/R yoki IFO bo'yicha)
        r_grp = None
        for gr in range(2, g_sheet.max_row + 1):
            if str(g_sheet.cell(gr, 1).value or '').strip() == str(tr_val):
                r_grp = gr
                break
        if not r_grp:
            # Agar T/R guruh ichida 1..N bo'lsa, Col 2 (IFO) yoki Col 8 (FISH) bo'yicha tekshirish
            old_names = ['Saidaliyeva', 'Saydalimova', 'Mustafayeva', 'Qodirova', 'Quchqarova', 'Nazirova', 'Narziyeva', 'Tuychiyeva']
            for gr in range(2, g_sheet.max_row + 1):
                cell_v = str(g_sheet.cell(gr, 2).value or '') + ' ' + str(g_sheet.cell(gr, 8).value or '')
                for on in old_names:
                    if on.lower() in cell_v.lower() and str(tr_val) in str(g_sheet.cell(gr, 1).value or ''):
                        r_grp = gr
                        break
        if r_grp:
            apply_update_to_sheet(g_sheet, r_grp, updates)
            print(f"  [{group_name}] Row {r_grp} yangilandi.")

wb.save(f_main)
print(f"✅ {f_main} muvaffaqiyatli saqlandi!")

# 3. FOYDALANUVCHINING ASL JADVALI: to'liq 1-KURS 2026-2027.xlsx NI HAM TUZATISH
wb_orig = openpyxl.load_workbook(f_orig)
ws_orig = wb_orig['Talabalar']

for r in range(13, ws_orig.max_row + 1):
    group = str(ws_orig.cell(r, 2).value or '').strip()
    tr = str(ws_orig.cell(r, 3).value or '').strip()
    fam = str(ws_orig.cell(r, 4).value or '').strip()
    ism = str(ws_orig.cell(r, 5).value or '').strip()
    ota = str(ws_orig.cell(r, 6).value or '').strip()
    ifo = str(ws_orig.cell(r, 7).value or '').strip()

    # Case 1: 26-02, T/R 17 (Saidaliyeva -> Saydalimova)
    if group == '26-02' and ('saidaliyeva' in fam.lower() or 'saydalimova' in fam.lower()):
        ws_orig.cell(r, 4, 'Saydalimova')
        ws_orig.cell(r, 7, 'Saydalimova Ruxshona Sharofiddin qizi')
        ws_orig.cell(r, 8, '18.07.2008')
        ws_orig.cell(r, 9, 'AE3352294')
        ws_orig.cell(r, 10, '61807085590034')
        ws_orig.cell(r, 11, '10.07.2025')
        ws_orig.cell(r, 12, 'UM 03754198')
        ws_orig.cell(r, 13, "29-sonli umumiy o'rta ta'lim maktabi")
        ws_orig.cell(r, 15, '2026')
        ws_orig.cell(r, 16, '74')
        print(f"  [to'liq 1-KURS] Row {r} (Saydalimova Ruxshona) tuzatildi va to'ldirildi.")

    # Case 4: 26-06, T/R 19 (Quchqarova Zebiniso Zebiniso qizi -> Farhod qizi)
    elif group == '26-06' and 'quchqarova' in fam.lower():
        ws_orig.cell(r, 6, 'Farhod qizi')
        ws_orig.cell(r, 7, 'Quchqarova Zebiniso Farhod qizi')
        ws_orig.cell(r, 11, '10.01.2025')
        ws_orig.cell(r, 16, '187')
        print(f"  [to'liq 1-KURS] Row {r} (Quchqarova Zebiniso Farhod qizi) otasining ismi tuzatildi.")

    # Case 5: 26-07, T/R 5 (Nazirova Intizor -> Narziyeva Intizor)
    elif group == '26-07' and ('nazirova' in fam.lower() or 'narziyeva' in fam.lower()):
        ws_orig.cell(r, 4, 'Narziyeva')
        ws_orig.cell(r, 7, 'Narziyeva Intizor Iskandar qizi')
        ws_orig.cell(r, 10, '40211995680023')
        print(f"  [to'liq 1-KURS] Row {r} (Narziyeva Intizor Iskandar qizi) familiyasi va PINFL tuzatildi.")

    # Case 6: 26-04, T/R 21 (Tuychiyeva Farzona)
    elif group == '26-04' and ('tuychiyeva' in fam.lower() or "to'ychiyeva" in fam.lower()):
        ws_orig.cell(r, 8, '05.08.2009')
        ws_orig.cell(r, 9, 'AE4114421')
        ws_orig.cell(r, 11, '09.09.2025')
        ws_orig.cell(r, 12, 'UM 00064269')
        ws_orig.cell(r, 13, "34-sonli umumiy o'rta ta'lim maktabi")
        ws_orig.cell(r, 15, '2026')
        ws_orig.cell(r, 16, '280')
        print(f"  [to'liq 1-KURS] Row {r} (Tuychiyeva Farzona) hujjat ma'lumotlari to'ldirildi.")

wb_orig.save(f_orig)
print(f"✅ {f_orig} muvaffaqiyatli saqlandi!")

# 4. SCRATCH/VERIFY_RESULTS.JSON NI HAM YANGILASH
res_path = os.path.join(BASE_DIR, 'arxiv', 'scratch', 'verify_results.json')
if os.path.exists(res_path):
    with open(res_path, 'r', encoding='utf-8') as f:
        v_data = json.load(f)
    for it in v_data:
        r_num = it.get('row')
        # Row 33
        if r_num == 33:
            it['ifo'] = 'Saydalimova Ruxshona'
            it['file'] = 'Saydalimova Ruxshona 74.docx'
            it['err'] = None
            if 'excel' in it:
                it['excel']['pass'] = 'AE3352294'
                it['excel']['pinfl'] = '61807085590034'
                it['excel']['ber'] = '10.07.2025'
                it['excel']['dob'] = '18.07.2008'
                it['excel']['cert'] = 'UM 03754198'
                it['excel']['maktab'] = "29-sonli umumiy o'rta ta'lim maktabi"
                it['excel']['doc_tur'] = 'Shahodatnoma'
                it['excel']['yil'] = '2026'
        # Row 55
        elif r_num == 55:
            it['file'] = None
            it['err'] = 'FAYL_YOQ'
            if 'excel' in it:
                for k in ['pass', 'pinfl', 'ber', 'dob', 'cert', 'qr', 'maktab', 'doc_tur', 'yil']:
                    it['excel'][k] = ''
        # Row 92
        elif r_num == 92:
            it['file'] = None
            it['err'] = 'FAYL_YOQ'
            if 'excel' in it:
                for k in ['pass', 'pinfl', 'ber', 'dob', 'cert', 'qr', 'maktab', 'doc_tur', 'yil']:
                    it['excel'][k] = ''
        # Row 96
        elif r_num == 96:
            it['file'] = 'To’ychiyeva  Farzona 280.docx'
            it['err'] = None
            if 'excel' in it:
                it['excel']['pass'] = 'AE4114421'
                it['excel']['ber'] = '09.09.2025'
                it['excel']['dob'] = '05.08.2009'
                it['excel']['cert'] = 'UM 00064269'
                it['excel']['maktab'] = "34-sonli umumiy o'rta ta'lim maktabi"
                it['excel']['doc_tur'] = 'Shahodatnoma'
                it['excel']['yil'] = '2026'
        # Row 153
        elif r_num == 153:
            it['ifo'] = 'Quchqarova Zebiniso'
            it['file'] = 'Kuchqarova Zebiniso 187.docx'
            it['err'] = None
            if 'excel' in it:
                it['excel']['pass'] = 'AE1143288'
                it['excel']['pinfl'] = '62911085630012'
                it['excel']['ber'] = '10.01.2025'
                it['excel']['dob'] = '29.11.2008'
        # Row 166
        elif r_num == 166:
            it['ifo'] = 'Narziyeva Intizor'
            it['file'] = 'Nazirova  Intizor 212.docx'
            it['err'] = None
            if 'excel' in it:
                it['excel']['pass'] = 'AB5359926'
                it['excel']['pinfl'] = '40211995680023'
                it['excel']['ber'] = '12.12.2016'
                it['excel']['dob'] = '02.11.1999'
                it['excel']['cert'] = 'K 5047841'
                it['excel']['maktab'] = "Kitob san'at va xizmat ko'rsatish kasb-hunar kolleji"
                it['excel']['doc_tur'] = 'Diplom'
                it['excel']['yil'] = '2018'

    with open(res_path, 'w', encoding='utf-8') as f:
        json.dump(v_data, f, ensure_ascii=False, indent=2)
    print(f"✅ {res_path} yangilandi!")
