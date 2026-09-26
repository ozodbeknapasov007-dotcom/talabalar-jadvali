# -*- coding: utf-8 -*-
"""
BOSQICH 3: Tasdiqlangan 33 ta talaba faylini aniq OCR qilish
============================================================
QR + Barcode + AI Vision (Pasport, PINFL, Berilgan sana, Ota ismi, Shahodatnoma, Maktab, Yil)
"""

import openpyxl, docx, os, sys, re, io, json, base64, urllib.request, time
from PIL import Image

try:
    import zxingcpp
except:
    zxingcpp = None
try:
    import pypdf
except:
    pypdf = None

sys.stdout.reconfigure(encoding='utf-8')

OPENROUTER_API_KEY = "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"
FILES_DIR = 'hujjatlar/shartnomalar'
EXCEL_PATH = 'data/Talabalar_Toliq_Royxati.xlsx'

# Tasdiqlangan 33 ta talaba va ularga biriktirilgan fayllar xaritasi
TARGET_MAPPINGS = [
    {"shnum": "3",   "ism": "Sayfillayeva Shodiya",   "file": "Sayfillayeva Shodiyona 03.docx"},
    {"shnum": "4",   "ism": "Karamova Dilroz",        "file": "Koramova Dilroz 04.docx"},
    {"shnum": "20",  "ism": "Temitboyeva Iroda",      "file": "Temirboyeva Iroda 20.docx"},
    {"shnum": "28",  "ism": "Xaydarova Xayriniso",    "file": "Haydarova Xayriniso 28.docx"},
    {"shnum": "41",  "ism": "Ergashova Sayyora",      "file": "Ergasheva Sayyora 41.docx"},
    {"shnum": "43",  "ism": "Muxammadiyeva Fazlina",  "file": "Muhammadiyeva Fazlina 43 chala.docx"},
    {"shnum": "53",  "ism": "Rerajaboliyeva Marjona", "file": "Rejaboliyeva Marjona 53.docx"},
    {"shnum": "59",  "ism": "Jumayeva Muhlisa",       "file": "Jumayeva Muxlisa 60.docx"},
    {"shnum": "60",  "ism": "Meyliqulova Hilola",     "file": "Meliqulova Hilola 58.docx"},
    {"shnum": "66",  "ism": "Isomiddinova Ruxshona",  "file": "Isomiddinova Ruhshona 65.docx"},
    {"shnum": "75",  "ism": "Norboboyeva Mahliyo",    "file": "Yorboboyeva Mahliyo 75.docx"},
    {"shnum": "89",  "ism": "Turayeva Dilnoz",        "file": "Jurayeva Dilnoz 89.docx"},
    {"shnum": "93",  "ism": "Zarifjonova Marjona",    "file": "Zarifova Marjona 93.docx"},
    {"shnum": "100", "ism": "Ruzimurodova Larisa",    "file": "Ro'zimurodova Larisa 100.docx"},
    {"shnum": "111", "ism": "Turayeva Adiba",         "file": "Jurayeva Adiba 111.docx"},
    {"shnum": "115", "ism": "Absattorova Nafisa",     "file": "Absottorova Nafisa 115.docx"},
    {"shnum": "123", "ism": "Ergashova Shabbona",     "file": "Ergasheva Shabbona 123.docx"},
    {"shnum": "161", "ism": "To'xtayeva Durdona",     "file": "Tuxtayeva Durdona 161.docx"},
    {"shnum": "162", "ism": "Shukurova Nariye",       "file": "Shukriyeva Nariye 162.docx"},
    {"shnum": "165", "ism": "Temurqulova Shahzoda",   "file": "Temirqulova Shahzoda 165.docx"},
    {"shnum": "177", "ism": "Sa'dullayeva Shahzoda",  "file": "Sa’dullayeva Shahzoda 177.docx"},
    {"shnum": "185", "ism": "Qurbonova Charos",       "file": "Qurbon0ova Charos 185.docx"},
    {"shnum": "186", "ism": "Abdug'apporova E'zoza",  "file": "Abdug’apporova Ezoza186.docx"},
    {"shnum": "187", "ism": "Quchqarova Zebiniso",    "file": "Kuchqarova Zebiniso 187.docx"},
    {"shnum": "194", "ism": "Umarova Lobar",          "file": "Umarovqa Lobar 194.docx"},
    {"shnum": "204", "ism": "Abdukarimova Zahrobonu", "file": "Abdurahmonova Zahrabonu 204.docx"},
    {"shnum": "206", "ism": "Mustafayeva Nargiza",    "file": "Mustafoyeva Nargiza    205.docx"},
    {"shnum": "216", "ism": "Hamdamov Shahzod",       "file": "Hamdamov Shahzod 216.docx"},
    {"shnum": "221", "ism": "Eshnazarova Dilnoza",    "file": "Eshnazarova Dilnoza 221.docx"},
    {"shnum": "232", "ism": "Xo'jamurotova Durdona",  "file": "Xujamuratova Durdona 232.docx"},
    {"shnum": "383", "ism": "Izzatullayeva Vazira",   "file": "Izatullayeva Vazira 383.docx"},
    {"shnum": "390", "ism": "Mahmamurodova Sevara",   "file": "Maxmudmurodova Sevara 390.docx"},
    {"shnum": "391", "ism": "Ovlayeva Ma'mura",       "file": "Ovlayeva Ma'mura  391 chala.docx"},
]

def ai_vision_deep(img_bytes):
    b64 = base64.b64encode(img_bytes).decode('utf-8')
    payload = {
        "model": "google/gemini-2.5-flash",
        "messages": [{"role": "user", "content": [
            {"type": "text", "text": """Siz O'zbekiston hujjatlarini tahlil qiluvchi AI mutaxassisisiz.
Quyidagi JSON formatda juda aniq javob qaytaring:
{
  "tur": "pasport" yoki "id" yoki "shahodatnoma" yoki "diplom",
  "pass_ser": "Pasport seriya va 7 raqam (masalan AE1234567, AD9876543)",
  "pinfl": "14 xonali JSHSHIR raqami",
  "dob": "Tug'ilgan sana (DD.MM.YYYY)",
  "pass_ber": "FAQAT Pasport yoki ID-karta berilgan sanasi (DD.MM.YYYY). Shahodatnoma sanasini HECH QACHON yozmang!",
  "ota": "Otasining ismi (masalan: ALISHER QIZI, SHOKIR O'G'LI)",
  "sh_ser": "Shahodatnoma seriyasi (UM yoki K)",
  "sh_num": "Shahodatnoma yoki Diplom raqami",
  "maktab": "Tugatgan maktab yoki kollejning to'liq nomi",
  "yil": "Bitirgan yili (YYYY)"
}
Faqat JSON qaytaring. Ma'lumot yo'q bo'lsa null yozing."""},
            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}
        ]}],
        "temperature": 0.05
    }
    try:
        req = urllib.request.Request(
            "https://openrouter.ai/api/v1/chat/completions",
            data=json.dumps(payload).encode('utf-8'),
            headers={"Authorization": f"Bearer {OPENROUTER_API_KEY}", "Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req, timeout=25) as resp:
            res = json.loads(resp.read().decode('utf-8'))
            txt = res['choices'][0]['message']['content'].strip()
            txt = re.sub(r'^```json\s*', '', txt)
            txt = re.sub(r'\s*```$', '', txt)
            return json.loads(txt)
    except Exception as e:
        print(f"    [AI xatolik]: {e}", flush=True)
        return {}

def qr_cert_pdf(url):
    r = {}
    if not pypdf: return r
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=12) as resp:
            txt = pypdf.PdfReader(io.BytesIO(resp.read())).pages[0].extract_text()
        m = re.search(r'UM\s*[№#\.\s]*(\d{7,8})', txt)
        if not m: m = re.search(r'\b(03\d{6}|01\d{6}|\d{8})\b', txt)
        if m: r['sh'] = f"UM {m.group(1)}"
        m = re.search(r'(\d{4})\s*-yilda', txt)
        if m: r['yil'] = m.group(1)
        m = re.search(r'(\d{1,3}\s*-\s*sonli\s+[\w\s\'\-]+(?:maktabini|maktabi|maktab))', txt, re.I)
        if not m: m = re.search(r'([\w\s\'\-]+(?:umumiy o\'rta ta\'lim maktabi|maktabini|maktab))', txt, re.I)
        if m:
            s = re.sub(r'^\d{4}\s*','', m.group(1).strip()); s = re.sub(r'ini$','i',s)
            r['maktab'] = s
    except: pass
    return r

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

print(f"\n{'='*70}")
print(f"  TASDIQLANGAN 33 TA TALABANI QAYTA OCR QILISH BOSHLANDI")
print(f"{'='*70}\n")

processed_count = 0

for idx, item in enumerate(TARGET_MAPPINGS, 1):
    target_sh = str(item['shnum']).strip()
    target_ism = str(item['ism']).strip()
    target_file = item['file']
    
    # Exceldagi qatorini topamiz
    target_row = None
    for r in range(2, ws.max_row + 1):
        sh_val = str(ws.cell(row=r, column=5).value or '').strip()
        ism_val = str(ws.cell(row=r, column=2).value or '').strip()
        
        if sh_val == target_sh or (target_ism.lower() in ism_val.lower()):
            target_row = r
            break
            
    if not target_row:
        print(f"[{idx:2d}/33] [OGOHLANTIRISH] Excelda talaba topilmadi: #{target_sh} {target_ism}", flush=True)
        continue

    file_path = os.path.join(FILES_DIR, target_file)
    if not os.path.exists(file_path):
        print(f"[{idx:2d}/33] [XATO] Fayl topilmadi: {file_path}", flush=True)
        continue

    print(f"\n[{idx:2d}/33] #{target_sh} {target_ism} (Qator {target_row})", flush=True)
    print(f"    [FAYL] {target_file}", flush=True)

    try:
        doc = docx.Document(file_path)
    except Exception as e:
        print(f"    [XATO] Docx ochishda xato: {e}", flush=True)
        continue

    img_blobs = [rel.target_part.blob for rel in doc.part.rels.values() if 'image' in rel.target_ref and len(rel.target_part.blob) > 8000]

    pass_val = pinfl = dob = pass_ber = ota = sh_val = sh_qr = maktab = yil = cert_tur = ''

    # 1. QR kod va Barcode tahlili
    if zxingcpp:
        for b in img_blobs:
            im = Image.open(io.BytesIO(b)).convert('RGB')
            for angle in [0, 90, 180, 270]:
                bcs = zxingcpp.read_barcodes(im.rotate(angle, expand=True))
                if bcs:
                    for bc in bcs:
                        t = bc.text
                        if 'shahodatnoma' in t.lower() or 'e-shahodatnoma' in t.lower():
                            sh_qr = t
                            pd = qr_cert_pdf(t)
                            if pd.get('sh'):     sh_val  = pd['sh']
                            if pd.get('maktab'): maktab  = pd['maktab']
                            if pd.get('yil'):    yil     = pd['yil']
                            print(f"    [QR] Shahodatnoma: {sh_val} | {maktab} | {yil}", flush=True)
                        elif 'IUUZB' in t:
                            m = re.search(r'IUUZB([A-Z]{2}\d{7})\d?([3-6]\d{13})', t)
                            if m:
                                pass_val = m.group(1)
                                pinfl    = m.group(2)
                                p = m.group(2)
                                cent = '19' if p[0] in ['3','4'] else '20'
                                dob = f"{p[1:3]}.{p[3:5]}.{cent}{p[5:7]}"
                                print(f"    [PASPORT QR]: {pass_val} | PINFL: {pinfl} | DOB: {dob}", flush=True)
                    break

    # 2. AI Vision tahlili (agar to'liq bo'lmasa yoki pasport berilgan sana / ota ismi kerak bo'lsa)
    if not pass_val or not sh_val or not ota or not pass_ber:
        print(f"    [AI] Vision tahlili ishga tushmoqda...", flush=True)
        for b in img_blobs:
            im = Image.open(io.BytesIO(b)).convert('RGB')
            buf = io.BytesIO()
            im.save(buf, format='JPEG', quality=90)
            ai = ai_vision_deep(buf.getvalue())
            if not ai: continue
            
            h = str(ai.get('tur') or '').lower()
            if 'pasport' in h or 'id' in h or ai.get('pass_ser'):
                if not pass_val and ai.get('pass_ser') and str(ai['pass_ser']).lower() != 'null':
                    pass_val = str(ai['pass_ser']).replace(' ','').upper()
                if not pinfl and ai.get('pinfl') and str(ai['pinfl']).lower() != 'null':
                    pinfl = str(ai['pinfl']).strip()
                if not dob and ai.get('dob') and str(ai['dob']).lower() != 'null':
                    dob = str(ai['dob']).strip()
                if not pass_ber and ai.get('pass_ber') and str(ai['pass_ber']).lower() != 'null':
                    pass_ber = str(ai['pass_ber']).strip()
                if not ota and ai.get('ota') and str(ai['ota']).lower() != 'null':
                    ota = str(ai['ota']).strip()
                print(f"    AI Pasport: {pass_val} | Ota: {ota} | PINFL: {pinfl} | Berilgan: {pass_ber}", flush=True)
            elif 'shahodatnoma' in h or 'diplom' in h or ai.get('sh_num'):
                if not sh_val:
                    ss = str(ai.get('sh_ser') or '').strip()
                    sn = str(ai.get('sh_num') or '').strip()
                    if sn and sn.lower() != 'null':
                        sh_val = f"{ss} {sn}".strip()
                if not maktab and ai.get('maktab') and str(ai['maktab']).lower() != 'null':
                    maktab = str(ai['maktab']).strip()
                if not yil and ai.get('yil') and str(ai['yil']).lower() != 'null':
                    yil = str(ai['yil']).strip()
                if 'diplom' in h: cert_tur = 'Diplom'
                print(f"    AI Hujjat: {sh_val} | Maktab: {maktab} | Yil: {yil}", flush=True)
            time.sleep(0.4)

    # PINFL dan DOB hisoblash
    if not dob and pinfl and len(pinfl) == 14 and pinfl.isdigit():
        p = pinfl
        cent = '19' if p[0] in ['3','4'] else '20'
        dob = f"{p[1:3]}.{p[3:5]}.{cent}{p[5:7]}"

    # Excelga yozish
    fish = f"{target_ism} {ota}".strip() if ota else target_ism
    ws.cell(row=target_row, column=7,  value=ota)
    ws.cell(row=target_row, column=8,  value=fish)
    # 9-ustun (Passport bo'yicha F.I.SH) — bu yerda yozilmaydi,
    # uni scripts/verify_passport_names.py boshqaradi.
    ws.cell(row=target_row, column=10, value=pass_val)
    ws.cell(row=target_row, column=11, value=pinfl)
    ws.cell(row=target_row, column=12, value=pass_ber)
    ws.cell(row=target_row, column=13, value=dob)
    ws.cell(row=target_row, column=15, value=sh_val)
    ws.cell(row=target_row, column=16, value=sh_qr)
    ws.cell(row=target_row, column=17, value=maktab)
    ws.cell(row=target_row, column=18, value="Umumiy o'rta maktab" if cert_tur != 'Diplom' else 'Kollej')
    ws.cell(row=target_row, column=19, value=yil)
    ws.cell(row=target_row, column=21, value="TOPILDI" if pass_val or sh_val else "CHALA")
    ws.cell(row=target_row, column=22, value="Qayta tahlil: Fayl moslandi")
    
    wb.save(EXCEL_PATH)
    processed_count += 1
    print(f"    [OK] Saqlandi: Pass={pass_val or 'yoq'} | Sh={sh_val or 'yoq'} | Ota={ota or 'yoq'}", flush=True)

print(f"\n{'='*70}")
print(f"  [OK] Jami {processed_count} ta talaba ma'lumotlari muvaffaqiyatli yangilandi!")
print(f"{'='*70}")
