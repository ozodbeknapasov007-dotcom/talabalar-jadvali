import filecmp
import glob
import json
import os
import re
import shutil
import openpyxl
from reportlab import rl_config
# Ma'lumot o'zgarmagan bo'lsa PDF baytlari ham o'zgarmasligi shart.
# Aks holda ReportLab har safar yangi CreationDate/ModDate va /ID yozadi,
# natijada har bir qayta generatsiyada PDFlar "o'zgargan" bo'lib
# git ga tushadi va har bir commit ~800 KB keraksiz ma'lumot bilan shishadi.
rl_config.invariant = 1
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# qayta_tekshiruv ichida bo'lsa, asosiy loyiha papkasi bir pog'ona yuqorida:
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(CURRENT_DIR) if os.path.basename(CURRENT_DIR) == 'qayta_tekshiruv' else CURRENT_DIR

# Fontlarni ro'yxatdan o'tkazish
try:
    pdfmetrics.registerFont(TTFont('Times', 'C:/Windows/Fonts/times.ttf'))
    pdfmetrics.registerFont(TTFont('Times-Bold', 'C:/Windows/Fonts/timesbd.ttf'))
    FONT_NORMAL = 'Times'
    FONT_BOLD = 'Times-Bold'
except Exception as e:
    FONT_NORMAL = 'Helvetica'
    FONT_BOLD = 'Helvetica-Bold'

def load_group_info():
    """data/guruhlar.json dan guruhlar ma'lumotlarini yuklaydi."""
    guruhlar_path = os.path.join(BASE_DIR, 'data', 'guruhlar.json')
    leaders = {}
    titles = {}
    if os.path.exists(guruhlar_path):
        try:
            with open(guruhlar_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                for k, v in data.items():
                    leaders[k] = v.get('rahbar', '—')
                    titles[k] = v.get('yonalish', 'Hamshiralik ishi')
        except Exception as e:
            print(f"guruhlar.json yuklashda xatolik: {e}")
    return leaders, titles

GROUP_LEADERS, GROUP_TITLES = load_group_info()

def export_pdf_to_images(pdf_path, dpi=250):
    """PDF sahifalarini yuqori sifatli JPG rasm formatida saqlaydi."""
    try:
        import pymupdf
        doc = pymupdf.open(pdf_path)
        base = os.path.splitext(pdf_path)[0]
        img_paths = []
        for idx, page in enumerate(doc):
            pix = page.get_pixmap(dpi=dpi)
            out_img = f"{base}.jpg" if len(doc) == 1 else f"{base}_page_{idx+1}.jpg"
            pix.save(out_img)
            img_paths.append(out_img)
        doc.close()
        return img_paths
    except Exception as e:
        print(f"Rasmga eksport qilishda xatolik ({pdf_path}): {e}")
        return []

def build_group_flowables(group_code, students, avail_width):
    is_withdrawn = ('chiqaril' in str(group_code).lower() or str(group_code) in ['N', 'n', 'WITHDRAWN'])
    is_akademik = ('akademik' in str(group_code).lower())

    if is_withdrawn:
        g_students = [s for s in students if 'chiqaril' in str(s.get('group', '')).lower() or str(s.get('group', '')) in ['N', 'n', 'WITHDRAWN']]
        leader = "Texnikum ma'muriyati"
        g_title = "Safdan chiqarilganlar"
    elif is_akademik:
        g_students = [s for s in students if 'akademik' in str(s.get('group', '')).lower()]
        leader = "Texnikum ma'muriyati"
        g_title = "Akademik ta'til olganlar"
    else:
        g_students = [s for s in students if s.get('group') == group_code]
        leader = GROUP_LEADERS.get(group_code, "—")
        g_title = GROUP_TITLES.get(group_code, "Hamshiralik ishi")
    g_students.sort(key=lambda x: re.sub(r"['`‘’ʻʼ´\-_.]", "", str(x.get('fio') or x.get('ism', '')).lower()))

    elements = []
    count = len(g_students)

    # 1 TA VAROQQA (A4) ANIQ SIG'DIRISH LOGIKASI:
    # 35 tagacha talaba to'liq 1 ta A4 varoqqa sig'adi
    if count <= 20:
        row_height = 25.0
        font_size = 10.5
        header_font_size = 11.0
        padding_size = 3.5
    elif count <= 27:
        row_height = 22.5
        font_size = 10.0
        header_font_size = 10.5
        padding_size = 2.8
    else:
        row_height = 19.5
        font_size = 9.0
        header_font_size = 9.5
        padding_size = 2.0

    title_style = ParagraphStyle(
        'DocTitle',
        fontName=FONT_BOLD,
        fontSize=15,
        leading=18,
        alignment=1,
        textColor=colors.HexColor('#0f172a')
    )

    info_left_style = ParagraphStyle(
        'InfoLeft',
        fontName=FONT_BOLD,
        fontSize=10,
        leading=13,
        alignment=0,
        textColor=colors.HexColor('#1e293b')
    )

    info_right_style = ParagraphStyle(
        'InfoRight',
        fontName=FONT_BOLD,
        fontSize=10,
        leading=13,
        alignment=2,
        textColor=colors.HexColor('#1e293b')
    )

    header_para_style = ParagraphStyle(
        'HeaderCell',
        fontName=FONT_BOLD,
        fontSize=header_font_size,
        leading=header_font_size + 2,
        alignment=1,
        textColor=colors.HexColor('#0f172a')
    )

    cell_name_style = ParagraphStyle(
        'CellName',
        fontName=FONT_BOLD,
        fontSize=font_size,
        leading=font_size + 2,
        alignment=0,
        textColor=colors.HexColor('#0f172a')
    )

    cell_center_style = ParagraphStyle(
        'CellCenter',
        fontName=FONT_NORMAL,
        fontSize=font_size,
        leading=font_size + 2,
        alignment=1,
        textColor=colors.HexColor('#334155')
    )

    if is_withdrawn:
        header_table_data = [
            [
                Paragraph("Shahrisabz Tibbiyot Texnikumi", title_style),
                ""
            ],
            [
                Paragraph("Maxsus ro'yxat: <b>Talabalar safidan chiqarilganlar</b>", info_left_style),
                Paragraph("Holati: <b>Safdan chiqarilgan</b>", info_right_style)
            ]
        ]
    elif is_akademik:
        header_table_data = [
            [
                Paragraph("Shahrisabz Tibbiyot Texnikumi", title_style),
                ""
            ],
            [
                Paragraph("Maxsus ro'yxat: <b>Akademik ta'til olganlar</b>", info_left_style),
                Paragraph("Holati: <b>Akademik ta'tilda</b>", info_right_style)
            ]
        ]
    else:
        header_table_data = [
            [
                Paragraph("Shahrisabz Tibbiyot Texnikumi", title_style),
                ""
            ],
            [
                Paragraph(f"Yo'nalish: <b>{g_title}</b> | Guruh: <b>{group_code}</b>", info_left_style),
                Paragraph(f"Guruh rahbari: <b>{leader}</b>", info_right_style)
            ]
        ]

    ht = Table(header_table_data, colWidths=[avail_width * 0.62, avail_width * 0.38])
    ht.setStyle(TableStyle([
        ('SPAN', (0, 0), (1, 0)),
        ('ALIGN', (0, 0), (1, 0), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (1, 0), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 1), (-1, 1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(ht)

    w_no = 26
    w_sh = 65
    w_dob = 80
    w_imzo = 95
    w_name = avail_width - (w_no + w_sh + w_dob + w_imzo)

    col_widths = [w_no, w_sh, w_name, w_dob, w_imzo]

    headers = [
        Paragraph("№", header_para_style),
        Paragraph("Shartnoma", header_para_style),
        Paragraph("F.I.SH (To'liq)", header_para_style),
        Paragraph("Tug'ilgan sana", header_para_style),
        Paragraph("Imzo", header_para_style)
    ]

    data = [headers]

    for idx, st in enumerate(g_students, start=1):
        sh_val = str(st.get('shnum', '')).strip()
        if sh_val and sh_val not in ['—', '-', 'None', 'nan']:
            sh_num = f"#{sh_val}" if not sh_val.startswith('#') else sh_val
        else:
            sh_num = "—"
        row = [
            Paragraph(str(idx), cell_center_style),
            Paragraph(sh_num, cell_center_style),
            Paragraph(st.get('fio', st.get('ism', '')), cell_name_style),
            Paragraph(st.get('dob') or "—", cell_center_style),
            Paragraph("", cell_center_style)
        ]
        data.append(row)

    t = Table(data, colWidths=col_widths, rowHeights=[row_height + 3] + [row_height] * len(g_students))
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), padding_size),
        ('BOTTOMPADDING', (0, 0), (-1, -1), padding_size),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.6, colors.HexColor('#94a3b8')),
        ('BOX', (0, 0), (-1, -1), 1.0, colors.HexColor('#334155')),
        ('LINEBELOW', (0, 0), (-1, 0), 1.2, colors.HexColor('#0f172a')),
    ]

    for row_idx in range(1, len(data)):
        if row_idx % 2 == 0:
            t_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor('#f8fafc')))

    t.setStyle(TableStyle(t_style))
    elements.append(t)
    return elements

def create_single_group_pdf(group_code, students, output_pdf_path):
    os.makedirs(os.path.dirname(output_pdf_path), exist_ok=True)
    doc = SimpleDocTemplate(
        output_pdf_path,
        pagesize=A4,
        leftMargin=24,
        rightMargin=24,
        topMargin=16,
        bottomMargin=14
    )
    avail_width = A4[0] - 48
    elements = build_group_flowables(group_code, students, avail_width)
    doc.build(elements)
    return output_pdf_path

def create_all_groups_combined_pdf(groups, students, output_pdf_path):
    """
    Barcha guruhlarni 1 ta umumiy PDF hujjatga birlashtiradi.
    Har bir guruh navbati bilan alohida-alohida A4 varoqda chiqadi (PageBreak orqali).
    """
    os.makedirs(os.path.dirname(output_pdf_path), exist_ok=True)
    doc = SimpleDocTemplate(
        output_pdf_path,
        pagesize=A4,
        leftMargin=24,
        rightMargin=24,
        topMargin=16,
        bottomMargin=14
    )
    avail_width = A4[0] - 48
    elements = []
    for idx, g in enumerate(groups):
        g_elements = build_group_flowables(g, students, avail_width)
        elements.extend(g_elements)
        if idx < len(groups) - 1:
            elements.append(PageBreak())
    doc.build(elements)
    return output_pdf_path

def _replace_if_changed(tmp_path, path):
    """tmp_path ni path o'rniga qo'yadi; mazmun avvalgidek bo'lsa False qaytaradi."""
    if os.path.exists(path) and filecmp.cmp(tmp_path, path, shallow=False):
        os.remove(tmp_path)
        return False
    os.replace(tmp_path, path)
    return True

def _jpgs_of(pdf_path):
    """export_pdf_to_images yaratgan rasmlar: <nom>.jpg yoki <nom>_page_N.jpg"""
    base = os.path.splitext(pdf_path)[0]
    return sorted(glob.glob(glob.escape(base) + '.jpg') + glob.glob(glob.escape(base) + '_page_*.jpg'))

def _mirror(src, dirs):
    """Faylni boshqa papkalarga nusxalash (mazmun bir xil bo'lsa tegmaydi)."""
    for d in dirs:
        dst = os.path.join(d, os.path.basename(src))
        if not (os.path.exists(dst) and filecmp.cmp(src, dst, shallow=False)):
            shutil.copyfile(src, dst)

def build_all_group_pdfs():
    # students.json yoki Excel bazasidan talabalarni o'qish
    students_json_path = os.path.join(BASE_DIR, 'data', 'students.json')
    excel_path = os.path.join(BASE_DIR, 'data', 'Talabalar_Toliq_Royxati.xlsx')

    students = []
    if os.path.exists(students_json_path):
        try:
            with open(students_json_path, 'r', encoding='utf-8') as f:
                raw_st = json.load(f)
                for s in raw_st:
                    students.append({
                        'ism': s.get('ism', ''),
                        'ota': s.get('ota', ''),
                        'fio': s.get('fish') or f"{s.get('ism', '')} {s.get('ota', '')}".strip(),
                        'dob': s.get('dob', ''),
                        'group': s.get('group', ''),
                        'shnum': s.get('shnum', '')
                    })
        except Exception as e:
            print(f"students.json o'qishda xatolik: {e}")

    if not students and os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb.active
        for r in range(2, ws.max_row + 1):
            ism = str(ws.cell(row=r, column=2).value or '').strip()
            shnum = str(ws.cell(row=r, column=5).value or '').strip()
            ota = str(ws.cell(row=r, column=7).value or '').strip()
            fish = str(ws.cell(row=r, column=8).value or '').strip()
            dob = str(ws.cell(row=r, column=13).value or '').strip()
            group = str(ws.cell(row=r, column=23).value or '').strip()
            if not ism and not shnum: continue
            full_fio = fish or f"{ism} {ota}".strip()
            students.append({
                'ism': ism,
                'ota': ota,
                'fio': full_fio,
                'dob': dob,
                'group': group,
                'shnum': shnum
            })

    if not students:
        print("[XATO] Talabalar ro'yxati topilmadi!")
        return {}

    global GROUP_LEADERS, GROUP_TITLES
    GROUP_LEADERS, GROUP_TITLES = load_group_info()

    # Barcha rasmiy guruhlar (guruhlar.json dan tartiblangan holda)
    official_groups = sorted(GROUP_LEADERS.keys())
    if not official_groups:
        official_groups = sorted(list(set(
            s['group'] for s in students
            if s.get('group') and 'chiqaril' not in s['group'].lower() and 'akademik' not in s['group'].lower()
        )))

    groups = list(official_groups)
    withdrawn_students = [s for s in students if 'chiqaril' in str(s.get('group', '')).lower() or str(s.get('group', '')) in ['N', 'n', 'WITHDRAWN']]
    if withdrawn_students:
        groups.append("Talabalar safidan chiqarilganlar")
    akademik_students = [s for s in students if 'akademik' in str(s.get('group', '')).lower()]
    if akademik_students:
        groups.append("Akademik ta'til olganlar")

    target_dirs = [
        os.path.join(BASE_DIR, 'hisobotlar', 'pdf_jurnallar'),
        os.path.join(BASE_DIR, 'qayta_tekshiruv', 'pdf_jurnallar')
    ]

    primary, mirrors = target_dirs[0], target_dirs[1:]
    for tdir in target_dirs:
        os.makedirs(tdir, exist_ok=True)
        # Eski bo'lingan sahifalar (_page_*.jpg) ni tozalash
        for old_page in glob.glob(os.path.join(tdir, '*_page_*.jpg')):
            try:
                os.remove(old_page)
            except Exception:
                pass

    generated_files = {}

    # 1. Alohida guruh PDF lari va ularning yuqori sifatli 1 ta varoqli JPG rasmlari
    for g in groups:
        if 'chiqaril' in g.lower():
            fname = "Guruh_Talabalar_safidan_chiqarilganlar.pdf"
        elif 'akademik' in g.lower():
            fname = "Guruh_Akademik_tatil_olganlar.pdf"
        else:
            fname = f"Guruh_{g}.pdf"

        out_file = os.path.join(primary, fname)
        tmp_file = out_file + '.tmp'
        create_single_group_pdf(g, students, tmp_file)
        changed = _replace_if_changed(tmp_file, out_file)
        generated_files[g] = out_file

        jpg_path = os.path.splitext(out_file)[0] + '.jpg'
        if changed or not os.path.exists(jpg_path):
            export_pdf_to_images(out_file, dpi=250)
        for p in [out_file] + _jpgs_of(out_file):
            _mirror(p, mirrors)

    # 2. Barcha rasmiy guruhlarni birlashtirgan YAGONA 1 ta A4 PDF (har bir guruh alohida varoqda)
    combined_file = os.path.join(primary, "Barcha_Guruhlar_Jurnali.pdf")
    create_all_groups_combined_pdf(official_groups, students, combined_file + '.tmp')
    _replace_if_changed(combined_file + '.tmp', combined_file)
    _mirror(combined_file, mirrors)
    generated_files["ALL"] = combined_file

    print(f"[OK] Barcha {len(groups)} ta guruh uchun 1-sahifali A4 PDF jurnallar va yuqori sifatli rasmlar yaratildi!")
    return generated_files

if __name__ == '__main__':
    build_all_group_pdfs()
