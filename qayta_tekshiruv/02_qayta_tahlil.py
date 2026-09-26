# -*- coding: utf-8 -*-
"""
BOSQICH 2 (v2): 30 tadan guruhlab - TEZKOR versiya
=======================================================
Strategiya:
1. AVVAL: Barcha QR kodni o'qiy + PDF dan ma'lumot oladi (AI yo'q, tez)
2. KEYIN: Faqat QR'siz yoki to'liq bo'lmaganlarga AI Flash ishlatadi
3. Har bir talaba natijasi DARHOL Excelga yoziladi
"""

import openpyxl, docx, os, sys, re, io, json, base64, urllib.request, time
from PIL import Image
try:
    import zxingcpp
except:
    zxingcpp = None
try:
    import pypdf
except:
    pypdf = None

sys.stdout.reconfigure(encoding='utf-8')

OPENROUTER_API_KEY = "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"
FILES_DIR = 'hujjatlar/shartnomalar'
EXCEL_PATH = 'data/Talabalar_Toliq_Royxati.xlsx'
OUT_DIR = 'qayta_tekshiruv'

GROUP = int(sys.argv[1]) if len(sys.argv) > 1 else 1
START = (GROUP - 1) * 30 + 1
END   = GROUP * 30

print(f"\n{'='*60}")
print(f"  GURUH {GROUP}: Talaba #{START} — #{END}")
print(f"{'='*60}\n")

def ai_vision_flash(img_bytes):
    """Flash model - tez, arzon"""
    b64 = base64.b64encode(img_bytes).decode('utf-8')
    payload = {
        "model": "google/gemini-2.5-flash",
        "messages": [{"role": "user", "content": [
            {"type": "text", "text": """O'zbekiston hujjatini o'qi. JSON qaytarish:
{
  "tur": "pasport|id|shahodatnoma|diplom",
  "pass_ser": "AE1234567",
  "pinfl": "14 xonali",
  "dob": "DD.MM.YYYY",
  "pass_ber": "DD.MM.YYYY (FAQAT pasport/ID-karta berilgan sanasi)",
  "ota": "... qizi / ... o'g'li",
  "sh_ser": "UM",
  "sh_num": "12345678",
  "maktab": "N-sonli maktab nomi",
  "yil": "YYYY"
}
Faqat JSON. Bo'sh bo'lsa null."""},
            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}
        ]}],
        "temperature": 0.05
    }
    try:
        req = urllib.request.Request(
            "https://openrouter.ai/api/v1/chat/completions",
            data=json.dumps(payload).encode('utf-8'),
            headers={"Authorization": f"Bearer {OPENROUTER_API_KEY}", "Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req, timeout=20) as resp:
            res = json.loads(resp.read().decode('utf-8'))
            txt = res['choices'][0]['message']['content'].strip()
            txt = re.sub(r'^```json\s*', '', txt); txt = re.sub(r'\s*```$', '', txt)
            return json.loads(txt)
    except Exception as e:
        print(f"    [AI xato]: {e}")
        return {}

def qr_cert_pdf(url):
    """QR → PDF → ma'lumot"""
    r = {}
    if not pypdf: return r
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as resp:
            txt = pypdf.PdfReader(io.BytesIO(resp.read())).pages[0].extract_text()
        m = re.search(r'UM\s*[№#\.\s]*(\d{7,8})', txt)
        if not m: m = re.search(r'\b(03\d{6}|01\d{6}|\d{8})\b', txt)
        if m: r['sh'] = f"UM {m.group(1)}"
        m = re.search(r'(\d{4})\s*-yilda', txt)
        if m: r['yil'] = m.group(1)
        m = re.search(r'(\d{1,3}\s*-\s*sonli\s+[\w\s\'\-]+(?:maktabini|maktabi|maktab))', txt, re.I)
        if not m: m = re.search(r'([\w\s\'\-]+(?:umumiy o\'rta ta\'lim maktabi|maktabini|maktab))', txt, re.I)
        if m:
            s = re.sub(r'^\d{4}\s*','', m.group(1).strip()); s = re.sub(r'ini$','i',s)
            r['maktab'] = s
    except: pass
    return r

def find_docx(ism, shnum, all_files):
    ic = ism.lower().replace("'","").replace("'","").replace("'","")
    parts = [p for p in ic.split() if len(p) >= 3]
    # 1. Familiya + shartnoma raqami
    if shnum and parts:
        for f in all_files:
            fc = f.lower().replace("'","").replace("'","")
            if fc.endswith('.docx') and parts[0] in fc:
                if fc.endswith(f" {shnum}.docx"):
                    return f
    # 2. Familiya + Ism
    if len(parts) >= 2:
        for f in all_files:
            fc = f.lower().replace("'","").replace("'","")
            if fc.endswith('.docx') and parts[0] in fc and parts[1] in fc:
                return f
    return None

# Excel
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active
all_files = os.listdir(FILES_DIR)

row_count = 0
done = 0

for r in range(2, ws.max_row + 1):
    ism = str(ws.cell(row=r, column=2).value or '').strip()
    if not ism: continue
    row_count += 1
    if row_count < START or row_count > END: continue

    shnum = str(ws.cell(row=r, column=5).value or '').strip()
    print(f"\n[{row_count:3d}/{END}] #{shnum} {ism}", flush=True)

    docx_file = find_docx(ism, shnum, all_files)
    if not docx_file:
        print(f"    [OGOHLANTIRISH] Fayl topilmadi")
        ws.cell(row=r, column=21, value="FAYL_YOQ")
        ws.cell(row=r, column=22, value=f"G{GROUP}: Fayl topilmadi")
        wb.save(EXCEL_PATH)
        continue

    print(f"    [FAYL] {docx_file}", flush=True)
    try:
        doc = docx.Document(os.path.join(FILES_DIR, docx_file))
    except:
        print("    [XATO] Ochilmadi"); continue

    img_blobs = [rel.target_part.blob for rel in doc.part.rels.values() if 'image' in rel.target_ref and len(rel.target_part.blob) > 8000]

    pass_val = pinfl = dob = pass_ber = ota = sh_val = sh_qr = maktab = yil = cert_tur = ''

    # === 1-BOSQICH: QR KOD (tez, AI yo'q) ===
    if zxingcpp:
        for b in img_blobs:
            im = Image.open(io.BytesIO(b)).convert('RGB')
            for angle in [0,90,180,270]:
                bcs = zxingcpp.read_barcodes(im.rotate(angle, expand=True))
                if bcs:
                    for bc in bcs:
                        t = bc.text
                        if 'shahodatnoma' in t.lower() or 'e-shahodatnoma' in t.lower():
                            sh_qr = t
                            pd = qr_cert_pdf(t)
                            if pd.get('sh'):     sh_val  = pd['sh']
                            if pd.get('maktab'): maktab  = pd['maktab']
                            if pd.get('yil'):    yil     = pd['yil']
                            print(f"    [QR] {sh_val} | {maktab} | {yil}", flush=True)
                        elif 'IUUZB' in t:
                            m = re.search(r'IUUZB([A-Z]{2}\d{7})\d?([3-6]\d{13})', t)
                            if m:
                                pass_val = m.group(1)
                                pinfl    = m.group(2)
                                p = m.group(2)
                                cent = '19' if p[0] in ['3','4'] else '20'
                                dob = f"{p[1:3]}.{p[3:5]}.{cent}{p[5:7]}"
                                print(f"    [PASPORT QR] {pass_val} | {pinfl} | {dob}", flush=True)
                    break

    # === 2-BOSQICH: AI (faqat kerak bo'lganda) ===
    if not pass_val or not sh_val or not ota:
        print(f"    [AI] Flash ishga tushmoqda...", flush=True)
        for b in img_blobs:
            if pass_val and sh_val and ota: break
            im = Image.open(io.BytesIO(b)).convert('RGB')
            buf = io.BytesIO()
            im.save(buf, format='JPEG', quality=85)
            ai = ai_vision_flash(buf.getvalue())
            if not ai: continue
            h = str(ai.get('tur') or '').lower()
            if 'pasport' in h or 'id' in h:
                if not pass_val and ai.get('pass_ser') and ai['pass_ser'] != 'null':
                    pass_val = str(ai['pass_ser']).replace(' ','').upper()
                if not pinfl and ai.get('pinfl') and ai['pinfl'] != 'null':
                    pinfl = str(ai['pinfl']).strip()
                if not dob and ai.get('dob') and ai['dob'] != 'null':
                    dob = str(ai['dob']).strip()
                if not pass_ber and ai.get('pass_ber') and ai['pass_ber'] != 'null':
                    pass_ber = str(ai['pass_ber']).strip()
                if not ota and ai.get('ota') and ai['ota'] != 'null':
                    ota = str(ai['ota']).strip()
                print(f"    AI pasport: {pass_val} | Ota: {ota} | Ber: {pass_ber}", flush=True)
            elif 'shahodatnoma' in h or 'diplom' in h:
                if not sh_val:
                    ss = str(ai.get('sh_ser') or '').strip()
                    sn = str(ai.get('sh_num') or '').strip()
                    if sn and sn != 'null':
                        sh_val = f"{ss} {sn}".strip()
                if not maktab and ai.get('maktab') and ai['maktab'] != 'null':
                    maktab = str(ai['maktab']).strip()
                if not yil and ai.get('yil') and ai['yil'] != 'null':
                    yil = str(ai['yil']).strip()
                if 'diplom' in h: cert_tur = 'Diplom'
                print(f"    AI hujjat: {sh_val} | {maktab} | {yil}", flush=True)
            time.sleep(0.5)

    # PINFL dan DOB hisoblash
    if not dob and pinfl and len(pinfl) == 14 and pinfl.isdigit():
        p = pinfl
        cent = '19' if p[0] in ['3','4'] else '20'
        dob = f"{p[1:3]}.{p[3:5]}.{cent}{p[5:7]}"

    # Excelga yozish (darhol)
    fish = f"{ism} {ota}".strip() if ota else ism
    ws.cell(row=r, column=7,  value=ota)
    ws.cell(row=r, column=8,  value=fish)
    # 9-ustun (Passport bo'yicha F.I.SH) — bu yerda yozilmaydi,
    # uni scripts/verify_passport_names.py boshqaradi.
    ws.cell(row=r, column=10, value=pass_val)
    ws.cell(row=r, column=11, value=pinfl)
    ws.cell(row=r, column=12, value=pass_ber)
    ws.cell(row=r, column=13, value=dob)
    ws.cell(row=r, column=15, value=sh_val)
    ws.cell(row=r, column=16, value=sh_qr)
    ws.cell(row=r, column=17, value=maktab)
    ws.cell(row=r, column=18, value="Umumiy o'rta maktab" if cert_tur != 'Diplom' else 'Kollej')
    ws.cell(row=r, column=19, value=yil)
    ws.cell(row=r, column=21, value="TOPILDI" if pass_val or sh_val else "CHALA")
    ws.cell(row=r, column=22, value=f"G{GROUP}: Qayta tahlil")
    wb.save(EXCEL_PATH)
    done += 1
    print(f"    [OK] Saqlandi ({done}/{min(30, END-START+1)})", flush=True)

print(f"\n{'='*60}")
print(f"  [OK] Guruh {GROUP} yakunlandi: {done} ta talaba tahlil qilindi")
print(f"{'='*60}")
