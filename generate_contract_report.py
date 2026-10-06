import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import re
import shutil
from difflib import SequenceMatcher
from collections import defaultdict

def clean_short(s):
    if not s: return []
    s = s.lower()
    s = re.sub(r'[\'ʻʼ`‘´’\u2018\u2019\u02bb\u02bc]', '', s)
    s = re.sub(r'[\s\-\.\,]+', ' ', s)
    s = s.replace('sh', 's').replace('ch', 'c').replace('dj', 'j')
    s = s.replace('o‘', 'o').replace('g‘', 'g').replace("o'", "o").replace("g'", "g")
    s = s.replace('o’', 'o').replace('g’', 'g')
    s = s.replace('ya', 'ia').replace('yu', 'iu').replace('yo', 'io').replace('ye', 'e')
    s = s.replace('h', 'x').replace('q', 'k').replace('p', 'f').replace('b', 'p')
    s = s.replace('d', 't').replace('z', 's').replace('v', 'w')
    words = []
    for w in s.split():
        for suf in ['yeva', 'yev', 'ova', 'ov', 'qizi', 'ogli', 'kizi']:
            if w.endswith(suf) and len(w) > len(suf) + 2:
                w = w[:-len(suf)]
                break
        words.append(w)
    return words[:2]

def create_report():
    # Load DB
    with open('data/students.json', 'r', encoding='utf-8') as f:
        db = json.load(f)

    for s in db:
        s_fio = s.get('fish') or f"{s.get('ism', '')} {s.get('ota', '')}".strip()
        s['fio'] = s_fio
        s['parts'] = clean_short(s_fio)

    # Load Contracts
    wb_in = openpyxl.load_workbook('Kontraktlar/29.09.2026_GACHA_KONTRAKTLAR.xlsx', data_only=True)

    # Parse KONTRAKTLAR sheet
    ws_k = wb_in['KONTRAKTLAR']
    kontrakt_rows = []
    for r in range(19, ws_k.max_row + 1):
        fio = ws_k.cell(row=r, column=3).value
        grp = ws_k.cell(row=r, column=1).value
        req = ws_k.cell(row=r, column=4).value or 0
        paid = ws_k.cell(row=r, column=5).value or 0
        debt = ws_k.cell(row=r, column=6).value or 0
        if fio and str(fio).strip().lower() != 'jami':
            fio_str = str(fio).strip()
            kontrakt_rows.append({
                'sheet': 'KONTRAKTLAR',
                'row': r,
                'fio': fio_str,
                'grp': str(grp).strip(),
                'req': float(req) if isinstance(req, (int, float)) else 0.0,
                'paid': float(paid) if isinstance(paid, (int, float)) else 0.0,
                'debt': float(debt) if isinstance(debt, (int, float)) else 0.0,
                'parts': clean_short(fio_str)
            })

    # Parse Akademik.TSCH
    ws_a = wb_in['Akademik.TSCH']
    akademik_rows = []
    for r in range(4, ws_a.max_row + 1):
        fio = ws_a.cell(row=r, column=3).value
        grp = ws_a.cell(row=r, column=1).value
        req = ws_a.cell(row=r, column=4).value or 0
        paid = ws_a.cell(row=r, column=5).value or 0
        debt = ws_a.cell(row=r, column=6).value or 0
        if fio and str(fio).strip().lower() != 'jami':
            fio_str = str(fio).strip()
            akademik_rows.append({
                'sheet': 'Akademik.TSCH',
                'row': r,
                'fio': fio_str,
                'grp': str(grp).strip(),
                'req': float(req) if isinstance(req, (int, float)) else 0.0,
                'paid': float(paid) if isinstance(paid, (int, float)) else 0.0,
                'debt': float(debt) if isinstance(debt, (int, float)) else 0.0,
                'parts': clean_short(fio_str)
            })

    # Bijective Matching
    matched_k_to_db = {}
    matched_db_to_k = {}
    used_db_rows = set()

    # Pass 1: exact parts match in same group
    for kr in kontrakt_rows:
        for s in db:
            if s['row'] in used_db_rows: continue
            if kr['parts'] == s['parts'] and kr['grp'] == s.get('group'):
                matched_k_to_db[kr['row']] = (s, 'Aniq moslik')
                matched_db_to_k[s['row']] = kr
                used_db_rows.add(s['row'])
                break

    # Pass 2: exact parts match in any group
    for kr in kontrakt_rows:
        if kr['row'] in matched_k_to_db: continue
        for s in db:
            if s['row'] in used_db_rows: continue
            if kr['parts'] == s['parts']:
                matched_k_to_db[kr['row']] = (s, 'Guruh farqi bilan moslik')
                matched_db_to_k[s['row']] = kr
                used_db_rows.add(s['row'])
                break

    # Pass 3: fuzzy in same group
    for kr in kontrakt_rows:
        if kr['row'] in matched_k_to_db: continue
        best_sim = 0
        best_s = None
        for s in db:
            if s['row'] in used_db_rows: continue
            if kr['grp'] == s.get('group'):
                sim = SequenceMatcher(None, ' '.join(kr['parts']), ' '.join(s['parts'])).ratio()
                if sim > best_sim and sim >= 0.70:
                    best_sim = sim
                    best_s = s
        if best_s:
            matched_k_to_db[kr['row']] = (best_s, f'O\'xshashlik ({best_sim:.2f})')
            matched_db_to_k[best_s['row']] = kr
            used_db_rows.add(best_s['row'])

    # Pass 4: fuzzy cross group
    for kr in kontrakt_rows:
        if kr['row'] in matched_k_to_db: continue
        best_sim = 0
        best_s = None
        for s in db:
            if s['row'] in used_db_rows: continue
            sim = SequenceMatcher(None, ' '.join(kr['parts']), ' '.join(s['parts'])).ratio()
            if sim > best_sim and sim >= 0.72:
                best_sim = sim
                best_s = s
        if best_s:
            matched_k_to_db[kr['row']] = (best_s, f'Kross-guruh o\'xshashlik ({best_sim:.2f})')
            matched_db_to_k[best_s['row']] = kr
            used_db_rows.add(best_s['row'])

    # Match Akademik.TSCH rows
    matched_a_to_db = {}
    matched_db_to_a = {}
    for ar in akademik_rows:
        for s in db:
            if s['row'] in matched_db_to_k or s['row'] in matched_db_to_a: continue
            if ar['parts'] == s['parts']:
                matched_a_to_db[ar['row']] = s
                matched_db_to_a[s['row']] = ar
                break
        if ar['row'] not in matched_a_to_db:
            best_sim = 0
            best_s = None
            for s in db:
                if s['row'] in matched_db_to_k or s['row'] in matched_db_to_a: continue
                sim = SequenceMatcher(None, ' '.join(ar['parts']), ' '.join(s['parts'])).ratio()
                if sim > best_sim and sim >= 0.70:
                    best_sim = sim
                    best_s = s
            if best_s:
                matched_a_to_db[ar['row']] = best_s
                matched_db_to_a[best_s['row']] = ar

    print(f"KONTRAKTLAR matched: {len(matched_k_to_db)} / {len(kontrakt_rows)}")
    print(f"Akademik.TSCH matched: {len(matched_a_to_db)} / {len(akademik_rows)}")

    # Build Excel Workbook
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    # Styles
    font_title = Font(name='Calibri', size=16, bold=True, color='1E293B')
    font_subtitle = Font(name='Calibri', size=11, italic=True, color='64748B')
    font_sec_header = Font(name='Calibri', size=13, bold=True, color='1E3A8A')
    font_tbl_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    font_bold = Font(name='Calibri', size=11, bold=True, color='1E293B')
    font_regular = Font(name='Calibri', size=11, color='334155')
    font_small = Font(name='Calibri', size=9, italic=True, color='64748B')
    
    font_success = Font(name='Calibri', size=11, bold=True, color='047857')
    font_warning = Font(name='Calibri', size=11, bold=True, color='B45309')
    font_danger = Font(name='Calibri', size=11, bold=True, color='B91C1C')
    font_info = Font(name='Calibri', size=11, bold=True, color='1D4ED8')

    fill_header = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_header_sub = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    fill_header_danger = PatternFill(start_color='B91C1C', end_color='B91C1C', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_total = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    fill_badge_green = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
    fill_badge_red = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    fill_badge_yellow = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    fill_badge_blue = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='CBD5E1')
    thick_bottom_side = Side(border_style='medium', color='1E3A8A')
    double_bottom_side = Side(border_style='double', color='1E293B')
    
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)
    total_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_bottom_side)

    align_left = Alignment(horizontal='left', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_header = Alignment(horizontal='center', vertical='center', wrap_text=True)

    NUM_CURRENCY = '#,##0 "so\'m"'
    NUM_INTEGER = '#,##0'
    NUM_PERCENT = '0.0%'

    # ==========================================
    # SHEET 1: Umumiy Tahlil (Dashboard)
    # ==========================================
    ws1 = wb_out.create_sheet(title="Umumiy Tahlil")
    ws1.views.sheetView[0].showGridLines = True

    ws1['B2'] = "TALABALAR RO'YXATI VA BUXGALTERIYA SHARTNOMALARI TAQQOSLASH HISOBOTI"
    ws1['B2'].font = font_title
    ws1['B3'] = "Buxgalteriya fayli: 29.09.2026_GACHA_KONTRAKTLAR.xlsx | Tahlil sanasi: 2026-yil 30-sentyabr"
    ws1['B3'].font = font_subtitle

    db_active_students = [s for s in db if s.get('group', '').startswith(('24-', '25-'))]
    db_1kurs_students = [s for s in db if s.get('group', '').startswith('26-')]

    kpi_cards = [
        ("Jami Talabalar (Bazada)", len(db), "Barcha guruhlar, akademik va chetlatilganlar", "1E3A8A"),
        ("Aktiv 2-3 Kurs Talabalari", len(db_active_students), "13 ta guruh bo'yicha", "2563EB"),
        ("Shartnoma Mavjud (Aktiv)", len(kontrakt_rows), "Buxgalteriya KONTRAKTLAR varag'ida", "0D9488"),
        ("Jami Shartnoma Qiymati", sum(kr['req'] for kr in kontrakt_rows), "2-3 kurs kontrakt summasi", "1E293B"),
        ("Jami To'langan Mablag'", sum(kr['paid'] for kr in kontrakt_rows), "Kassaga va bank orqali kelib tushgan", "16A34A"),
        ("Jami Haqiqiy Qarzdorlik", sum(kr['debt'] for kr in kontrakt_rows if kr['debt'] > 0), "Qarzdor talabalar bo'yicha", "DC2626"),
        ("Ortiqcha To'lov (Avans)", abs(sum(kr['debt'] for kr in kontrakt_rows if kr['debt'] < 0)), "Ortiqcha to'lagan talabalar bo'yicha", "D97706"),
        ("1-Kurs Yangi Talabalar", len(db_1kurs_students), "Hali kontrakt varag'iga kiritilmagan", "7C3AED"),
    ]

    row_kpi = 5
    for i, (title, val, desc, color_code) in enumerate(kpi_cards):
        col_start = 2 + (i % 4) * 2
        r_start = row_kpi if i < 4 else row_kpi + 4
        
        ws1.merge_cells(start_row=r_start, start_column=col_start, end_row=r_start, end_column=col_start+1)
        ws1.merge_cells(start_row=r_start+1, start_column=col_start, end_row=r_start+1, end_column=col_start+1)
        ws1.merge_cells(start_row=r_start+2, start_column=col_start, end_row=r_start+2, end_column=col_start+1)
        
        c_title = ws1.cell(row=r_start, column=col_start, value=title)
        c_title.font = Font(name='Calibri', size=10, bold=True, color='64748B')
        c_title.alignment = align_center

        c_val = ws1.cell(row=r_start+1, column=col_start, value=val)
        c_val.font = Font(name='Calibri', size=15, bold=True, color=color_code)
        c_val.alignment = align_center
        if isinstance(val, float) or val > 1000:
            if any(k in title for k in ["Summasi", "Qiymati", "Mablag'", "Qarzdorlik", "To'lov"]):
                c_val.number_format = NUM_CURRENCY
            else:
                c_val.number_format = NUM_INTEGER

        c_desc = ws1.cell(row=r_start+2, column=col_start, value=desc)
        c_desc.font = Font(name='Calibri', size=8, italic=True, color='94A3B8')
        c_desc.alignment = align_center

        card_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
        for r in range(r_start, r_start+3):
            for c in range(col_start, col_start+2):
                ws1.cell(row=r, column=c).fill = card_fill
                ws1.cell(row=r, column=c).border = cell_border

    tbl_start_row = 15
    ws1.cell(row=tbl_start_row, column=2, value="GURUHLAR KESIMIDA SHARTNOMA VA TO'LOVLAR TAHLILI").font = font_sec_header

    headers_ws1 = [
        ("T/r", 6, align_center),
        ("Guruh", 10, align_center),
        ("Kurs", 8, align_center),
        ("Yo'nalish", 25, align_left),
        ("Bazada Soni", 14, align_right),
        ("Shartnomada", 14, align_right),
        ("Farq", 10, align_right),
        ("Jami Shartnoma (so'm)", 24, align_right),
        ("Jami To'langan (so'm)", 24, align_right),
        ("Qarzdorlik (so'm)", 22, align_right),
        ("Ortiqcha To'lov (so'm)", 22, align_right),
        ("Qarzdorlar", 13, align_right),
        ("To'lov %", 12, align_right),
        ("Holat / Xulosa", 28, align_left)
    ]

    header_row = tbl_start_row + 2
    for col_idx, (h_name, width, align) in enumerate(headers_ws1, start=2):
        cell = ws1.cell(row=header_row, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = max(width + 2, 12)
    ws1.row_dimensions[header_row].height = 28

    active_groups = [
        ('24-11', '3-kurs', 'Hamshiralik ishi'),
        ('24-12', '3-kurs', 'Hamshiralik ishi'),
        ('24-13', '3-kurs', 'Hamshiralik ishi'),
        ('24-15', '3-kurs', 'Hamshiralik ishi'),
        ('24-16', '3-kurs', 'Hamshiralik ishi'),
        ('25-16', '2-kurs', 'Hamshiralik ishi'),
        ('25-17', '2-kurs', 'Hamshiralik ishi'),
        ('25-18', '2-kurs', 'Hamshiralik ishi'),
        ('25-19', '2-kurs', 'Hamshiralik ishi'),
        ('25-20', '2-kurs', 'Hamshiralik ishi'),
        ('25-21', '2-kurs', 'Hamshiralik ishi'),
        ('25-22', '2-kurs', 'Davolash ishi'),
        ('25-23', '2-kurs', 'Davolash ishi'),
    ]

    cur_row = header_row + 1
    for tr, (grp, kurs, yon) in enumerate(active_groups, start=1):
        db_cnt = sum(1 for s in db_active_students if s.get('group') == grp)
        k_grp_recs = [kr for kr in kontrakt_rows if kr['grp'] == grp]
        k_cnt = len(k_grp_recs)
        diff = db_cnt - k_cnt
        req_sum = sum(kr['req'] for kr in k_grp_recs)
        paid_sum = sum(kr['paid'] for kr in k_grp_recs)
        pos_debt = sum(kr['debt'] for kr in k_grp_recs if kr['debt'] > 0)
        neg_debt = abs(sum(kr['debt'] for kr in k_grp_recs if kr['debt'] < 0))
        debtor_cnt = sum(1 for kr in k_grp_recs if kr['debt'] > 0)
        pay_pct = (paid_sum / req_sum) if req_sum > 0 else 0.0

        if diff == 0:
            status_text = "To'liq mos keladi"
        elif diff > 0:
            status_text = f"Bazada {diff} ta ko'p (shartnomada yo'q)"
        else:
            status_text = f"Shartnomada {abs(diff)} ta ko'p"

        vals = [
            (tr, align_center, NUM_INTEGER, font_regular),
            (grp, align_center, None, font_bold),
            (kurs, align_center, None, font_regular),
            (yon, align_left, None, font_regular),
            (db_cnt, align_right, NUM_INTEGER, font_bold),
            (k_cnt, align_right, NUM_INTEGER, font_bold),
            (diff, align_right, NUM_INTEGER, font_danger if diff != 0 else font_success),
            (req_sum, align_right, NUM_CURRENCY, font_regular),
            (paid_sum, align_right, NUM_CURRENCY, font_regular),
            (pos_debt, align_right, NUM_CURRENCY, font_danger if pos_debt > 0 else font_success),
            (neg_debt, align_right, NUM_CURRENCY, font_regular),
            (debtor_cnt, align_right, NUM_INTEGER, font_warning if debtor_cnt > 0 else font_regular),
            (pay_pct, align_right, NUM_PERCENT, font_bold),
            (status_text, align_left, None, font_regular)
        ]

        row_fill = fill_zebra if tr % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
            cell = ws1.cell(row=cur_row, column=col_idx, value=val)
            cell.alignment = align
            cell.font = f_style
            cell.border = cell_border
            if row_fill.fill_type: cell.fill = row_fill
            if num_fmt: cell.number_format = num_fmt

        ws1.row_dimensions[cur_row].height = 20
        cur_row += 1

    tot_row = cur_row
    ws1.merge_cells(start_row=tot_row, start_column=2, end_row=tot_row, end_column=5)
    ws1.cell(row=tot_row, column=2, value="JAMI (2-3 KURS)").alignment = align_center
    ws1.cell(row=tot_row, column=2).font = font_bold
    ws1.cell(row=tot_row, column=6, value=f"=SUM(F{header_row+1}:F{tot_row-1})").number_format = NUM_INTEGER
    ws1.cell(row=tot_row, column=7, value=f"=SUM(G{header_row+1}:G{tot_row-1})").number_format = NUM_INTEGER
    ws1.cell(row=tot_row, column=8, value=f"=SUM(H{header_row+1}:H{tot_row-1})").number_format = NUM_INTEGER
    ws1.cell(row=tot_row, column=9, value=f"=SUM(I{header_row+1}:I{tot_row-1})").number_format = NUM_CURRENCY
    ws1.cell(row=tot_row, column=10, value=f"=SUM(J{header_row+1}:J{tot_row-1})").number_format = NUM_CURRENCY
    ws1.cell(row=tot_row, column=11, value=f"=SUM(K{header_row+1}:K{tot_row-1})").number_format = NUM_CURRENCY
    ws1.cell(row=tot_row, column=12, value=f"=SUM(L{header_row+1}:L{tot_row-1})").number_format = NUM_CURRENCY
    ws1.cell(row=tot_row, column=13, value=f"=SUM(M{header_row+1}:M{tot_row-1})").number_format = NUM_INTEGER
    ws1.cell(row=tot_row, column=14, value=f"=J{tot_row}/I{tot_row}").number_format = NUM_PERCENT

    for c in range(2, len(headers_ws1)+2):
        cell = ws1.cell(row=tot_row, column=c)
        cell.font = font_bold
        cell.fill = fill_total
        cell.border = total_border
        if cell.alignment.horizontal is None:
            cell.alignment = align_right
    ws1.row_dimensions[tot_row].height = 24
    ws1.freeze_panes = 'C18'


    # ==========================================
    # SHEET 2: Farqlar va Muammolar (Discrepancies)
    # ==========================================
    ws2 = wb_out.create_sheet(title="Farqlar va Muammolar")
    ws2.views.sheetView[0].showGridLines = True

    ws2['B2'] = "BUXGALTERIYA SHARTNOMALARI VA TALABALAR BAZASI O'RTASIDAGI FARQLAR"
    ws2['B2'].font = font_title
    ws2['B3'] = "Mazkur sahifada zudlik bilan e'tibor qaratilishi va to'g'irlanishi lozim bo'lgan barcha nomuvofiqliklar keltirilgan."
    ws2['B3'].font = font_subtitle

    # Section 1: Active students without contract (Dynamic)
    unmatched_active_students = [
        s for s in db_active_students
        if s['row'] not in matched_db_to_k and s['row'] not in matched_db_to_a
    ]
    unmatched_active_students.sort(key=lambda x: (x.get('group', ''), x.get('fish', '')))

    r2 = 5
    ws2.cell(row=r2, column=2, value=f"1. AKTIV 2-3 KURS TALABALARI: SHARTNOMALAR VARAG'IDA YO'Q ({len(unmatched_active_students)} NAFAR)").font = font_sec_header

    headers_s1 = [
        ("T/r", 6, align_center),
        ("Baza ID", 10, align_center),
        ("F.I.Sh (Bazada)", 32, align_left),
        ("Guruhi", 10, align_center),
        ("PINFL", 18, align_center),
        ("Telefon", 16, align_center),
        ("Baza (HEMIS)", 15, align_center),
        ("Buxgalteriya faylidagi holati", 35, align_left),
        ("Tavsiya / Amal", 40, align_left)
    ]
    r2 += 2
    for col_idx, (h_name, width, align) in enumerate(headers_s1, start=2):
        cell = ws2.cell(row=r2, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header_danger
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = max(ws2.column_dimensions[col_letter].width or 0, width + 2)
    ws2.row_dimensions[r2].height = 25

    r2 += 1
    for tr, s in enumerate(unmatched_active_students, start=1):
        fio = s.get('fish') or s.get('ism', '')
        b_baza = s.get('baza', 'KIRITILDI')
        if s['row'] == 424:
            k_stat = "Eski '1 chi kurs' varag'ida (114-guruh) qolib ketgan"
            rec = "Faol 'KONTRAKTLAR' varag'iga 25-19 guruhga o'tkazish"
        elif s['row'] in (558, 559, 560):
            k_stat = "Yangi qo'shilgan (shartnoma ochilmagan)"
            rec = "Buxgalteriyada yangi 25-16 shartnoma kartasi ochish"
        elif s['row'] == 563:
            k_stat = "Yangi kiritilgan (buxgalteriyada shartnoma ochilmagan)"
            rec = "Buxgalteriyada 25-17 shartnoma ochish va HEMIS bazasiga kiritish"
        else:
            k_stat = "Shartnoma faylida umuman topilmadi"
            rec = "Buxgalteriyadan shartnoma holati va to'lovlarini aniqlash"

        vals = [
            (tr, align_center, NUM_INTEGER, font_regular),
            (s['row'], align_center, NUM_INTEGER, font_bold),
            (fio, align_left, None, font_bold),
            (s.get('group', ''), align_center, None, font_bold),
            (s.get('pinfl', ''), align_center, None, font_regular),
            (s.get('tel', ''), align_center, None, font_regular),
            (b_baza, align_center, None, font_danger if b_baza == 'KIRITILMAGAN' else font_success),
            (k_stat, align_left, None, font_danger),
            (rec, align_left, None, font_regular)
        ]
        row_fill = fill_badge_red if "umuman topilmadi" in k_stat or b_baza == 'KIRITILMAGAN' else fill_badge_yellow
        for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
            cell = ws2.cell(row=r2, column=col_idx, value=val)
            cell.alignment = align
            cell.font = f_style
            cell.border = cell_border
            cell.fill = row_fill
            if num_fmt: cell.number_format = num_fmt
        ws2.row_dimensions[r2].height = 22
        r2 += 1

    # Section 2: Active in DB but in Akademik.TSCH in contract (Dynamic)
    active_in_tsch = [
        s for s in db_active_students
        if s['row'] in matched_db_to_a
    ]
    r2 += 2
    ws2.cell(row=r2, column=2, value=f"2. BAZADA AKTIV, LEKIN KONTRAKTDA 'AKADEMIK.TSCH' VARAG'IDA BO'LGANLAR ({len(active_in_tsch)} NAFAR)").font = font_sec_header

    headers_s2 = [
        ("T/r", 6, align_center),
        ("Baza ID", 10, align_center),
        ("F.I.Sh (Bazada)", 32, align_left),
        ("Bazadagi Guruh", 14, align_center),
        ("TSCH dagi Guruh", 14, align_center),
        ("TSCH Summasi", 18, align_right),
        ("TSCH To'langan", 18, align_right),
        ("TSCH Qarzi", 16, align_right),
        ("Tavsiya / Izoh", 40, align_left)
    ]
    r2 += 2
    for col_idx, (h_name, width, align) in enumerate(headers_s2, start=2):
        cell = ws2.cell(row=r2, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header_sub
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = max(ws2.column_dimensions[col_letter].width or 0, width + 2)
    ws2.row_dimensions[r2].height = 25

    if active_in_tsch:
        r2 += 1
        for tr, s in enumerate(active_in_tsch, start=1):
            ar = matched_db_to_a[s['row']]
            fio = s.get('fish') or s.get('ism', '')
            vals = [
                (tr, align_center, NUM_INTEGER, font_regular),
                (s['row'], align_center, NUM_INTEGER, font_bold),
                (fio, align_left, None, font_bold),
                (s.get('group', ''), align_center, None, font_bold),
                (ar['grp'], align_center, None, font_regular),
                (ar['req'], align_right, NUM_CURRENCY, font_regular),
                (ar['paid'], align_right, NUM_CURRENCY, font_regular),
                (ar['debt'], align_right, NUM_CURRENCY, font_regular),
                ("Talaba darslarga qaytgan bo'lsa, 'KONTRAKTLAR' varag'iga qaytarish lozim", align_left, None, font_regular)
            ]
            for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
                cell = ws2.cell(row=r2, column=col_idx, value=val)
                cell.alignment = align
                cell.font = f_style
                cell.border = cell_border
                cell.fill = fill_badge_yellow
                if num_fmt: cell.number_format = num_fmt
            ws2.row_dimensions[r2].height = 22
            r2 += 1
    else:
        r2 += 1
        ws2.merge_cells(start_row=r2, start_column=2, end_row=r2, end_column=10)
        c_ok = ws2.cell(row=r2, column=2, value="Bunday nomuvofiqlik mavjud emas! Barcha faol talabalar shartnoma varag'iga to'g'ri o'tkazilgan.")
        c_ok.font = font_success
        c_ok.alignment = align_center
        c_ok.fill = fill_badge_green
        for col_idx in range(2, 11):
            ws2.cell(row=r2, column=col_idx).border = cell_border
        ws2.row_dimensions[r2].height = 24

    # Section 3: In Kontrakt active sheet, but in DB Akademik (Dynamic)
    akademik_in_k = [
        s for s in db
        if s.get('group') == "Akademik ta'til olganlar" and s['row'] in matched_db_to_k
    ]
    r2 += 2
    ws2.cell(row=r2, column=2, value=f"3. BUXGALTERIYADA AKTIV KONTRAKTDA, LEKIN BAZADA 'AKADEMIK TA'TIL'DA ({len(akademik_in_k)} NAFAR)").font = font_sec_header

    headers_s3 = [
        ("T/r", 6, align_center),
        ("F.I.Sh (Bazada)", 32, align_left),
        ("Bazadagi Holat", 20, align_center),
        ("Kontrakt Guruhi", 14, align_center),
        ("Kontrakt Qatori", 14, align_center),
        ("Talab qilingan", 18, align_right),
        ("To'langan", 18, align_right),
        ("Qoldiq Qarzi", 16, align_right),
        ("Tavsiya / Izoh", 40, align_left)
    ]
    r2 += 2
    for col_idx, (h_name, width, align) in enumerate(headers_s3, start=2):
        cell = ws2.cell(row=r2, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header_sub
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = max(ws2.column_dimensions[col_letter].width or 0, width + 2)
    ws2.row_dimensions[r2].height = 25

    if akademik_in_k:
        r2 += 1
        for tr, s in enumerate(akademik_in_k, start=1):
            kr = matched_db_to_k[s['row']]
            fio = s.get('fish') or s.get('ism', '')
            vals = [
                (tr, align_center, NUM_INTEGER, font_regular),
                (fio, align_left, None, font_bold),
                ("Akademik ta'tilda", align_center, None, font_warning),
                (kr['grp'], align_center, None, font_bold),
                (f"Qator {kr['row']}", align_center, None, font_regular),
                (kr['req'], align_right, NUM_CURRENCY, font_regular),
                (kr['paid'], align_right, NUM_CURRENCY, font_regular),
                (kr['debt'], align_right, NUM_CURRENCY, font_danger),
                ("Buxgalteriya varag'ida 'Akademik.TSCH' ro'yxatiga o'tkazish kerak", align_left, None, font_regular)
            ]
            for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
                cell = ws2.cell(row=r2, column=col_idx, value=val)
                cell.alignment = align
                cell.font = f_style
                cell.border = cell_border
                cell.fill = fill_badge_yellow
                if num_fmt: cell.number_format = num_fmt
            ws2.row_dimensions[r2].height = 22
            r2 += 1
    else:
        r2 += 1
        ws2.merge_cells(start_row=r2, start_column=2, end_row=r2, end_column=10)
        c_ok = ws2.cell(row=r2, column=2, value="Bunday nomuvofiqlik bartaraf etilgan! Akademik ta'tildagi talabalar 'Akademik.TSCH' varag'iga to'g'ri o'tkazilgan.")
        c_ok.font = font_success
        c_ok.alignment = align_center
        c_ok.fill = fill_badge_green
        for col_idx in range(2, 11):
            ws2.cell(row=r2, column=col_idx).border = cell_border
        ws2.row_dimensions[r2].height = 24

    # Section 4: Group mismatches between DB and Kontrakt (Dynamic)
    group_mismatches = []
    for s in db_active_students:
        if s['row'] in matched_db_to_k:
            kr = matched_db_to_k[s['row']]
            if kr['grp'] != s.get('group'):
                group_mismatches.append((s, kr))

    r2 += 2
    ws2.cell(row=r2, column=2, value=f"4. GURUH NOMUVOFIQLIKLARI: BAZADA BOSHQA, SHARTNOMADA BOSHQA GURUH ({len(group_mismatches)} NAFAR)").font = font_sec_header

    headers_s4 = [
        ("T/r", 6, align_center),
        ("F.I.Sh (Bazada)", 32, align_left),
        ("Bazadagi Guruhi", 15, align_center),
        ("Shartnomadagi Guruhi", 18, align_center),
        ("Shartnoma Summasi", 18, align_right),
        ("To'langan", 18, align_right),
        ("Qarzdorlik", 16, align_right),
        ("Holati va Tavsiya", 40, align_left)
    ]
    r2 += 2
    for col_idx, (h_name, width, align) in enumerate(headers_s4, start=2):
        cell = ws2.cell(row=r2, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = max(ws2.column_dimensions[col_letter].width or 0, width + 2)
    ws2.row_dimensions[r2].height = 25

    if group_mismatches:
        r2 += 1
        for tr, (s, kr) in enumerate(group_mismatches, start=1):
            fio = s.get('fish') or s.get('ism', '')
            b_grp = s.get('group', '')
            k_grp = kr['grp']
            rec = f"Talaba {b_grp} guruhga ko'chirilgan, buxgalteriyada {k_grp} guruhda turibdi"
            vals = [
                (tr, align_center, NUM_INTEGER, font_regular),
                (fio, align_left, None, font_bold),
                (b_grp, align_center, None, font_bold),
                (k_grp, align_center, None, font_warning),
                (kr['req'], align_right, NUM_CURRENCY, font_regular),
                (kr['paid'], align_right, NUM_CURRENCY, font_regular),
                (kr['debt'], align_right, NUM_CURRENCY, font_regular),
                (rec, align_left, None, font_regular)
            ]
            for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
                cell = ws2.cell(row=r2, column=col_idx, value=val)
                cell.alignment = align
                cell.font = f_style
                cell.border = cell_border
                cell.fill = fill_badge_blue
                if num_fmt: cell.number_format = num_fmt
            ws2.row_dimensions[r2].height = 22
            r2 += 1
    else:
        r2 += 1
        ws2.merge_cells(start_row=r2, start_column=2, end_row=r2, end_column=9)
        c_ok = ws2.cell(row=r2, column=2, value="Guruh nomuvofiqliklari mavjud emas! Barcha talabalar o'z guruhlari bilan 100% to'liq mos keladi.")
        c_ok.font = font_success
        c_ok.alignment = align_center
        c_ok.fill = fill_badge_green
        for col_idx in range(2, 10):
            ws2.cell(row=r2, column=col_idx).border = cell_border
        ws2.row_dimensions[r2].height = 24


    # ==========================================
    # SHEET 3: Aktiv Talabalar (Batafsil) - 326 students
    # ==========================================
    ws3 = wb_out.create_sheet(title="Aktiv Talabalar (Batafsil)")
    ws3.views.sheetView[0].showGridLines = True

    ws3['B2'] = "2-3 KURS AKTIV TALABALARNING TO'LIQ SHARTNOMA VA TO'LOV HISOBOTI"
    ws3['B2'].font = font_title
    ws3['B3'] = f"Ushbu ro'yxatda bazadagi barcha {len(db_active_students)} nafar faol talabaning shartnoma ko'rsatkichlari aks ettirilgan."
    ws3['B3'].font = font_subtitle

    headers_ws3 = [
        ("T/r", 6, align_center),
        ("Baza ID", 9, align_center),
        ("F.I.Sh (Talaba)", 34, align_left),
        ("Guruhi (Baza)", 12, align_center),
        ("Guruhi (Kontrakt)", 14, align_center),
        ("PINFL", 16, align_center),
        ("Telefon", 16, align_center),
        ("Moslik Holati", 22, align_left),
        ("Shartnoma Summasi", 20, align_right),
        ("To'langan Summa", 20, align_right),
        ("Qoldiq Qarzdorlik", 20, align_right),
        ("To'lov %", 10, align_right),
        ("Moliyaviy Holat", 18, align_center),
        ("Manba Varag'i & Qator", 22, align_left)
    ]

    r3 = 5
    for col_idx, (h_name, width, align) in enumerate(headers_ws3, start=2):
        cell = ws3.cell(row=r3, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws3.column_dimensions[col_letter].width = max(width + 2, 12)
    ws3.row_dimensions[r3].height = 26

    db_active_sorted = sorted(db_active_students, key=lambda x: (x.get('group', ''), x.get('fish', '')))

    r3 += 1
    for tr, s in enumerate(db_active_sorted, start=1):
        s_row = s['row']
        b_grp = s.get('group', '')
        pinfl = s.get('pinfl', '')
        tel = s.get('tel', '')
        fio = s.get('fish') or s.get('ism', '')

        if s_row in matched_db_to_k:
            kr = matched_db_to_k[s_row]
            k_grp = kr['grp']
            req_s = kr['req']
            paid_s = kr['paid']
            debt_s = kr['debt']
            source_loc = f"KONTRAKTLAR (qator {kr['row']})"
            
            if k_grp == b_grp:
                m_stat = "To'liq mos keladi"
            else:
                m_stat = f"Guruh farqi ({k_grp})"

            pay_pct = (paid_s / req_s) if req_s > 0 else (1.0 if paid_s >= req_s else 0.0)

            if debt_s > 0:
                fin_stat = "Qarzdor"
                fin_fill = fill_badge_red
                fin_font = font_danger
            elif debt_s < 0:
                fin_stat = "Ortiqcha to'lov"
                fin_fill = fill_badge_blue
                fin_font = font_info
            else:
                fin_stat = "To'liq to'langan"
                fin_fill = fill_badge_green
                fin_font = font_success

        elif s_row in matched_db_to_a:
            ar = matched_db_to_a[s_row]
            k_grp = ar['grp']
            req_s = ar['req']
            paid_s = ar['paid']
            debt_s = ar['debt']
            source_loc = f"Akademik.TSCH (qator {ar['row']})"
            m_stat = "TSCH varag'ida!"
            pay_pct = (paid_s / req_s) if req_s > 0 else 0.0
            fin_stat = "TSCH ro'yxatida"
            fin_fill = fill_badge_yellow
            fin_font = font_warning

        else:
            k_grp = "-"
            req_s = 0.0
            paid_s = 0.0
            debt_s = 0.0
            source_loc = "-"
            m_stat = "Shartnomasi yo'q"
            pay_pct = 0.0
            fin_stat = "Shartnoma kiritilmagan"
            fin_fill = fill_badge_red
            fin_font = font_danger

        vals = [
            (tr, align_center, NUM_INTEGER, font_regular),
            (s_row, align_center, NUM_INTEGER, font_regular),
            (fio, align_left, None, font_bold),
            (b_grp, align_center, None, font_bold),
            (k_grp, align_center, None, font_regular),
            (pinfl, align_center, None, font_regular),
            (tel, align_center, None, font_regular),
            (m_stat, align_left, None, font_warning if "!" in m_stat or "farqi" in m_stat else (font_danger if "yo'q" in m_stat else font_regular)),
            (req_s, align_right, NUM_CURRENCY, font_regular),
            (paid_s, align_right, NUM_CURRENCY, font_regular),
            (debt_s, align_right, NUM_CURRENCY, fin_font),
            (pay_pct, align_right, NUM_PERCENT, font_bold),
            (fin_stat, align_center, None, fin_font),
            (source_loc, align_left, None, font_small)
        ]

        row_fill = fill_zebra if tr % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
            cell = ws3.cell(row=r3, column=col_idx, value=val)
            cell.alignment = align
            cell.font = f_style
            cell.border = cell_border
            if col_idx == 14:
                cell.fill = fin_fill
            elif row_fill.fill_type:
                cell.fill = row_fill
            if num_fmt: cell.number_format = num_fmt

        ws3.row_dimensions[r3].height = 20
        r3 += 1

    ws3.merge_cells(start_row=r3, start_column=2, end_row=r3, end_column=9)
    ws3.cell(row=r3, column=2, value=f"JAMI ({len(db_active_students)} TALABA)").alignment = align_center
    ws3.cell(row=r3, column=2).font = font_bold
    ws3.cell(row=r3, column=10, value=f"=SUM(J6:J{r3-1})").number_format = NUM_CURRENCY
    ws3.cell(row=r3, column=11, value=f"=SUM(K6:K{r3-1})").number_format = NUM_CURRENCY
    ws3.cell(row=r3, column=12, value=f"=SUM(L6:L{r3-1})").number_format = NUM_CURRENCY
    ws3.cell(row=r3, column=13, value=f"=K{r3}/J{r3}").number_format = NUM_PERCENT

    for c in range(2, len(headers_ws3)+2):
        cell = ws3.cell(row=r3, column=c)
        cell.font = font_bold
        cell.fill = fill_total
        cell.border = total_border
        if cell.alignment.horizontal is None: cell.alignment = align_right
    ws3.row_dimensions[r3].height = 24
    ws3.freeze_panes = 'E6'


    # ==========================================
    # SHEET 4: 1-Kurs (Yangi Qabul) - 181 students
    # ==========================================
    ws4 = wb_out.create_sheet(title="1-Kurs (Yangi Qabul)")
    ws4.views.sheetView[0].showGridLines = True

    ws4['B2'] = "1-KURS TALABALARI RO'YXATI (2026-2027 O'QUV YILI YANGI QABUL)"
    ws4['B2'].font = font_title
    ws4['B3'] = "Bu talabalar yangi qabul qilingan bo'lib (181 nafar), buxgalteriya shartnomalar bazasiga hali kiritilmagan."
    ws4['B3'].font = font_subtitle

    headers_ws4 = [
        ("T/r", 6, align_center),
        ("Baza ID", 10, align_center),
        ("F.I.Sh (Talaba)", 32, align_left),
        ("Guruhi", 10, align_center),
        ("PINFL", 18, align_center),
        ("Telefon Raqami", 16, align_center),
        ("Hujjat Raqami (Sh. raqam)", 18, align_center),
        ("Shartnoma Holati", 35, align_left),
        ("Tavsiya / Amal", 40, align_left)
    ]

    r4 = 5
    for col_idx, (h_name, width, align) in enumerate(headers_ws4, start=2):
        cell = ws4.cell(row=r4, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws4.column_dimensions[col_letter].width = max(ws4.column_dimensions[col_letter].width or 0, width + 2)
    ws4.row_dimensions[r4].height = 26

    db_1kurs_sorted = sorted(db_1kurs_students, key=lambda x: (x.get('group', ''), x.get('fish', '')))

    r4 += 1
    for tr, s in enumerate(db_1kurs_sorted, start=1):
        vals = [
            (tr, align_center, NUM_INTEGER, font_regular),
            (s['row'], align_center, NUM_INTEGER, font_regular),
            (s.get('fish') or s.get('ism', ''), align_left, None, font_bold),
            (s.get('group', ''), align_center, None, font_bold),
            (s.get('pinfl', ''), align_center, None, font_regular),
            (s.get('tel', ''), align_center, None, font_regular),
            (s.get('shnum', '') or "-", align_center, None, font_regular),
            ("Buxgalteriyaga kiritilishi kutilmoqda", align_left, None, font_warning),
            ("2026-2027 o'quv yili shartnomasi rasmiylashtirilishi kerak", align_left, None, font_regular)
        ]
        row_fill = fill_zebra if tr % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
            cell = ws4.cell(row=r4, column=col_idx, value=val)
            cell.alignment = align
            cell.font = f_style
            cell.border = cell_border
            if row_fill.fill_type: cell.fill = row_fill
            if num_fmt: cell.number_format = num_fmt
        ws4.row_dimensions[r4].height = 20
        r4 += 1
    ws4.freeze_panes = 'E6'


    # ==========================================
    # SHEET 5: Qarzdorlar Ro'yxati (Debtors)
    # ==========================================
    ws5 = wb_out.create_sheet(title="Qarzdorlar Ro'yxati")
    ws5.views.sheetView[0].showGridLines = True

    debtor_records = [kr for kr in kontrakt_rows if kr['debt'] > 0]
    debtor_records.sort(key=lambda x: x['debt'], reverse=True)
    total_debt_amount = sum(kr['debt'] for kr in debtor_records)

    ws5['B2'] = "SHARTNOMA TO'LOVIDAN QARZDOR TALABALAR RO'YXATI"
    ws5['B2'].font = font_title
    ws5['B3'] = f"Jami qarzdorlar soni: {len(debtor_records)} nafar | Jami qarzdorlik summasi: {total_debt_amount:,.0f} so'm"
    ws5['B3'].font = font_subtitle

    headers_ws5 = [
        ("T/r", 6, align_center),
        ("Guruhi", 10, align_center),
        ("F.I.Sh (Talaba)", 32, align_left),
        ("PINFL", 16, align_center),
        ("Telefon", 16, align_center),
        ("Shartnoma Summasi", 22, align_right),
        ("To'langan Summa", 22, align_right),
        ("Qoldiq Qarzdorlik", 22, align_right),
        ("Qarz Ulushi %", 14, align_right),
        ("Kontrakt Faylidagi Qator", 22, align_center)
    ]

    r5 = 5
    for col_idx, (h_name, width, align) in enumerate(headers_ws5, start=2):
        cell = ws5.cell(row=r5, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header_danger
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws5.column_dimensions[col_letter].width = max(ws5.column_dimensions[col_letter].width or 0, width + 2)
    ws5.row_dimensions[r5].height = 26

    r5 += 1
    for tr, kr in enumerate(debtor_records, start=1):
        s_match = matched_k_to_db.get(kr['row'])
        s = s_match[0] if s_match else None
        pinfl = s.get('pinfl', '') if s else ''
        tel = s.get('tel', '') if s else ''
        debt_share = (kr['debt'] / kr['req']) if kr['req'] > 0 else 0.0

        vals = [
            (tr, align_center, NUM_INTEGER, font_regular),
            (kr['grp'], align_center, None, font_bold),
            (kr['fio'], align_left, None, font_bold),
            (pinfl, align_center, None, font_regular),
            (tel, align_center, None, font_regular),
            (kr['req'], align_right, NUM_CURRENCY, font_regular),
            (kr['paid'], align_right, NUM_CURRENCY, font_regular),
            (kr['debt'], align_right, NUM_CURRENCY, font_danger),
            (debt_share, align_right, NUM_PERCENT, font_danger),
            (f"Qator {kr['row']}", align_center, None, font_small)
        ]
        row_fill = fill_zebra if tr % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
            cell = ws5.cell(row=r5, column=col_idx, value=val)
            cell.alignment = align
            cell.font = f_style
            cell.border = cell_border
            if row_fill.fill_type: cell.fill = row_fill
            if num_fmt: cell.number_format = num_fmt
        ws5.row_dimensions[r5].height = 20
        r5 += 1

    ws5.merge_cells(start_row=r5, start_column=2, end_row=r5, end_column=6)
    ws5.cell(row=r5, column=2, value=f"JAMI QARZDORLIK ({len(debtor_records)} TALABA)").alignment = align_center
    ws5.cell(row=r5, column=2).font = font_bold
    ws5.cell(row=r5, column=7, value=f"=SUM(G6:G{r5-1})").number_format = NUM_CURRENCY
    ws5.cell(row=r5, column=8, value=f"=SUM(H6:H{r5-1})").number_format = NUM_CURRENCY
    ws5.cell(row=r5, column=9, value=f"=SUM(I6:I{r5-1})").number_format = NUM_CURRENCY
    ws5.cell(row=r5, column=10, value=f"=I{r5}/G{r5}").number_format = NUM_PERCENT

    for c in range(2, len(headers_ws5)+2):
        cell = ws5.cell(row=r5, column=c)
        cell.font = font_bold
        cell.fill = fill_total
        cell.border = total_border
        if cell.alignment.horizontal is None: cell.alignment = align_right
    ws5.row_dimensions[r5].height = 24
    ws5.freeze_panes = 'E6'


    # ==========================================
    # SHEET 6: Ortiqcha To'lovlar (Overpayments)
    # ==========================================
    ws6 = wb_out.create_sheet(title="Ortiqcha To'lovlar (Avans)")
    ws6.views.sheetView[0].showGridLines = True

    overpay_records = [kr for kr in kontrakt_rows if kr['debt'] < 0]
    overpay_records.sort(key=lambda x: x['debt'])
    total_overpay_amount = abs(sum(kr['debt'] for kr in overpay_records))

    ws6['B2'] = "SHARTNOMADAN ORTIQCHA TO'LOV QILGAN TALABALAR (AVANS)"
    ws6['B2'].font = font_title
    ws6['B3'] = f"Jami ortiqcha to'lov qilganlar: {len(overpay_records)} nafar | Jami avans summasi: {total_overpay_amount:,.0f} so'm"
    ws6['B3'].font = font_subtitle

    headers_ws6 = [
        ("T/r", 6, align_center),
        ("Guruhi", 10, align_center),
        ("F.I.Sh (Talaba)", 32, align_left),
        ("PINFL", 16, align_center),
        ("Telefon", 16, align_center),
        ("Shartnoma Summasi", 22, align_right),
        ("To'langan Summa", 22, align_right),
        ("Ortiqcha To'lov (Avans)", 22, align_right),
        ("To'lov Koeffitsiyenti", 18, align_right),
        ("Kontrakt Faylidagi Qator", 22, align_center)
    ]

    r6 = 5
    for col_idx, (h_name, width, align) in enumerate(headers_ws6, start=2):
        cell = ws6.cell(row=r6, column=col_idx, value=h_name)
        cell.font = font_tbl_header
        cell.fill = fill_header_sub
        cell.alignment = align_header
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws6.column_dimensions[col_letter].width = max(ws6.column_dimensions[col_letter].width or 0, width + 2)
    ws6.row_dimensions[r6].height = 26

    r6 += 1
    for tr, kr in enumerate(overpay_records, start=1):
        s_match = matched_k_to_db.get(kr['row'])
        s = s_match[0] if s_match else None
        pinfl = s.get('pinfl', '') if s else ''
        tel = s.get('tel', '') if s else ''
        overpay_amount = abs(kr['debt'])
        pay_coeff = (kr['paid'] / kr['req']) if kr['req'] > 0 else 0.0

        vals = [
            (tr, align_center, NUM_INTEGER, font_regular),
            (kr['grp'], align_center, None, font_bold),
            (kr['fio'], align_left, None, font_bold),
            (pinfl, align_center, None, font_regular),
            (tel, align_center, None, font_regular),
            (kr['req'], align_right, NUM_CURRENCY, font_regular),
            (kr['paid'], align_right, NUM_CURRENCY, font_regular),
            (overpay_amount, align_right, NUM_CURRENCY, font_info),
            (pay_coeff, align_right, NUM_PERCENT, font_info),
            (f"Qator {kr['row']}", align_center, None, font_small)
        ]
        row_fill = fill_zebra if tr % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, (val, align, num_fmt, f_style) in enumerate(vals, start=2):
            cell = ws6.cell(row=r6, column=col_idx, value=val)
            cell.alignment = align
            cell.font = f_style
            cell.border = cell_border
            if row_fill.fill_type: cell.fill = row_fill
            if num_fmt: cell.number_format = num_fmt
        ws6.row_dimensions[r6].height = 20
        r6 += 1

    ws6.merge_cells(start_row=r6, start_column=2, end_row=r6, end_column=6)
    ws6.cell(row=r6, column=2, value=f"JAMI ORTIQCHA TO'LOV ({len(overpay_records)} TALABA)").alignment = align_center
    ws6.cell(row=r6, column=2).font = font_bold
    ws6.cell(row=r6, column=7, value=f"=SUM(G6:G{r6-1})").number_format = NUM_CURRENCY
    ws6.cell(row=r6, column=8, value=f"=SUM(H6:H{r6-1})").number_format = NUM_CURRENCY
    ws6.cell(row=r6, column=9, value=f"=SUM(I6:I{r6-1})").number_format = NUM_CURRENCY

    for c in range(2, len(headers_ws6)+2):
        cell = ws6.cell(row=r6, column=c)
        cell.font = font_bold
        cell.fill = fill_total
        cell.border = total_border
        if cell.alignment.horizontal is None: cell.alignment = align_right
    ws6.row_dimensions[r6].height = 24
    ws6.freeze_panes = 'E6'

    # Save report
    out_file1 = 'Kontraktlar/Talabalar_va_Kontraktlar_Taqqoslash_Hisoboti.xlsx'
    wb_out.save(out_file1)
    out_file2 = 'Talabalar_va_Kontraktlar_Taqqoslash_Hisoboti.xlsx'
    shutil.copyfile(out_file1, out_file2)
    print(f"Successfully saved to:\n  1. {out_file1}\n  2. {out_file2}")

if __name__ == '__main__':
    create_report()
