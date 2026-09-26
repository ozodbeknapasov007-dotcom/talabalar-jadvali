# -*- coding: utf-8 -*-
"""
SHAHRISABZ ABU ALI IBN SINO NOMIDAGI JAMOAT SALOMATLIGI TEXNIKUMI (Sh.JSST)
2026/2027-O'QUV YILI YAKUNIY TANLOV NATIJALARI QAYDNOMASI PROTSESSORI
==========================================================================
Ushbu skript 'Шахрисабз ЖСТ.pdf' faylidan barcha 843 ta abituriyentning
natijalarini ajratib oladi, tahlil qiladi, Excel jadvali (.xlsx) va
interaktiv HTML hisobotini yaratadi.
"""

import os
import sys
import re
import json
import pymupdf
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict, Counter

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PDF_PATH = os.path.join(BASE_DIR, 'Шахрисабз ЖСТ.pdf')
EXCEL_PATH = os.path.join(BASE_DIR, 'Shahrisabz_JSST_Qaydnomasi.xlsx')
HTML_PATH = os.path.join(BASE_DIR, 'hisobot.html')
INDEX_HTML_PATH = os.path.join(BASE_DIR, 'index.html')
JSON_PATH = os.path.join(BASE_DIR, 'talabalar_bazasi.json')

def clean_uz_name(text):
    if not text:
        return '', '', '', '', 'Noma\'lum'
    text = str(text).strip()
    text = re.sub(r'\s+', ' ', text)
    # Fix OCR and character typos
    text = re.sub(r"0[\'‘\’\s]*G[\'‘\’\s]*LI", "o'g'li", text, flags=re.I)
    text = re.sub(r"O[\'‘\’\s]*G[\'‘\’\s]*LI", "o'g'li", text, flags=re.I)
    text = re.sub(r"[\'‘\’´`\?]", "'", text)
    text = re.sub(r"0\'RAL", "O'ral", text, flags=re.I)
    text = re.sub(r"0\'KTAM", "O'ktam", text, flags=re.I)
    text = re.sub(r"BAXT1GUL", "Baxtigul", text, flags=re.I)
    text = re.sub(r"HUS\s*AN", "Husan", text, flags=re.I)
    text = re.sub(r"A\'ZAM", "A'zam", text, flags=re.I)
    text = re.sub(r"G\'OLIB", "G'olib", text, flags=re.I)
    text = re.sub(r"G\'AYBULLA", "G'aybulla", text, flags=re.I)
    text = re.sub(r"G\'ULOM", "G'ulom", text, flags=re.I)
    
    words = text.split()
    cleaned_words = []
    for w in words:
        wl = w.lower()
        if wl in ['qizi', 'kizi']:
            cleaned_words.append('qizi')
        elif wl in ["o'g'li", "o‘g‘li", "o’g’li", "ogli", "ugli", "o'gli"]:
            cleaned_words.append("o'g'li")
        elif "'" in w:
            parts = w.split("'")
            fp = [parts[0].capitalize()] + [p.lower() for p in parts[1:]]
            cleaned_words.append("'".join(fp))
        else:
            cleaned_words.append(w.capitalize())
            
    full_fish = ' '.join(cleaned_words).strip()
    
    fam = cleaned_words[0] if len(cleaned_words) > 0 else ''
    ism = cleaned_words[1] if len(cleaned_words) > 1 else ''
    ota = ' '.join(cleaned_words[2:]) if len(cleaned_words) > 2 else ''
    
    gender = 'Noma\'lum'
    if 'qizi' in full_fish.lower() or full_fish.lower().endswith(('ovna', 'evna', 'ova', 'eva', 'kizi')):
        gender = 'Ayol'
    if "o'g'li" in full_fish.lower() or full_fish.lower().endswith(('ovich', 'evich', 'ov', 'ev', 'uli', 'ogli')):
        gender = 'Erkak'
        
    return fam, ism, ota, full_fish, gender

def parse_score(val):
    v = str(val).replace(',', '.').strip()
    try:
        return float(v)
    except:
        return 0.0

def extract_pdf_data(pdf_path):
    print(f"📖 PDF o'qilmoqda: {pdf_path}")
    doc = pymupdf.open(pdf_path)
    students = []
    current_direction = ''
    current_section = ''

    for pno, page in enumerate(doc):
        text = page.get_text('text')
        
        # Yo'nalishni aniqlash
        dir_m = re.search(r'TA[\'\‘\’\s]*LIM YO[\'\‘\’\s]*NALISHI:\s*([0-9]{8}\s*-\s*[^\n]+)', text)
        if not dir_m:
            dir_m = re.search(r'([0-9]{8}\s*-\s*[^\n]+)', text)
        if dir_m:
            d_clean = dir_m.group(1).strip()
            d_clean = re.sub(r'TA[\'\‘\’\s]*LIM YO[\'\‘\’\s]*NALISHI:?', '', d_clean, flags=re.I).strip()
            current_direction = d_clean
            
        words = page.get_text('words')
        
        # 1. Bo'limlarni topish (Grant / Kontrakt / Chegara)
        section_markers = []
        for w in words:
            if '===' in w[4] or '=====' in w[4]:
                y = w[1]
                nearby = [w2[4] for w2 in words if abs(w2[1] - y) < 6]
                nearby_str = ' '.join(nearby)
                if 'Davlat grant' in nearby_str:
                    section_markers.append((y, 'Davlat granti'))
                elif 'To\'lov' in nearby_str or 'kontrakt' in nearby_str or 'shartnoma' in nearby_str:
                    section_markers.append((y, 'To\'lov-kontrakt'))
                elif 'chegarasi' in nearby_str:
                    section_markers.append((y, 'Qabul chegarasidan tashqari'))
                    
        section_markers = sorted(list({round(m[0], 1): m[1] for m in section_markers}.items()))
        
        # 2. 7 xonali ID raqamlarini topish
        id_words = [w for w in words if 315 <= w[0] <= 365 and re.match(r'^\d{7}$', w[4]) and 190 <= w[1] <= 1080]
        id_words = sorted(id_words, key=lambda x: x[1])
        
        for iw in id_words:
            y_id = iw[1]
            
            for sm_y, sm_name in section_markers:
                if y_id >= sm_y - 8:
                    current_section = sm_name
                    
            row_w = [w for w in words if abs(w[1] - y_id) <= 6.0]
            row_w = sorted(row_w, key=lambda x: x[0])
            
            tr_w = [w[4] for w in row_w if w[0] < 75]
            name_w = [w[4] for w in row_w if 75 <= w[0] < 315]
            ball_w = [w[4] for w in row_w if 365 <= w[0] < 410]
            maj_w = [w[4] for w in row_w if 410 <= w[0] < 450]
            f1_w = [w[4] for w in row_w if 450 <= w[0] < 515]
            f2_w = [w[4] for w in row_w if 515 <= w[0] < 580]
            milliy_w = [w[4] for w in row_w if 580 <= w[0] < 640]
            chet_w = [w[4] for w in row_w if 640 <= w[0] < 690]
            kasb_w = [w[4] for w in row_w if 690 <= w[0] < 750]
            imtiyoz_w = [w[4] for w in row_w if 750 <= w[0] < 820]
            
            if not tr_w:
                tr_near = [w[4] for w in words if abs(w[1] - y_id) <= 9.0 and w[0] < 75]
                if tr_near:
                    tr_w = tr_near
                    
            name_str = ' '.join(name_w).strip()
            if len(name_w) < 2:
                name_near_sorted = sorted([w for w in words if abs(w[1] - y_id) <= 8.5 and 75 <= w[0] < 315], key=lambda x: (round(x[1]/4), x[0]))
                name_str = ' '.join([w[4] for w in name_near_sorted]).strip()
                
            fam, ism, ota, full_fish, gender = clean_uz_name(name_str)
            
            total_b_str = ' '.join(ball_w).strip()
            maj_b_str = ' '.join(maj_w).strip()
            f1_b_str = ' '.join(f1_w).strip()
            f2_b_str = ' '.join(f2_w).strip()
            milliy_b_str = ' '.join(milliy_w).strip()
            chet_b_str = ' '.join(chet_w).strip()
            kasb_b_str = ' '.join(kasb_w).strip()
            imtiyoz_b_str = ' '.join(imtiyoz_w).strip()
            
            # Extract direction code and name
            dir_parts = current_direction.split('-', 1)
            dir_code = dir_parts[0].strip() if len(dir_parts) > 1 else ''
            dir_name = dir_parts[1].strip() if len(dir_parts) > 1 else current_direction
            
            students.append({
                'global_tr': len(students) + 1,
                'page': pno + 1,
                'direction_full': current_direction,
                'direction_code': dir_code,
                'direction_name': dir_name,
                'section': current_section,
                'tr_in_section': ' '.join(tr_w).strip(),
                'name_raw': name_str,
                'fam': fam,
                'ism': ism,
                'ota': ota,
                'full_fish': full_fish,
                'gender': gender,
                'id_number': iw[4],
                'total_ball_str': total_b_str,
                'total_ball': parse_score(total_b_str),
                'majburiy': parse_score(maj_b_str) if maj_b_str else 0.0,
                'fan1': parse_score(f1_b_str) if f1_b_str else 0.0,
                'fan2': parse_score(f2_b_str) if f2_b_str else 0.0,
                'milliy': parse_score(milliy_b_str) if milliy_b_str else 0.0,
                'chet_tili': chet_b_str,
                'kasbiy': parse_score(kasb_b_str) if kasb_b_str else 0.0,
                'imtiyoz': parse_score(imtiyoz_b_str) if imtiyoz_b_str else 0.0,
            })

    print(f"✅ Jami ajratib olingan abituriyentlar soni: {len(students)} ta")
    return students

def generate_excel_file(students, excel_path):
    print(f"📊 Excel fayli yaratilmoqda: {excel_path}")
    wb = openpyxl.Workbook()
    
    # Styles
    font_header = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    font_bold = Font(name='Arial', size=10, bold=True)
    font_regular = Font(name='Arial', size=10)
    font_small = Font(name='Arial', size=9, color='555555')
    
    fill_header_main = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Navy
    fill_header_grant = PatternFill(start_color='065F46', end_color='065F46', fill_type='solid') # Green
    fill_header_kontrakt = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid') # Blue
    fill_header_chegara = PatternFill(start_color='9A3412', end_color='9A3412', fill_type='solid') # Amber/Rust
    fill_header_stats = PatternFill(start_color='334155', end_color='334155', fill_type='solid') # Slate
    
    fill_grant_row = PatternFill(start_color='F0FDF4', end_color='F0FDF4', fill_type='solid')
    fill_kontrakt_row = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')
    fill_chegara_row = PatternFill(start_color='FFF7ED', end_color='FFF7ED', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    
    border_thin = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    border_header = Border(left=border_thin, right=border_thin, top=border_thin, bottom=Side(border_style="medium", color="0F172A"))
    
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    
    headers = [
        "T/r", "Abituriyent ID", "To'liq F.I.SH", "Familiyasi", "Ismi", "Otasining ismi (Sharifi)",
        "Ta'lim Yo'nalishi", "Yo'nalish Kodi", "Qabul Holati (Status)", "Bo'lim T/r",
        "Umumiy Ball", "Majburiy Fanlar (1.1)", "1-Fan (Biologiya/Kimyo)", "2-Fan (Kimyo/Biologiya)",
        "Milliy Sertifikat", "Chet Tili", "Kasbiy Imtihon", "Imtiyoz Bali", "Jinsi", "PDF Sahifa"
    ]
    
    # -------------------------------------------------------------
    # SHEET 1: BARCHA ABITURIYENTLAR (843)
    # -------------------------------------------------------------
    ws_all = wb.active
    ws_all.title = "Barcha Abituriyentlar (843)"
    ws_all.views.sheetView[0].showGridLines = True
    
    # Title Row
    ws_all.merge_cells('A1:T1')
    title_cell = ws_all['A1']
    title_cell.value = "SHAHRISABZ ABU ALI IBN SINO NOMIDAGI JAMOAT SALOMATLIGI TEXNIKUMI — 2026/2027-O'QUV YILI YAKUNIY QABUL QAYDNOMASI"
    title_cell.font = Font(name='Arial', size=13, bold=True, color='1E3A8A')
    title_cell.alignment = align_center
    ws_all.row_dimensions[1].height = 30
    
    # Header Row
    ws_all.row_dimensions[2].height = 26
    for col_idx, h in enumerate(headers, 1):
        c = ws_all.cell(row=2, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_main
        c.alignment = align_center
        c.border = border_header
        
    for r_idx, s in enumerate(students, 3):
        ws_all.row_dimensions[r_idx].height = 20
        row_vals = [
            s['global_tr'], s['id_number'], s['full_fish'], s['fam'], s['ism'], s['ota'],
            s['direction_name'], s['direction_code'], s['section'], s['tr_in_section'],
            s['total_ball'], s['majburiy'], s['fan1'], s['fan2'],
            s['milliy'], s['chet_tili'], s['kasbiy'], s['imtiyoz'], s['gender'], s['page']
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_all.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx in [1, 2, 8, 10, 19, 20]:
                cell.alignment = align_center
            elif c_idx in [3, 4, 5, 6, 7, 9]:
                cell.alignment = align_left
            else:
                cell.alignment = align_right
                if isinstance(val, float):
                    cell.number_format = '0.00'
            
            # Conditional row color
            if s['section'] == 'Davlat granti':
                cell.fill = fill_grant_row
            elif s['section'] == 'To\'lov-kontrakt':
                cell.fill = fill_kontrakt_row
            else:
                if r_idx % 2 == 0:
                    cell.fill = fill_zebra

    # -------------------------------------------------------------
    # SHEET 2: DAVLAT GRANTI (108)
    # -------------------------------------------------------------
    ws_grant = wb.create_sheet(title="Davlat Granti (108)")
    ws_grant.views.sheetView[0].showGridLines = True
    ws_grant.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers, 1):
        c = ws_grant.cell(row=1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_grant
        c.alignment = align_center
        c.border = border_header
        
    grant_students = [s for s in students if s['section'] == 'Davlat granti']
    for r_idx, s in enumerate(grant_students, 2):
        ws_grant.row_dimensions[r_idx].height = 20
        row_vals = [
            r_idx - 1, s['id_number'], s['full_fish'], s['fam'], s['ism'], s['ota'],
            s['direction_name'], s['direction_code'], s['section'], s['tr_in_section'],
            s['total_ball'], s['majburiy'], s['fan1'], s['fan2'],
            s['milliy'], s['chet_tili'], s['kasbiy'], s['imtiyoz'], s['gender'], s['page']
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_grant.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            cell.fill = fill_grant_row
            if c_idx in [1, 2, 8, 10, 19, 20]:
                cell.alignment = align_center
            elif c_idx in [3, 4, 5, 6, 7, 9]:
                cell.alignment = align_left
            else:
                cell.alignment = align_right
                if isinstance(val, float):
                    cell.number_format = '0.00'

    # -------------------------------------------------------------
    # SHEET 3: TO'LOV-KONTRAKT (312)
    # -------------------------------------------------------------
    ws_kontrakt = wb.create_sheet(title="To'lov-Kontrakt (312)")
    ws_kontrakt.views.sheetView[0].showGridLines = True
    ws_kontrakt.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers, 1):
        c = ws_kontrakt.cell(row=1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_kontrakt
        c.alignment = align_center
        c.border = border_header
        
    kontrakt_students = [s for s in students if s['section'] == 'To\'lov-kontrakt']
    for r_idx, s in enumerate(kontrakt_students, 2):
        ws_kontrakt.row_dimensions[r_idx].height = 20
        row_vals = [
            r_idx - 1, s['id_number'], s['full_fish'], s['fam'], s['ism'], s['ota'],
            s['direction_name'], s['direction_code'], s['section'], s['tr_in_section'],
            s['total_ball'], s['majburiy'], s['fan1'], s['fan2'],
            s['milliy'], s['chet_tili'], s['kasbiy'], s['imtiyoz'], s['gender'], s['page']
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_kontrakt.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            cell.fill = fill_kontrakt_row
            if c_idx in [1, 2, 8, 10, 19, 20]:
                cell.alignment = align_center
            elif c_idx in [3, 4, 5, 6, 7, 9]:
                cell.alignment = align_left
            else:
                cell.alignment = align_right
                if isinstance(val, float):
                    cell.number_format = '0.00'

    # -------------------------------------------------------------
    # SHEET 4: QABUL CHEGARASIDAN TASHQARI (423)
    # -------------------------------------------------------------
    ws_chegara = wb.create_sheet(title="Chegaradan Tashqari (423)")
    ws_chegara.views.sheetView[0].showGridLines = True
    ws_chegara.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers, 1):
        c = ws_chegara.cell(row=1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_chegara
        c.alignment = align_center
        c.border = border_header
        
    chegara_students = [s for s in students if s['section'] == 'Qabul chegarasidan tashqari']
    for r_idx, s in enumerate(chegara_students, 2):
        ws_chegara.row_dimensions[r_idx].height = 20
        row_vals = [
            r_idx - 1, s['id_number'], s['full_fish'], s['fam'], s['ism'], s['ota'],
            s['direction_name'], s['direction_code'], s['section'], s['tr_in_section'],
            s['total_ball'], s['majburiy'], s['fan1'], s['fan2'],
            s['milliy'], s['chet_tili'], s['kasbiy'], s['imtiyoz'], s['gender'], s['page']
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_chegara.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if r_idx % 2 == 0: cell.fill = fill_chegara_row
            if c_idx in [1, 2, 8, 10, 19, 20]:
                cell.alignment = align_center
            elif c_idx in [3, 4, 5, 6, 7, 9]:
                cell.alignment = align_left
            else:
                cell.alignment = align_right
                if isinstance(val, float):
                    cell.number_format = '0.00'

    # -------------------------------------------------------------
    # SHEET 5: YO'NALISHLAR BO'YICHA STATISTIKA
    # -------------------------------------------------------------
    ws_stats = wb.create_sheet(title="Yo'nalishlar Statistikasi")
    ws_stats.views.sheetView[0].showGridLines = True
    
    stat_headers = [
        "T/r", "Ta'lim Yo'nalishi Kodi", "Ta'lim Yo'nalishi Nomi", "Jami Abituriyentlar",
        "Grant Kvotasi (Nafar)", "Grant Min Ball (O'tish bali)", "Grant Max Ball", "Grant O'rtacha Ball",
        "Kontrakt Kvotasi (Nafar)", "Kontrakt Min Ball (O'tish bali)", "Kontrakt Max Ball", "Kontrakt O'rtacha Ball",
        "Chegaradan Tashqari (Nafar)", "Chegara Min Ball", "Chegara Max Ball", "Chegara O'rtacha Ball"
    ]
    ws_stats.row_dimensions[1].height = 28
    for col_idx, h in enumerate(stat_headers, 1):
        c = ws_stats.cell(row=1, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_stats
        c.alignment = align_center
        c.border = border_header
        
    stats_map = defaultdict(lambda: {
        'code': '', 'name': '',
        'Davlat granti': [], 'To\'lov-kontrakt': [], 'Qabul chegarasidan tashqari': []
    })
    for s in students:
        d_full = s['direction_full']
        stats_map[d_full]['code'] = s['direction_code']
        stats_map[d_full]['name'] = s['direction_name']
        stats_map[d_full][s['section']].append(s['total_ball'])
        
    for r_idx, (d_full, d_data) in enumerate(stats_map.items(), 2):
        ws_stats.row_dimensions[r_idx].height = 22
        g_sc = d_data['Davlat granti']
        k_sc = d_data['To\'lov-kontrakt']
        ch_sc = d_data['Qabul chegarasidan tashqari']
        tot = len(g_sc) + len(k_sc) + len(ch_sc)
        
        row_vals = [
            r_idx - 1, d_data['code'], d_data['name'], tot,
            len(g_sc), min(g_sc) if g_sc else '-', max(g_sc) if g_sc else '-', sum(g_sc)/len(g_sc) if g_sc else '-',
            len(k_sc), min(k_sc) if k_sc else '-', max(k_sc) if k_sc else '-', sum(k_sc)/len(k_sc) if k_sc else '-',
            len(ch_sc), min(ch_sc) if ch_sc else '-', max(ch_sc) if ch_sc else '-', sum(ch_sc)/len(ch_sc) if ch_sc else '-',
        ]
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_stats.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            if c_idx in [1, 2]: cell.alignment = align_center
            elif c_idx == 3: cell.alignment = align_left
            else:
                cell.alignment = align_right
                if isinstance(val, float): cell.number_format = '0.00'

    # Auto adjust column widths for all sheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # skip merged title
                if ws == ws_all and cell.row == 1: continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    wb.save(excel_path)
    print(f"✅ Excel fayli muvaffaqiyatli saqlandi!")

def generate_html_dashboard(students, html_path, index_path):
    print(f"🌐 Interaktiv HTML Dashboard yaratilmoqda...")
    
    # Precompute statistics for HTML dashboard
    total_count = len(students)
    grant_count = len([s for s in students if s['section'] == 'Davlat granti'])
    kontrakt_count = len([s for s in students if s['section'] == 'To\'lov-kontrakt'])
    chegara_count = len([s for s in students if s['section'] == 'Qabul chegarasidan tashqari'])
    
    scores = [s['total_ball'] for s in students]
    max_score = max(scores) if scores else 0
    min_score = min(scores) if scores else 0
    avg_score = sum(scores) / len(scores) if scores else 0
    
    # Direction breakdown
    dir_stats = {}
    for s in students:
        d = s['direction_name']
        if d not in dir_stats:
            dir_stats[d] = {'code': s['direction_code'], 'total': 0, 'grant': 0, 'kontrakt': 0, 'chegara': 0, 'scores': []}
        dir_stats[d]['total'] += 1
        dir_stats[d]['scores'].append(s['total_ball'])
        if s['section'] == 'Davlat granti': dir_stats[d]['grant'] += 1
        elif s['section'] == 'To\'lov-kontrakt': dir_stats[d]['kontrakt'] += 1
        elif s['section'] == 'Qabul chegarasidan tashqari': dir_stats[d]['chegara'] += 1

    students_json_str = json.dumps(students, ensure_ascii=False)
    
    html_content = f"""<!DOCTYPE html>
<html lang="uz">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Shahrisabz Abu Ali ibn Sino JSST — 2026/2027 Yakuniy Qabul Qaydnomasi</title>
    <!-- Google Fonts -->
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&display=swap" rel="stylesheet">
    <!-- Lucide Icons -->
    <script src="https://unpkg.com/lucide@latest"></script>
    <!-- Chart.js -->
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <!-- SheetJS for Export -->
    <script src="https://cdn.jsdelivr.net/npm/xlsx/dist/xlsx.full.min.js"></script>
    <style>
        :root {{
            --primary: #2563eb;
            --primary-dark: #1d4ed8;
            --primary-light: #dbeafe;
            --success: #10b981;
            --success-dark: #059669;
            --success-bg: #d1fae5;
            --warning: #f59e0b;
            --warning-dark: #d97706;
            --warning-bg: #fef3c7;
            --danger: #ef4444;
            --danger-dark: #dc2626;
            --danger-bg: #fee2e2;
            --info: #06b6d4;
            --slate-50: #f8fafc;
            --slate-100: #f1f5f9;
            --slate-200: #e2e8f0;
            --slate-300: #cbd5e1;
            --slate-600: #475569;
            --slate-700: #334155;
            --slate-800: #1e293b;
            --slate-900: #0f172a;
            --card-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.05), 0 8px 10px -6px rgba(0, 0, 0, 0.03);
            --card-shadow-hover: 0 20px 30px -10px rgba(37, 99, 235, 0.12), 0 10px 15px -5px rgba(0, 0, 0, 0.04);
        }}

        * {{
            margin: 0;
            padding: 0;
            box-sizing: border-box;
            font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, sans-serif;
        }}

        body {{
            background: linear-gradient(135deg, #f0f4f8 0%, #e2e8f0 100%);
            color: var(--slate-800);
            min-height: 100vh;
            padding-bottom: 60px;
        }}

        .navbar {{
            background: rgba(255, 255, 255, 0.95);
            backdrop-filter: blur(12px);
            border-bottom: 1px solid var(--slate-200);
            position: sticky;
            top: 0;
            z-index: 50;
            padding: 16px 32px;
            display: flex;
            justify-content: space-between;
            align-items: center;
            box-shadow: 0 2px 10px rgba(0,0,0,0.03);
        }}

        .brand-box {{
            display: flex;
            align-items: center;
            gap: 14px;
        }}

        .brand-icon {{
            background: linear-gradient(135deg, #2563eb 0%, #1d4ed8 100%);
            color: white;
            width: 46px;
            height: 46px;
            border-radius: 12px;
            display: flex;
            align-items: center;
            justify-content: center;
            box-shadow: 0 4px 12px rgba(37, 99, 235, 0.3);
        }}

        .brand-title {{
            font-size: 1.15rem;
            font-weight: 800;
            color: var(--slate-900);
            line-height: 1.2;
        }}

        .brand-subtitle {{
            font-size: 0.8rem;
            font-weight: 500;
            color: var(--slate-600);
        }}

        .nav-actions {{
            display: flex;
            align-items: center;
            gap: 12px;
        }}

        .btn {{
            display: inline-flex;
            align-items: center;
            gap: 8px;
            padding: 9px 18px;
            border-radius: 10px;
            font-size: 0.875rem;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.2s ease;
            border: 1px solid transparent;
            text-decoration: none;
        }}

        .btn-primary {{
            background: var(--primary);
            color: white;
        }}
        .btn-primary:hover {{
            background: var(--primary-dark);
            transform: translateY(-1px);
            box-shadow: 0 4px 12px rgba(37, 99, 235, 0.25);
        }}

        .btn-success {{
            background: #10b981;
            color: white;
        }}
        .btn-success:hover {{
            background: #059669;
            transform: translateY(-1px);
            box-shadow: 0 4px 12px rgba(16, 185, 129, 0.25);
        }}

        .btn-secondary {{
            background: white;
            color: var(--slate-700);
            border-color: var(--slate-300);
        }}
        .btn-secondary:hover {{
            background: var(--slate-50);
            border-color: var(--slate-400);
        }}

        .container {{
            max-width: 1440px;
            margin: 0 auto;
            padding: 28px 24px;
        }}

        /* Hero Banner */
        .hero-banner {{
            background: linear-gradient(135deg, #1e3a8a 0%, #1e40af 50%, #2563eb 100%);
            border-radius: 20px;
            padding: 32px 36px;
            color: white;
            margin-bottom: 28px;
            box-shadow: 0 15px 30px rgba(30, 58, 138, 0.2);
            position: relative;
            overflow: hidden;
        }}

        .hero-banner::after {{
            content: '';
            position: absolute;
            top: -50px;
            right: -50px;
            width: 250px;
            height: 250px;
            background: radial-gradient(circle, rgba(255,255,255,0.15) 0%, rgba(255,255,255,0) 70%);
            border-radius: 50%;
        }}

        .hero-title {{
            font-size: 1.85rem;
            font-weight: 800;
            margin-bottom: 8px;
            letter-spacing: -0.5px;
        }}

        .hero-desc {{
            font-size: 0.95rem;
            opacity: 0.9;
            max-width: 800px;
            line-height: 1.5;
        }}

        /* KPI Stats Grid */
        .stats-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
            gap: 18px;
            margin-bottom: 28px;
        }}

        .stat-card {{
            background: white;
            border-radius: 16px;
            padding: 22px;
            box-shadow: var(--card-shadow);
            border: 1px solid rgba(255,255,255,0.8);
            transition: all 0.25s ease;
            position: relative;
            overflow: hidden;
        }}

        .stat-card:hover {{
            transform: translateY(-3px);
            box-shadow: var(--card-shadow-hover);
        }}

        .stat-card::before {{
            content: '';
            position: absolute;
            top: 0;
            left: 0;
            width: 5px;
            height: 100%;
        }}

        .stat-card.total::before {{ background: var(--primary); }}
        .stat-card.grant::before {{ background: var(--success); }}
        .stat-card.kontrakt::before {{ background: #3b82f6; }}
        .stat-card.chegara::before {{ background: var(--warning); }}
        .stat-card.max::before {{ background: #8b5cf6; }}
        .stat-card.avg::before {{ background: var(--info); }}

        .stat-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 12px;
        }}

        .stat-label {{
            font-size: 0.8rem;
            font-weight: 700;
            text-transform: uppercase;
            letter-spacing: 0.5px;
            color: var(--slate-600);
        }}

        .stat-icon {{
            width: 36px;
            height: 36px;
            border-radius: 10px;
            display: flex;
            align-items: center;
            justify-content: center;
        }}

        .stat-card.total .stat-icon {{ background: var(--primary-light); color: var(--primary); }}
        .stat-card.grant .stat-icon {{ background: var(--success-bg); color: var(--success-dark); }}
        .stat-card.kontrakt .stat-icon {{ background: #eff6ff; color: #1d4ed8; }}
        .stat-card.chegara .stat-icon {{ background: var(--warning-bg); color: var(--warning-dark); }}
        .stat-card.max .stat-icon {{ background: #f3e8ff; color: #7c3aed; }}
        .stat-card.avg .stat-icon {{ background: #cffafe; color: #0891b2; }}

        .stat-value {{
            font-size: 1.9rem;
            font-weight: 800;
            color: var(--slate-900);
            line-height: 1.1;
        }}

        .stat-sub {{
            font-size: 0.8rem;
            font-weight: 600;
            color: var(--slate-600);
            margin-top: 6px;
        }}

        /* Section Breakdown Cards */
        .direction-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
            gap: 18px;
            margin-bottom: 28px;
        }}

        .dir-card {{
            background: white;
            border-radius: 16px;
            padding: 20px;
            box-shadow: var(--card-shadow);
            border: 1px solid var(--slate-200);
            cursor: pointer;
            transition: all 0.2s ease;
        }}

        .dir-card:hover, .dir-card.active {{
            border-color: var(--primary);
            box-shadow: 0 8px 20px rgba(37,99,235,0.12);
            transform: translateY(-2px);
        }}

        .dir-code {{
            font-size: 0.75rem;
            font-weight: 700;
            color: var(--primary);
            background: var(--primary-light);
            padding: 3px 8px;
            border-radius: 6px;
            display: inline-block;
            margin-bottom: 8px;
        }}

        .dir-name {{
            font-size: 1rem;
            font-weight: 700;
            color: var(--slate-900);
            margin-bottom: 12px;
            line-height: 1.3;
        }}

        .dir-stats-row {{
            display: flex;
            justify-content: space-between;
            font-size: 0.825rem;
            padding: 4px 0;
            color: var(--slate-700);
            border-bottom: 1px dashed var(--slate-200);
        }}

        .dir-stats-row:last-child {{
            border-bottom: none;
        }}

        .dir-stats-row strong {{
            font-weight: 700;
        }}

        /* Visual Charts Container */
        .charts-row {{
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 20px;
            margin-bottom: 28px;
        }}

        @media (max-width: 992px) {{
            .charts-row {{ grid-template-columns: 1fr; }}
        }}

        .chart-box {{
            background: white;
            border-radius: 16px;
            padding: 24px;
            box-shadow: var(--card-shadow);
            border: 1px solid var(--slate-200);
        }}

        .chart-header {{
            font-size: 1.05rem;
            font-weight: 700;
            color: var(--slate-900);
            margin-bottom: 18px;
            display: flex;
            align-items: center;
            gap: 10px;
        }}

        /* Table Card & Filter Bar */
        .table-card {{
            background: white;
            border-radius: 18px;
            box-shadow: var(--card-shadow);
            border: 1px solid var(--slate-200);
            overflow: hidden;
        }}

        .table-toolbar {{
            padding: 20px 24px;
            border-bottom: 1px solid var(--slate-200);
            background: var(--slate-50);
            display: flex;
            flex-direction: column;
            gap: 16px;
        }}

        .toolbar-top {{
            display: flex;
            flex-wrap: wrap;
            justify-content: space-between;
            align-items: center;
            gap: 14px;
        }}

        .search-box {{
            position: relative;
            flex: 1;
            min-width: 280px;
            max-width: 450px;
        }}

        .search-box i {{
            position: absolute;
            left: 14px;
            top: 50%;
            transform: translateY(-50%);
            color: var(--slate-400);
        }}

        .search-input {{
            width: 100%;
            padding: 11px 16px 11px 42px;
            border-radius: 10px;
            border: 1px solid var(--slate-300);
            font-size: 0.9rem;
            background: white;
            outline: none;
            transition: all 0.2s ease;
        }}

        .search-input:focus {{
            border-color: var(--primary);
            box-shadow: 0 0 0 3px rgba(37, 99, 235, 0.15);
        }}

        .filter-tabs {{
            display: flex;
            flex-wrap: wrap;
            gap: 8px;
        }}

        .tab-btn {{
            padding: 8px 14px;
            border-radius: 8px;
            font-size: 0.825rem;
            font-weight: 600;
            border: 1px solid var(--slate-300);
            background: white;
            color: var(--slate-700);
            cursor: pointer;
            transition: all 0.2s ease;
            display: inline-flex;
            align-items: center;
            gap: 6px;
        }}

        .tab-btn:hover {{
            background: var(--slate-100);
            border-color: var(--slate-400);
        }}

        .tab-btn.active {{
            background: var(--primary);
            color: white;
            border-color: var(--primary);
        }}

        .tab-badge {{
            padding: 2px 6px;
            border-radius: 12px;
            font-size: 0.725rem;
            font-weight: 700;
            background: rgba(0,0,0,0.08);
        }}

        .tab-btn.active .tab-badge {{
            background: rgba(255,255,255,0.25);
            color: white;
        }}

        .filter-row-secondary {{
            display: flex;
            flex-wrap: wrap;
            align-items: center;
            gap: 12px;
            font-size: 0.85rem;
        }}

        .select-filter {{
            padding: 8px 12px;
            border-radius: 8px;
            border: 1px solid var(--slate-300);
            background: white;
            font-size: 0.85rem;
            font-weight: 500;
            color: var(--slate-700);
            outline: none;
        }}

        /* Table Design */
        .table-responsive {{
            overflow-x: auto;
            max-height: 650px;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
            text-align: left;
            font-size: 0.875rem;
        }}

        thead {{
            background: var(--slate-100);
            position: sticky;
            top: 0;
            z-index: 10;
        }}

        th {{
            padding: 13px 14px;
            font-weight: 700;
            color: var(--slate-700);
            border-bottom: 2px solid var(--slate-300);
            cursor: pointer;
            user-select: none;
            white-space: nowrap;
        }}

        th:hover {{
            background: var(--slate-200);
        }}

        td {{
            padding: 12px 14px;
            border-bottom: 1px solid var(--slate-200);
            color: var(--slate-800);
            vertical-align: middle;
            white-space: nowrap;
        }}

        tr:hover td {{
            background: #f8fafc !important;
        }}

        /* Badges */
        .badge {{
            padding: 4px 10px;
            border-radius: 20px;
            font-size: 0.75rem;
            font-weight: 700;
            display: inline-flex;
            align-items: center;
            gap: 4px;
        }}

        .badge-grant {{
            background: var(--success-bg);
            color: var(--success-dark);
            border: 1px solid #a7f3d0;
        }}

        .badge-kontrakt {{
            background: #eff6ff;
            color: #1d4ed8;
            border: 1px solid #bfdbfe;
        }}

        .badge-chegara {{
            background: var(--warning-bg);
            color: var(--warning-dark);
            border: 1px solid #fde68a;
        }}

        .score-pill {{
            font-weight: 800;
            font-size: 0.95rem;
            color: var(--slate-900);
        }}

        .id-badge {{
            font-family: monospace;
            font-weight: 700;
            background: var(--slate-100);
            color: var(--slate-700);
            padding: 3px 6px;
            border-radius: 4px;
            font-size: 0.85rem;
        }}

        /* Pagination & Footer */
        .table-footer {{
            padding: 16px 24px;
            border-top: 1px solid var(--slate-200);
            background: var(--slate-50);
            display: flex;
            justify-content: space-between;
            align-items: center;
            flex-wrap: wrap;
            gap: 12px;
            font-size: 0.875rem;
            color: var(--slate-600);
        }}

        .pagination-box {{
            display: flex;
            align-items: center;
            gap: 6px;
        }}

        .page-btn {{
            width: 34px;
            height: 34px;
            border-radius: 8px;
            border: 1px solid var(--slate-300);
            background: white;
            color: var(--slate-700);
            display: flex;
            align-items: center;
            justify-content: center;
            cursor: pointer;
            font-weight: 600;
            transition: all 0.2s;
        }}

        .page-btn:hover:not(:disabled) {{
            background: var(--slate-100);
            border-color: var(--slate-400);
        }}

        .page-btn.active {{
            background: var(--primary);
            color: white;
            border-color: var(--primary);
        }}

        .page-btn:disabled {{
            opacity: 0.4;
            cursor: not-allowed;
        }}

        /* Footer Info */
        .footer-credit {{
            text-align: center;
            margin-top: 36px;
            font-size: 0.85rem;
            color: var(--slate-600);
        }}
    </style>
</head>
<body>

    <!-- Top Navigation -->
    <header class="navbar">
        <div class="brand-box">
            <div class="brand-icon">
                <i data-lucide="award"></i>
            </div>
            <div>
                <div class="brand-title">Shahrisabz Abu Ali ibn Sino JSST</div>
                <div class="brand-subtitle">2026/2027-o'quv yili yakuniy qabul natijalari qaydnomasi</div>
            </div>
        </div>
        <div class="nav-actions">
            <button class="btn btn-success" onclick="exportToExcel()">
                <i data-lucide="file-spreadsheet"></i> Excel Yuklab Olish (.xlsx)
            </button>
            <button class="btn btn-secondary" onclick="window.print()">
                <i data-lucide="printer"></i> Chop Etish
            </button>
        </div>
    </header>

    <div class="container">

        <!-- Hero Banner -->
        <div class="hero-banner">
            <h1 class="hero-title">Shahrisabz Jamoat Salomatligi Texnikumi — Qabul 2026/2027</h1>
            <p class="hero-desc">
                O'zbekiston Respublikasi Oliy ta'lim, fan va innovatsiyalar vazirligi huzuridagi Bilim va malakalarni baholash agentligi (DTM) tomonidan tasdiqlangan yakuniy tanlov natijalari qaydnomasi.
            </p>
        </div>

        <!-- KPI Metrics Grid -->
        <div class="stats-grid">
            <div class="stat-card total">
                <div class="stat-header">
                    <span class="stat-label">Jami Abituriyentlar</span>
                    <div class="stat-icon"><i data-lucide="users"></i></div>
                </div>
                <div class="stat-value">{total_count}</div>
                <div class="stat-sub">5 ta tibbiy yo'nalish bo'yicha</div>
            </div>

            <div class="stat-card grant">
                <div class="stat-header">
                    <span class="stat-label">Davlat Granti</span>
                    <div class="stat-icon"><i data-lucide="check-circle-2"></i></div>
                </div>
                <div class="stat-value">{grant_count}</div>
                <div class="stat-sub">{grant_count/total_count*100:.1f}% qabul kvotasi</div>
            </div>

            <div class="stat-card kontrakt">
                <div class="stat-header">
                    <span class="stat-label">To'lov-Kontrakt</span>
                    <div class="stat-icon"><i data-lucide="credit-card"></i></div>
                </div>
                <div class="stat-value">{kontrakt_count}</div>
                <div class="stat-sub">{kontrakt_count/total_count*100:.1f}% qabul kvotasi</div>
            </div>

            <div class="stat-card chegara">
                <div class="stat-header">
                    <span class="stat-label">Chegaradan Tashqari</span>
                    <div class="stat-icon"><i data-lucide="alert-circle"></i></div>
                </div>
                <div class="stat-value">{chegara_count}</div>
                <div class="stat-sub">{chegara_count/total_count*100:.1f}% tavsiya etilmagan</div>
            </div>

            <div class="stat-card max">
                <div class="stat-header">
                    <span class="stat-label">Eng Yuqori Ball</span>
                    <div class="stat-icon"><i data-lucide="trophy"></i></div>
                </div>
                <div class="stat-value">{max_score:.1f}</div>
                <div class="stat-sub">Hamshiralik ishi (Grant)</div>
            </div>

            <div class="stat-card avg">
                <div class="stat-header">
                    <span class="stat-label">O'rtacha Ball</span>
                    <div class="stat-icon"><i data-lucide="trending-up"></i></div>
                </div>
                <div class="stat-value">{avg_score:.1f}</div>
                <div class="stat-sub">Barcha abituriyentlar bo'yicha</div>
            </div>
        </div>

        <!-- Direction Cards Grid -->
        <h2 style="font-size: 1.25rem; font-weight: 800; margin-bottom: 14px; color: var(--slate-900);">
            🩺 Ta'lim Yo'nalishlari Bo'yicha Qabul Ko'rsatkichlari
        </h2>
        <div class="direction-grid">
            <div class="dir-card" onclick="filterByDirection('50910208-HAMSHIRALIK ISHI')">
                <span class="dir-code">50910208</span>
                <div class="dir-name">Hamshiralik ishi</div>
                <div class="dir-stats-row"><span>Jami abituriyent:</span> <strong>467 nafar</strong></div>
                <div class="dir-stats-row"><span>Davlat granti:</span> <strong style="color:var(--success-dark)">54 ta (Min: 69.5 ball)</strong></div>
                <div class="dir-stats-row"><span>To'lov-kontrakt:</span> <strong style="color:var(--primary)">216 ta (Min: 49.4 ball)</strong></div>
                <div class="dir-stats-row"><span>Chegaradan tashqari:</span> <strong>197 nafar</strong></div>
            </div>

            <div class="dir-card" onclick="filterByDirection('50910204-DAVOLASH ISHI')">
                <span class="dir-code">50910204</span>
                <div class="dir-name">Davolash ishi</div>
                <div class="dir-stats-row"><span>Jami abituriyent:</span> <strong>177 nafar</strong></div>
                <div class="dir-stats-row"><span>Davlat granti:</span> <strong style="color:var(--success-dark)">12 ta (Min: 118.2 ball)</strong></div>
                <div class="dir-stats-row"><span>To'lov-kontrakt:</span> <strong style="color:var(--primary)">48 ta (Min: 57.9 ball)</strong></div>
                <div class="dir-stats-row"><span>Chegaradan tashqari:</span> <strong>117 nafar</strong></div>
            </div>

            <div class="dir-card" onclick="filterByDirection('50910402-FARMATSIYA')">
                <span class="dir-code">50910402</span>
                <div class="dir-name">Farmatsiya</div>
                <div class="dir-stats-row"><span>Jami abituriyent:</span> <strong>82 nafar</strong></div>
                <div class="dir-stats-row"><span>Davlat granti:</span> <strong style="color:var(--success-dark)">30 ta (Min: 51.5 ball)</strong></div>
                <div class="dir-stats-row"><span>To'lov-kontrakt:</span> <strong>0 ta</strong></div>
                <div class="dir-stats-row"><span>Chegaradan tashqari:</span> <strong>52 nafar</strong></div>
            </div>

            <div class="dir-card" onclick="filterByDirection('50910205-FUNKSIONAL DIAGNOSTIKA ISHI')">
                <span class="dir-code">50910205</span>
                <div class="dir-name">Funksional diagnostika ishi</div>
                <div class="dir-stats-row"><span>Jami abituriyent:</span> <strong>60 nafar</strong></div>
                <div class="dir-stats-row"><span>Davlat granti:</span> <strong style="color:var(--success-dark)">6 ta (Min: 59.0 ball)</strong></div>
                <div class="dir-stats-row"><span>To'lov-kontrakt:</span> <strong style="color:var(--primary)">24 ta (Min: 48.6 ball)</strong></div>
                <div class="dir-stats-row"><span>Chegaradan tashqari:</span> <strong>30 nafar</strong></div>
            </div>

            <div class="dir-card" onclick="filterByDirection('40910302-TIBBIY PROFILAKTIKA ISHI')">
                <span class="dir-code">40910302</span>
                <div class="dir-name">Tibbiy profilaktika ishi</div>
                <div class="dir-stats-row"><span>Jami abituriyent:</span> <strong>57 nafar</strong></div>
                <div class="dir-stats-row"><span>Davlat granti:</span> <strong style="color:var(--success-dark)">6 ta (Min: 49.9 ball)</strong></div>
                <div class="dir-stats-row"><span>To'lov-kontrakt:</span> <strong style="color:var(--primary)">24 ta (Min: 39.0 ball)</strong></div>
                <div class="dir-stats-row"><span>Chegaradan tashqari:</span> <strong>27 nafar</strong></div>
            </div>
        </div>

        <!-- Charts Row -->
        <div class="charts-row">
            <div class="chart-box">
                <div class="chart-header">
                    <i data-lucide="pie-chart" style="color:var(--primary)"></i> Qabul Turlari Bo'yicha Taqsimot
                </div>
                <div style="height: 250px;">
                    <canvas id="admissionPieChart"></canvas>
                </div>
            </div>

            <div class="chart-box">
                <div class="chart-header">
                    <i data-lucide="bar-chart-3" style="color:var(--primary)"></i> Yo'nalishlar Bo'yicha Abituriyentlar Soni
                </div>
                <div style="height: 250px;">
                    <canvas id="directionBarChart"></canvas>
                </div>
            </div>
        </div>

        <!-- Main Interactive Data Table Card -->
        <div class="table-card">
            <div class="table-toolbar">
                <div class="toolbar-top">
                    <!-- Search Input -->
                    <div class="search-box">
                        <i data-lucide="search" size="18"></i>
                        <input type="text" id="searchInput" class="search-input" placeholder="F.I.SH, ID raqam yoki sharifi bo'yicha qidiring..." onkeyup="handleSearch()">
                    </div>

                    <!-- Quick Status Filter Tabs -->
                    <div class="filter-tabs">
                        <button class="tab-btn active" onclick="setStatusFilter('ALL', this)">
                            Barchasi <span class="tab-badge">{total_count}</span>
                        </button>
                        <button class="tab-btn" onclick="setStatusFilter('Davlat granti', this)">
                            🟢 Davlat Granti <span class="tab-badge">{grant_count}</span>
                        </button>
                        <button class="tab-btn" onclick="setStatusFilter('To\'lov-kontrakt', this)">
                            🔵 To'lov-Kontrakt <span class="tab-badge">{kontrakt_count}</span>
                        </button>
                        <button class="tab-btn" onclick="setStatusFilter('Qabul chegarasidan tashqari', this)">
                            🟠 Chegaradan Tashqari <span class="tab-badge">{chegara_count}</span>
                        </button>
                    </div>
                </div>

                <!-- Secondary Filter Controls -->
                <div class="filter-row-secondary">
                    <div>
                        <label style="font-weight:600; margin-right:6px; color:var(--slate-700)">Yo'nalish:</label>
                        <select id="directionSelect" class="select-filter" onchange="handleFilterChange()">
                            <option value="ALL">Barcha yo'nalishlar (5 ta)</option>
                            <option value="50910208-HAMSHIRALIK ISHI">Hamshiralik ishi (467)</option>
                            <option value="50910204-DAVOLASH ISHI">Davolash ishi (177)</option>
                            <option value="50910402-FARMATSIYA">Farmatsiya (82)</option>
                            <option value="50910205-FUNKSIONAL DIAGNOSTIKA ISHI">Funksional diagnostika ishi (60)</option>
                            <option value="40910302-TIBBIY PROFILAKTIKA ISHI">Tibbiy profilaktika ishi (57)</option>
                        </select>
                    </div>

                    <div>
                        <label style="font-weight:600; margin-right:6px; color:var(--slate-700)">Jinsi:</label>
                        <select id="genderSelect" class="select-filter" onchange="handleFilterChange()">
                            <option value="ALL">Barchasi</option>
                            <option value="Ayol">Ayol talabalar</option>
                            <option value="Erkak">Erkak talabalar</option>
                        </select>
                    </div>

                    <div>
                        <label style="font-weight:600; margin-right:6px; color:var(--slate-700)">Minimal Ball:</label>
                        <input type="number" id="minScoreInput" class="select-filter" placeholder="Masalan: 50" style="width: 100px;" oninput="handleFilterChange()">
                    </div>

                    <div style="margin-left: auto;">
                        <span id="filteredStatsText" style="font-weight:700; color:var(--primary);">Ko'rsatilmoqda: {total_count} ta</span>
                    </div>
                </div>
            </div>

            <!-- Table -->
            <div class="table-responsive">
                <table id="studentsTable">
                    <thead>
                        <tr>
                            <th onclick="sortTable('global_tr')">T/r ↕</th>
                            <th onclick="sortTable('id_number')">Abituriyent ID ↕</th>
                            <th onclick="sortTable('full_fish')">F.I.SH (Familiya Ism Otasining ismi) ↕</th>
                            <th onclick="sortTable('direction_name')">Ta'lim Yo'nalishi ↕</th>
                            <th onclick="sortTable('section')">Qabul Holati ↕</th>
                            <th onclick="sortTable('total_ball')">Umumiy Ball ↕</th>
                            <th onclick="sortTable('majburiy')">Majburiy ↕</th>
                            <th onclick="sortTable('fan1')">1-Fan ↕</th>
                            <th onclick="sortTable('fan2')">2-Fan ↕</th>
                            <th onclick="sortTable('imtiyoz')">Imtiyoz ↕</th>
                            <th onclick="sortTable('gender')">Jinsi ↕</th>
                            <th>PDF</th>
                        </tr>
                    </thead>
                    <tbody id="tableBody">
                        <!-- Dynamic JS rendering -->
                    </tbody>
                </table>
            </div>

            <!-- Table Footer / Pagination -->
            <div class="table-footer">
                <div>
                    Qatorlar: 
                    <select id="pageSizeSelect" class="select-filter" onchange="changePageSize(this.value)">
                        <option value="25">25 ta</option>
                        <option value="50" selected>50 ta</option>
                        <option value="100">100 ta</option>
                        <option value="843">Barchasi (843)</option>
                    </select>
                </div>
                <div class="pagination-box" id="paginationControls">
                    <!-- Dynamic Pagination -->
                </div>
            </div>
        </div>

        <div class="footer-credit">
            Shahrisabz Abu Ali ibn Sino nomidagi Jamoat Salomatligi Texnikumi — 2026/2027-o'quv yili hisoboti
        </div>

    </div>

    <!-- Data Injection & JavaScript Logic -->
    <script>
        const RAW_DATA = {students_json_str};
        let currentFilteredData = [...RAW_DATA];
        let statusFilter = 'ALL';
        let directionFilter = 'ALL';
        let genderFilter = 'ALL';
        let minScore = 0;
        let searchQuery = '';
        let sortColumn = 'global_tr';
        let sortAsc = true;
        let currentPage = 1;
        let pageSize = 50;

        // Initialize Icons & Charts
        document.addEventListener('DOMContentLoaded', () => {{
            lucide.createIcons();
            initCharts();
            renderTable();
        }});

        function initCharts() {{
            // Pie Chart
            const ctxPie = document.getElementById('admissionPieChart').getContext('2d');
            new Chart(ctxPie, {{
                type: 'doughnut',
                data: {{
                    labels: ['Davlat Granti', 'To\'lov-Kontrakt', 'Chegaradan Tashqari'],
                    datasets: [{{
                        data: [{grant_count}, {kontrakt_count}, {chegara_count}],
                        backgroundColor: ['#10b981', '#2563eb', '#f59e0b'],
                        borderWidth: 2,
                        borderColor: '#ffffff'
                    }}]
                }},
                options: {{
                    responsive: true,
                    maintainAspectRatio: false,
                    plugins: {{
                        legend: {{ position: 'bottom', labels: {{ font: {{ family: 'Plus Jakarta Sans', weight: '600' }} }} }}
                    }}
                }}
            }});

            // Bar Chart
            const ctxBar = document.getElementById('directionBarChart').getContext('2d');
            new Chart(ctxBar, {{
                type: 'bar',
                data: {{
                    labels: ['Hamshiralik', 'Davolash', 'Farmatsiya', 'Diagnostika', 'Profilaktika'],
                    datasets: [
                        {{ label: 'Grant', data: [54, 12, 30, 6, 6], backgroundColor: '#10b981' }},
                        {{ label: 'Kontrakt', data: [216, 48, 0, 24, 24], backgroundColor: '#2563eb' }},
                        {{ label: 'Chegaradan tashqari', data: [197, 117, 52, 30, 27], backgroundColor: '#f59e0b' }}
                    ]
                }},
                options: {{
                    responsive: true,
                    maintainAspectRatio: false,
                    scales: {{
                        x: {{ stacked: true, grid: {{ display: false }} }},
                        y: {{ stacked: true }}
                    }},
                    plugins: {{
                        legend: {{ position: 'bottom', labels: {{ font: {{ family: 'Plus Jakarta Sans', weight: '600' }} }} }}
                    }}
                }}
            }});
        }}

        function setStatusFilter(status, btnElement) {{
            statusFilter = status;
            document.querySelectorAll('.filter-tabs .tab-btn').forEach(btn => btn.classList.remove('active'));
            if (btnElement) btnElement.classList.add('active');
            applyFilters();
        }}

        function filterByDirection(dirCodeName) {{
            document.getElementById('directionSelect').value = dirCodeName;
            handleFilterChange();
            // scroll to table
            document.querySelector('.table-card').scrollIntoView({{ behavior: 'smooth' }});
        }}

        function handleSearch() {{
            searchQuery = document.getElementById('searchInput').value.trim().toLowerCase();
            applyFilters();
        }}

        function handleFilterChange() {{
            directionFilter = document.getElementById('directionSelect').value;
            genderFilter = document.getElementById('genderSelect').value;
            const minScoreVal = parseFloat(document.getElementById('minScoreInput').value);
            minScore = isNaN(minScoreVal) ? 0 : minScoreVal;
            applyFilters();
        }}

        function applyFilters() {{
            currentFilteredData = RAW_DATA.filter(item => {{
                // Status Filter
                if (statusFilter !== 'ALL' && item.section !== statusFilter) return false;
                // Direction Filter
                if (directionFilter !== 'ALL' && item.direction_full !== directionFilter) return false;
                // Gender Filter
                if (genderFilter !== 'ALL' && item.gender !== genderFilter) return false;
                // Min Score Filter
                if (minScore > 0 && item.total_ball < minScore) return false;
                // Search Query
                if (searchQuery) {{
                    const s = searchQuery;
                    const matchFish = item.full_fish.toLowerCase().includes(s);
                    const matchId = item.id_number.includes(s);
                    const matchDir = item.direction_name.toLowerCase().includes(s);
                    const matchOta = (item.ota || '').toLowerCase().includes(s);
                    if (!matchFish && !matchId && !matchDir && !matchOta) return false;
                }}
                return true;
            }});

            currentPage = 1;
            document.getElementById('filteredStatsText').textContent = `Ko'rsatilmoqda: ${{currentFilteredData.length}} ta`;
            renderTable();
        }}

        function sortTable(column) {{
            if (sortColumn === column) {{
                sortAsc = !sortAsc;
            }} else {{
                sortColumn = column;
                sortAsc = true;
            }}

            currentFilteredData.sort((a, b) => {{
                let valA = a[column];
                let valB = b[column];
                if (typeof valA === 'string') {{
                    return sortAsc ? valA.localeCompare(valB) : valB.localeCompare(valA);
                }} else {{
                    return sortAsc ? (valA - valB) : (valB - valA);
                }}
            }});

            renderTable();
        }}

        function renderTable() {{
            const tbody = document.getElementById('tableBody');
            tbody.innerHTML = '';

            const startIdx = (currentPage - 1) * pageSize;
            const endIdx = Math.min(startIdx + pageSize, currentFilteredData.length);
            const pageData = currentFilteredData.slice(startIdx, endIdx);

            if (pageData.length === 0) {{
                tbody.innerHTML = `<tr><td colspan="12" style="text-align:center; padding:36px; color:var(--slate-600);">
                    <i data-lucide="search-x" style="width:36px; height:36px; margin-bottom:8px; stroke:var(--slate-400);"></i>
                    <div>Hech qanday ma'lumot topilmadi</div>
                </td></tr>`;
                lucide.createIcons();
                renderPagination(0);
                return;
            }}

            pageData.forEach(item => {{
                const tr = document.createElement('tr');
                
                let badgeHtml = '';
                if (item.section === 'Davlat granti') {{
                    badgeHtml = `<span class="badge badge-grant"><i data-lucide="check" size="12"></i> Davlat Granti</span>`;
                }} else if (item.section === 'To\'lov-kontrakt') {{
                    badgeHtml = `<span class="badge badge-kontrakt"><i data-lucide="file-check" size="12"></i> To'lov-kontrakt</span>`;
                }} else {{
                    badgeHtml = `<span class="badge badge-chegara"><i data-lucide="x" size="12"></i> Chegaradan tashqari</span>`;
                }}

                tr.innerHTML = `
                    <td style="text-align:center; font-weight:700; color:var(--slate-600);">${{item.global_tr}}</td>
                    <td><span class="id-badge">${{item.id_number}}</span></td>
                    <td><strong style="color:var(--slate-900); font-weight:700;">${{item.full_fish}}</strong></td>
                    <td><span style="font-weight:600; color:var(--slate-700);">${{item.direction_name}}</span></td>
                    <td>${{badgeHtml}}</td>
                    <td><span class="score-pill">${{item.total_ball.toFixed(2)}}</span></td>
                    <td style="color:var(--slate-600);">${{item.majburiy > 0 ? item.majburiy.toFixed(1) : '-'}}</td>
                    <td style="color:var(--slate-600);">${{item.fan1 > 0 ? item.fan1.toFixed(1) : '-'}}</td>
                    <td style="color:var(--slate-600);">${{item.fan2 > 0 ? item.fan2.toFixed(1) : '-'}}</td>
                    <td style="color:var(--slate-600);">${{item.imtiyoz > 0 ? item.imtiyoz.toFixed(1) : (item.milliy > 0 ? item.milliy.toFixed(1) : '-')}}</td>
                    <td><span style="font-size:0.8rem; font-weight:600; color:${{item.gender === 'Ayol' ? '#db2777' : '#2563eb'}};">${{item.gender}}</span></td>
                    <td style="text-align:center; color:var(--slate-600); font-size:0.8rem;">p.${{item.page}}</td>
                `;
                tbody.appendChild(tr);
            }});

            lucide.createIcons();
            renderPagination(currentFilteredData.length);
        }}

        function renderPagination(totalItems) {{
            const container = document.getElementById('paginationControls');
            container.innerHTML = '';
            const totalPages = Math.ceil(totalItems / pageSize);

            if (totalPages <= 1) return;

            // Prev button
            const prevBtn = document.createElement('button');
            prevBtn.className = 'page-btn';
            prevBtn.innerHTML = '<i data-lucide="chevron-left" size="16"></i>';
            prevBtn.disabled = currentPage === 1;
            prevBtn.onclick = () => {{ currentPage--; renderTable(); }};
            container.appendChild(prevBtn);

            // Pages
            let startPage = Math.max(1, currentPage - 2);
            let endPage = Math.min(totalPages, startPage + 4);
            if (endPage - startPage < 4) startPage = Math.max(1, endPage - 4);

            for (let p = startPage; p <= endPage; p++) {{
                const pageBtn = document.createElement('button');
                pageBtn.className = `page-btn ${{p === currentPage ? 'active' : ''}}`;
                pageBtn.textContent = p;
                pageBtn.onclick = () => {{ currentPage = p; renderTable(); }};
                container.appendChild(pageBtn);
            }}

            // Next button
            const nextBtn = document.createElement('button');
            nextBtn.className = 'page-btn';
            nextBtn.innerHTML = '<i data-lucide="chevron-right" size="16"></i>';
            nextBtn.disabled = currentPage === totalPages;
            nextBtn.onclick = () => {{ currentPage++; renderTable(); }};
            container.appendChild(nextBtn);

            lucide.createIcons();
        }}

        function changePageSize(val) {{
            pageSize = parseInt(val);
            currentPage = 1;
            renderTable();
        }}

        function exportToExcel() {{
            const exportRows = currentFilteredData.map(s => ({{
                'T/r': s.global_tr,
                'Abituriyent ID': s.id_number,
                'To\'liq F.I.SH': s.full_fish,
                'Familiyasi': s.fam,
                'Ismi': s.ism,
                'Otasining ismi': s.ota,
                'Ta\'lim Yo\'nalishi': s.direction_name,
                'Yo\'nalish Kodi': s.direction_code,
                'Qabul Holati': s.section,
                'Bo\'lim T/r': s.tr_in_section,
                'Umumiy Ball': s.total_ball,
                'Majburiy Fanlar': s.majburiy,
                '1-Fan': s.fan1,
                '2-Fan': s.fan2,
                'Milliy Sertifikat': s.milliy,
                'Chet Tili': s.chet_tili,
                'Kasbiy Imtihon': s.kasbiy,
                'Imtiyoz': s.imtiyoz,
                'Jinsi': s.gender,
                'PDF Sahifasi': s.page
            }}));

            const ws = XLSX.utils.json_to_sheet(exportRows);
            const wb = XLSX.utils.book_new();
            XLSX.utils.book_append_sheet(wb, ws, "Qabul Natijalari");
            XLSX.writeFile(wb, "Shahrisabz_JSST_Filtrlangan_Natijalar.xlsx");
        }}
    </script>
</body>
</html>
"""

    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    with open(index_path, 'w', encoding='utf-8') as f:
        f.write(html_content)
    print(f"✅ HTML Dashboard fayllari saqlandi: {html_path} va {index_path}")

def save_json_database(students, json_path):
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(students, f, ensure_ascii=False, indent=2)
    print(f"✅ JSON ma'lumotlar bazasi saqlandi: {json_path}")

def main():
    print("="*75)
    print("🚀 SHAHRISABZ JSST TAHLILI VA HISOBOT GENERATORI BOSHLANDI")
    print("="*75)
    
    if not os.path.exists(PDF_PATH):
        print(f"❌ Xatolik: {PDF_PATH} fayli topilmadi!")
        return

    # 1. Extract
    students = extract_pdf_data(PDF_PATH)
    
    # 2. Save JSON
    save_json_database(students, JSON_PATH)
    
    # 3. Generate Excel
    generate_excel_file(students, EXCEL_PATH)
    
    # 4. Generate HTML Dashboard
    generate_html_dashboard(students, HTML_PATH, INDEX_HTML_PATH)
    
    print("\n" + "="*75)
    print("🎉 BARCHA ISHLAR MUVAFFAQIYATLI YAKUNLANDI!")
    print(f"📁 Natijaviy fayllar:")
    print(f"  1. Excel bazasi: {EXCEL_PATH}")
    print(f"  2. Interaktiv Dashboard: {HTML_PATH}")
    print(f"  3. JSON bazasi: {JSON_PATH}")
    print("="*75)

if __name__ == '__main__':
    main()
