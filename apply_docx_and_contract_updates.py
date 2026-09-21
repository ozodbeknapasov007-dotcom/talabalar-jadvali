import openpyxl

wb = openpyxl.load_workbook('Talabalar_Toliq_Royxati.xlsx')
ws = wb.active

updates_unverified = {
    63: { # Rustamova Charos
        3: "Hamshiralik ishi - 3 yillik",
        5: "309",
        6: "2026-09-15 00:00:00",
        7: "O'ktam qizi",
        8: "Rustamova Charos O'ktam qizi",
        9: "Rustamova Charos O'ktam qizi",
        10: "AE4730990",
        11: "60810085680017",
        12: "22.10.2025",
        13: "08.10.2008",
        14: "Rustamova Charos O'ktam qizi",
        15: "UM 03728038",
        16: "https://e-shahodatnoma.uz/api/public/getcertpdf/4065248/3579915/251a81e0-79f2-4882-bc82-231a0a1ab956?preview=true",
        17: "4-sonli umumiy o'rta ta'lim maktabi",
        18: "Shahodatnoma",
        19: 2026,
        20: "880632608, 973360140",
        21: "TOPILDI",
        22: "Rustamova charos uktam qizi 309.docx faylidan to'liq kiritildi",
        24: "Pasport: MOS | Shahodatnoma: MOS",
        25: "TASDIQLANDI"
    },
    88: { # Maxmudova Musallam
        3: "Hamshiralik ishi - 3 yillik",
        5: "311",
        6: "2026-09-15 00:00:00",
        7: "Qamariddin qizi",
        8: "Maxmudova Musallam Qamariddin qizi",
        9: "Maxmudova Musallam Qamariddin qizi",
        10: "AE5297116",
        11: "61303085680079",
        12: "04.12.2025",
        13: "13.03.2008",
        14: "Maxmudova Musallam Qamariddin qizi",
        15: "UM 03728687",
        16: "https://e-shahodatnoma.uz/api/public/getcertpdf/3946429/3458161/ab44c5f0-66ab-4b7a-b270-81c64eb7130e?preview=true",
        17: "23-sonli umumiy o'rta ta'lim maktabi",
        18: "Shahodatnoma",
        19: 2026,
        20: "505899948",
        21: "TOPILDI",
        22: "Maxmudova Musallam qamariddin qizi 311.docx faylidan to'liq kiritildi",
        24: "Pasport: MOS | Shahodatnoma: MOS",
        25: "TASDIQLANDI"
    },
    92: { # Qodirova Shabnam
        3: "Hamshiralik ishi - 3 yillik",
        5: "137",
        6: "2026-08-04 00:00:00",
        7: "O'tkir qizi",
        8: "Qodirova Shabnam O'tkir qizi",
        9: "Qodirova Shabnam O'tkir qizi",
        10: "AE5242435",
        11: "61807085680045",
        12: "01.12.2025",
        13: "18.07.2008",
        14: "Qodirova Shabnam O'tkir qizi",
        15: "UM 03729278",
        16: "https://e-shahodatnoma.uz/api/public/getcertpdf/3934766/3452539/98763bcf-ef33-4ddb-b3e1-ddf8c7990688?preview=true",
        17: "44-sonli umumiy o'rta ta'lim maktabi",
        18: "Shahodatnoma",
        19: 2026,
        20: "88-827-07-18, 88-386-04-85",
        21: "TOPILDI",
        22: "Qodirova Shabnam 137.docx faylidan to'liq kiritildi",
        24: "Pasport: MOS | Shahodatnoma: MOS",
        25: "TASDIQLANDI"
    },
    99: { # Xo'jaqulova Navbahor
        3: "Hamshiralik ishi - 3 yillik",
        5: "302",
        6: "2026-09-14 00:00:00",
        7: "Azim qizi",
        8: "Xo'jaqulova Navbahor Azim qizi",
        9: "Xo'jaqulova Navbahor Azim qizi",
        10: "AD4186543",
        11: "62104025680040",
        12: "07.08.2023",
        13: "21.04.2002",
        14: "Xo'jaqulova Navbahor Azim qizi",
        15: "UM 0086260",
        17: "63-sonli umumiy o'rta ta'lim maktabi",
        18: "Shahodatnoma",
        19: 2019,
        20: "883851378, 884188278",
        21: "TOPILDI",
        22: "Xo'jaqulova Navbaxor Azim qizi 302.docx faylidan to'liq kiritildi",
        24: "Pasport: MOS | Shahodatnoma: MOS",
        25: "TASDIQLANDI"
    },
    100: { # Xolmurodova Dilnura
        3: "Hamshiralik ishi - 3 yillik",
        5: "307",
        6: "2026-09-15 00:00:00",
        7: "Dilmurod qizi",
        8: "Xolmurodova Dilnura Dilmurod qizi",
        9: "Xolmurodova Dilnura Dilmurod qizi",
        10: "AE1516812",
        11: "60110085590193",
        12: "06.02.2025",
        13: "01.10.2008",
        14: "Xolmurodova Dilnura Dilmurod qizi",
        15: "UM 03755333",
        16: "https://e-shahodatnoma.uz/api/public/getcertpdf/3646753/3265997/7168a42b-3c1c-4933-ac3c-05c228850986?preview=true",
        17: "60-sonli umumiy o'rta ta'lim maktabi",
        18: "Shahodatnoma",
        19: 2026,
        20: "879340301",
        21: "TOPILDI",
        22: "Xolmurodova Dilnura Dilmurod qizi 307.docx faylidan to'liq kiritildi",
        24: "Pasport: MOS | Shahodatnoma: MOS",
        25: "TASDIQLANDI"
    },
    169: { # Murtazayeva Kamola
        5: "25",
        6: "2026-06-29 00:00:00",
        7: "Iskandar qizi",
        8: "Murtazayeva Kamola Iskandar qizi",
        11: "42708915590028",
        17: "Ixtisoslashgan maktab",
        19: 2009,
        20: "88-803-91-93, 88-726-14-12",
        21: "TOPILDI",
        22: "Murtazayeva Kamola 25.docx faylidan to'liq kiritildi",
        24: "Pasport: MOS | Shahodatnoma: MOS",
        25: "TASDIQLANDI"
    }
}

# Contract updates for already-verified students (DO NOT touch any existing verified columns)
contract_only_updates = {
    22: {5: "294", 6: "2026-09-08 00:00:00"}, # Ilhomova Zuhra
    55: {5: "282", 6: "2026-09-07 00:00:00"}, # Mustafayeva Madina
    61: {5: "291", 6: "2026-09-08 00:00:00"}, # Rayimova Nozliya
    68: {5: "299", 6: "2026-09-10 00:00:00"}, # Temirova Rushana
    77: {5: "274", 6: "2026-09-07 00:00:00", 20: "93-019-06-86"}, # Abdixalilova Shohzoda
    83: {5: "275", 6: "2026-09-07 00:00:00", 20: "91-005-49-08, 93-553-33-51"}, # Hamroqulova Marjona
    86: {5: "293", 6: "2026-09-08 00:00:00"}, # Karimova Shaxina
    121: {5: "289", 6: "2026-09-08 00:00:00"}, # Ne'matova Sevinch
    135: {5: "286", 6: "2026-09-07 00:00:00", 20: "97-110-03-25, 91-071-06-83"}, # Abdixoliqova Elmira
    157: {5: "03", 6: "2026-06-19 00:00:00", 20: "93-152-25-77, 77-086-77-25"}  # Sayfillayeva Shodiyona
}

print("Applying unverified updates...")
for row_idx, col_dict in updates_unverified.items():
    fish = ws.cell(row_idx, 8).value or ws.cell(row_idx, 2).value
    print(f"  Row {row_idx}: {fish}")
    for c_idx, val in col_dict.items():
        ws.cell(row_idx, c_idx).value = val

print("\nApplying contract-only updates for verified students...")
for row_idx, col_dict in contract_only_updates.items():
    fish = ws.cell(row_idx, 8).value or ws.cell(row_idx, 2).value
    print(f"  Row {row_idx}: {fish}")
    for c_idx, val in col_dict.items():
        # Only set if cell is empty or if it's shnum/sana
        curr_val = ws.cell(row_idx, c_idx).value
        if c_idx in [5, 6] or not curr_val:
            ws.cell(row_idx, c_idx).value = val

wb.save('Talabalar_Toliq_Royxati.xlsx')
print("\nTalabalar_Toliq_Royxati.xlsx muvaffaqiyatli saqlandi!")
